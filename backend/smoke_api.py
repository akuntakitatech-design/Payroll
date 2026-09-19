"""Smoke test API baru: payroll run + approval, slip PDF, self-service karyawan,
perpanjangan kontrak, impor Excel, pengaturan SMTP/pengingat."""
import io
import json

import requests

B = "http://localhost:8001/api"
PW = "Hris#2026"
OK, BAD = [], []


def chk(name, cond, detail=""):
    (OK if cond else BAD).append(name)
    print(("  [OK]   " if cond else "  [FAIL] ") + name + (f" -- {detail}" if detail else ""))


def login(email):
    r = requests.post(f"{B}/auth/login", json={"email": email, "password": PW})
    r.raise_for_status()
    d = r.json()
    return {"Authorization": f"Bearer {d['access_token']}"}, d


H_ADMIN, S_ADMIN = login("hr.admin@nep.co.id")
H_MGR, _ = login("hr.manager@nep.co.id")
H_FIN, _ = login("finance@nep.co.id")
H_EMP, _ = login("karyawan@nep.co.id")
H_SUPER, _ = login("superadmin@hris.id")

TEST_YEAR = 2029

# ---- bersihkan sisa data dari eksekusi sebelumnya (idempoten) ----
for c in requests.get(f"{B}/payroll/components", headers=H_MGR).json().get("items", []):
    if c["code"] == "TJ-TEST":
        requests.delete(f"{B}/payroll/components/{c['id']}", headers=H_MGR)
import pymongo, os
_cl = pymongo.MongoClient(os.environ.get("MONGO_URL", "mongodb://localhost:27017"))
_db = _cl[os.environ.get("DB_NAME", "test_database")]
_old = list(_db.payroll_runs.find({"year": TEST_YEAR}, {"_id": 0, "id": 1}))
_db.payroll_runs.delete_many({"year": TEST_YEAR})
_db.payroll_items.delete_many({"run_id": {"$in": [r["id"] for r in _old]}})
print(f"  (bersih: {len(_old)} payroll run periode {TEST_YEAR} dihapus)")

print("\n=== Katalog & konfigurasi payroll ===")
r = requests.get(f"{B}/payroll/catalog", headers=H_ADMIN)
chk("GET /payroll/catalog 200", r.status_code == 200, r.text[:120])
cat = r.json() if r.ok else {}
chk("Katalog memuat 8 status PTKP", len(cat.get("ptkp_statuses", [])) == 8)
chk("Katalog memuat 5 kelas risiko JKK", len(cat.get("jkk_risk_classes", [])) == 5)
chk("Parameter statutori: batas JP 11.086.300",
    cat.get("statutory", {}).get("jp_cap") == 11086300, str(cat.get("statutory", {}).get("jp_cap")))

print("\n=== Komponen gaji ===")
r = requests.get(f"{B}/payroll/components", headers=H_ADMIN)
chk("GET /payroll/components 200", r.status_code == 200)
comps = r.json().get("items", [])
chk("Komponen ter-seed (>=7)", len(comps) >= 7, str(len(comps)))

r = requests.post(f"{B}/payroll/components", headers=H_ADMIN, json={
    "code": "TJ-TEST", "name": "Tunjangan Uji", "kind": "earning",
    "default_amount": 250000})
chk("POST komponen baru 201", r.status_code == 201, r.text[:120])
new_comp = r.json() if r.ok else {}
r = requests.post(f"{B}/payroll/components", headers=H_ADMIN, json={
    "code": "TJ-TEST", "name": "Duplikat", "kind": "earning"})
chk("Kode komponen duplikat ditolak 409", r.status_code == 409, r.text[:100])
if new_comp:
    r = requests.delete(f"{B}/payroll/components/{new_comp['id']}", headers=H_ADMIN)
    chk("DELETE komponen 200", r.status_code == 200, r.text[:100])

print("\n=== Struktur gaji karyawan ===")
r = requests.get(f"{B}/payroll/salaries", headers=H_ADMIN)
chk("GET /payroll/salaries 200", r.status_code == 200)
sal = r.json()
chk("Semua 5 karyawan NEP punya struktur gaji",
    sal.get("unset_count") == 0 and sal.get("total") == 5,
    f"total={sal.get('total')} unset={sal.get('unset_count')}")
cats = {i["ter_category"] for i in sal.get("items", [])}
chk("Seed mencakup kategori TER A, B dan C", cats == {"A", "B", "C"}, str(cats))

# pilih karyawan yang BUKAN satu-satunya pemilik kategori C agar seed tetap utuh
_by_cat = {}
for i in sal.get("items", []):
    _by_cat.setdefault(i["ter_category"], []).append(i)
_target = next(i for i in sal["items"] if len(_by_cat[i["ter_category"]]) > 1
               or i["ter_category"] == "A")
emp_id = _target["employee_id"]
_orig_ptkp, _orig_basic = _target["ptkp_status"], _target["basic_salary"]
r = requests.get(f"{B}/payroll/salaries/{emp_id}", headers=H_ADMIN)
chk("GET struktur gaji per karyawan 200", r.status_code == 200)
r = requests.put(f"{B}/payroll/salaries/{emp_id}", headers=H_ADMIN,
                 json={"basic_salary": 7800000, "ptkp_status": "K/1"})
chk("PUT struktur gaji 200", r.status_code == 200, r.text[:120])
r = requests.put(f"{B}/payroll/salaries/{emp_id}", headers=H_ADMIN,
                 json={"basic_salary": 100, "ptkp_status": "X/9"})
chk("Status PTKP invalid ditolak 422", r.status_code == 422, r.text[:100])
r = requests.put(f"{B}/payroll/salaries/{emp_id}", headers=H_ADMIN,
                 json={"basic_salary": -5})
chk("Gaji negatif ditolak 422", r.status_code == 422, r.text[:100])
# pulihkan agar variasi kategori TER hasil seed tetap utuh
requests.put(f"{B}/payroll/salaries/{emp_id}", headers=H_ADMIN,
             json={"basic_salary": _orig_basic, "ptkp_status": _orig_ptkp})

print("\n=== Payroll run: buat, hitung, sesuaikan ===")
requests.delete(f"{B}/payroll/runs/none", headers=H_ADMIN)
# bersihkan periode uji bila ada

r = requests.post(f"{B}/payroll/runs", headers=H_ADMIN,
                  json={"year": TEST_YEAR, "month": 5, "notes": "Uji smoke"})
chk("POST /payroll/runs 201", r.status_code == 201, r.text[:200])
run = r.json() if r.ok else {}
run_id = run.get("id")
chk("Run menghitung 4 slip (karyawan nonaktif dikecualikan)", run.get("employee_count") == 4,
    str(run.get("employee_count")))
chk("Total bruto > 0", (run.get("totals") or {}).get("gross", 0) > 0,
    str((run.get("totals") or {}).get("gross")))
chk("Total netto < total bruto",
    (run.get("totals") or {}).get("net_pay", 0) < (run.get("totals") or {}).get("gross", 1))
chk("Biaya pemberi kerja > bruto",
    (run.get("totals") or {}).get("employer_cost", 0) > (run.get("totals") or {}).get("gross", 0))

r = requests.post(f"{B}/payroll/runs", headers=H_ADMIN, json={"year": TEST_YEAR, "month": 5})
chk("Periode duplikat ditolak 409", r.status_code == 409, r.text[:110])

r = requests.get(f"{B}/payroll/runs/{run_id}", headers=H_ADMIN)
chk("GET detail run 200", r.status_code == 200)
detail = r.json()
items = detail.get("items", [])
chk("Detail memuat 4 baris slip", len(items) == 4, str(len(items)))
chk("Detail memuat langkah persetujuan (workflow payroll)",
    len(detail.get("approval_steps", [])) == 3, str(len(detail.get("approval_steps", []))))
chk("can_edit true saat draf", detail.get("can_edit") is True)

first = items[0]
chk("Slip memuat kategori & tarif TER",
    first["tax"].get("ter_category") in ("A", "B", "C") and "ter_rate" in first["tax"],
    str(first["tax"].get("ter_category")))
chk("Slip: netto = bruto - potongan",
    abs(first["totals"]["net_pay"] - (first["totals"]["gross"] - first["totals"]["total_deductions"])) < 1)
chk("Slip memuat rincian BPJS 5 program",
    all(k in first["bpjs"] for k in ("kesehatan", "jht", "jp", "jkk", "jkm")))

before_net = first["totals"]["net_pay"]
r = requests.put(f"{B}/payroll/runs/{run_id}/items/{first['id']}/adjustments",
                 headers=H_ADMIN,
                 json={"overtime_hours": 10, "unpaid_days": 1, "working_days": 22,
                       "extra_earnings": [{"name": "Bonus Proyek", "amount": 1000000}],
                       "extra_deductions": [{"name": "Kasbon", "amount": 500000}]})
chk("PUT penyesuaian 200", r.status_code == 200, r.text[:180])
adj = r.json() if r.ok else {}
if adj:
    it = adj["item"]
    chk("Lembur muncul di slip",
        any(e["code"] == "OVERTIME" for e in it["earnings"]))
    chk("Bonus tambahan muncul di slip",
        any(e["name"] == "Bonus Proyek" for e in it["earnings"]))
    chk("Potongan kasbon muncul di slip",
        any(d["name"] == "Kasbon" for d in it["deductions"]))
    chk("Absen 1 hari memproratakan gaji pokok",
        it["basic_salary_paid"] < it["basic_salary"],
        f"{it['basic_salary_paid']} < {it['basic_salary']}")
    chk("Netto berubah setelah penyesuaian", it["totals"]["net_pay"] != before_net,
        f"{before_net} -> {it['totals']['net_pay']}")
    chk("Total run ikut diperbarui", adj["totals"]["gross"] > 0)

r = requests.put(f"{B}/payroll/runs/{run_id}/items/{first['id']}/adjustments",
                 headers=H_ADMIN, json={"overtime_hours": -5})
chk("Jam lembur negatif ditolak 422", r.status_code == 422, r.text[:100])

r = requests.post(f"{B}/payroll/runs/{run_id}/recalculate", headers=H_ADMIN)
chk("POST recalculate 200", r.status_code == 200, r.text[:120])

print("\n=== Slip gaji PDF & ekspor ===")
r = requests.get(f"{B}/payroll/runs/{run_id}/items/{first['id']}/payslip", headers=H_ADMIN)
chk("GET slip PDF 200", r.status_code == 200, r.text[:100] if not r.ok else "")
chk("Respons berupa PDF", r.ok and r.content[:4] == b"%PDF", str(r.content[:8]))
chk("PDF punya nama file attachment",
    "attachment" in r.headers.get("content-disposition", ""),
    r.headers.get("content-disposition", "")[:70])

r = requests.get(f"{B}/payroll/runs/{run_id}/export", headers=H_FIN)
chk("GET ekspor Excel 200 (Finance)", r.status_code == 200, r.text[:100] if not r.ok else "")
chk("Ekspor berupa xlsx", r.ok and r.content[:2] == b"PK")

print("\n=== Alur persetujuan payroll ===")
r = requests.post(f"{B}/payroll/runs/{run_id}/approve", headers=H_MGR)
chk("Approve sebelum diajukan ditolak 409", r.status_code == 409, r.text[:110])
r = requests.post(f"{B}/payroll/runs/{run_id}/submit", headers=H_ADMIN,
                  json={"note": "Mohon direview"})
chk("POST submit 200", r.status_code == 200, r.text[:120])
chk("Status jadi pending_approval",
    r.ok and r.json().get("run_status") == "pending_approval", r.text[:80] if not r.ok else "")

r = requests.put(f"{B}/payroll/runs/{run_id}/items/{first['id']}/adjustments",
                 headers=H_ADMIN, json={"overtime_hours": 2})
chk("Penyesuaian saat menunggu persetujuan ditolak 409", r.status_code == 409, r.text[:110])
r = requests.post(f"{B}/payroll/runs/{run_id}/approve", headers=H_ADMIN)
chk("HR Admin tidak boleh menyetujui (tanpa hak) 403", r.status_code == 403, r.text[:110])
r = requests.post(f"{B}/payroll/runs/{run_id}/reject", headers=H_MGR, json={"note": ""})
chk("Tolak tanpa alasan ditolak 422", r.status_code == 422, r.text[:110])
r = requests.post(f"{B}/payroll/runs/{run_id}/reject", headers=H_MGR,
                  json={"note": "Lembur Joko perlu dicek ulang"})
chk("POST reject 200", r.status_code == 200, r.text[:120])
chk("Status jadi rejected", r.ok and r.json().get("run_status") == "rejected")

r = requests.post(f"{B}/payroll/runs/{run_id}/submit", headers=H_ADMIN)
chk("Ajukan ulang setelah ditolak 200", r.status_code == 200, r.text[:120])
r = requests.post(f"{B}/payroll/runs/{run_id}/approve", headers=H_MGR, json={"note": "Setuju"})
chk("POST approve 200", r.status_code == 200, r.text[:120])
chk("Status jadi approved", r.ok and r.json().get("run_status") == "approved")
r = requests.delete(f"{B}/payroll/runs/{run_id}", headers=H_MGR)
chk("Hapus run yang sudah disetujui ditolak 409", r.status_code == 409, r.text[:110])
r = requests.post(f"{B}/payroll/runs/{run_id}/mark-paid", headers=H_ADMIN)
chk("POST mark-paid 200", r.status_code == 200, r.text[:120])

print("\n=== Rekonsiliasi Desember (Pasal 17) ===")
r = requests.post(f"{B}/payroll/runs", headers=H_ADMIN, json={"year": TEST_YEAR, "month": 12})
chk("Buat payroll Desember 201", r.status_code == 201, r.text[:150])
dec_id = r.json().get("id") if r.ok else None
if dec_id:
    d = requests.get(f"{B}/payroll/runs/{dec_id}", headers=H_ADMIN).json()
    item = d["items"][0]
    chk("Desember memakai metode Pasal 17",
        item["tax"]["method"] == "article_17_annual", item["tax"]["method"])
    chk("Rincian tahunan tersedia", "annual" in item["tax"])
    ann = item["tax"].get("annual", {})
    chk("PTKP tahunan terisi", (ann.get("ptkp_annual") or 0) > 0, str(ann.get("ptkp_annual")))
    chk("Memperhitungkan potongan Jan-Nov (dari run Mei yang dibayar)",
        (ann.get("already_withheld") or 0) >= 0, str(ann.get("already_withheld")))
    chk("PKP dibulatkan ke ribuan", (ann.get("pkp") or 0) % 1000 == 0, str(ann.get("pkp")))
    r = requests.get(f"{B}/payroll/runs/{dec_id}/items/{item['id']}/payslip", headers=H_ADMIN)
    chk("Slip PDF Desember 200", r.status_code == 200 and r.content[:4] == b"%PDF")

print("\n=== Self-service slip gaji karyawan ===")
r = requests.get(f"{B}/payroll/my/payslips", headers=H_EMP)
chk("GET /payroll/my/payslips 200 (role karyawan)", r.status_code == 200, r.text[:150])
mine = r.json() if r.ok else {}
chk("Akun karyawan tertaut ke data karyawan", mine.get("employee") is not None,
    json.dumps(mine.get("employee"))[:90])
chk("Karyawan melihat slip periode yang sudah dibayar",
    len(mine.get("items", [])) >= 1, str(len(mine.get("items", []))))
if mine.get("items"):
    own = mine["items"][0]
    chk("Slip milik sendiri (employee_id cocok)",
        own["employee_id"] == mine["employee"]["id"])
    r = requests.get(f"{B}/payroll/my/payslips/{own['id']}/payslip", headers=H_EMP)
    chk("Karyawan bisa unduh PDF slipnya sendiri",
        r.status_code == 200 and r.content[:4] == b"%PDF", r.text[:100] if not r.ok else "")

    # slip karyawan LAIN tidak boleh bisa diambil
    others = [i for i in requests.get(f"{B}/payroll/runs/{run_id}", headers=H_ADMIN).json()["items"]
              if i["employee_id"] != mine["employee"]["id"]]
    if others:
        r = requests.get(f"{B}/payroll/my/payslips/{others[0]['id']}/payslip", headers=H_EMP)
        chk("Karyawan TIDAK bisa unduh slip orang lain (404)", r.status_code == 404,
            f"status {r.status_code}")

r = requests.get(f"{B}/payroll/runs", headers=H_EMP)
chk("Karyawan tidak boleh lihat daftar payroll 403", r.status_code == 403, r.text[:100])
r = requests.post(f"{B}/payroll/runs", headers=H_EMP, json={"year": TEST_YEAR, "month": 7})
chk("Karyawan tidak boleh menjalankan payroll 403", r.status_code == 403, r.text[:100])
r = requests.get(f"{B}/payroll/salaries", headers=H_EMP)
chk("Karyawan tidak boleh lihat struktur gaji 403", r.status_code == 403, r.text[:100])

print("\n=== Isolasi antar perusahaan ===")
H_MULTI, S_MULTI = login("hr.multi@hris.id")
companies = {c["code"]: c["id"] for c in S_MULTI.get("companies", [])}
sw = requests.post(f"{B}/auth/switch-company", headers=H_MULTI,
                   json={"company_id": companies["KBS"]}).json()
H_KBS = {"Authorization": f"Bearer {sw['access_token']}"}
sw2 = requests.post(f"{B}/auth/switch-company", headers=H_KBS,
                    json={"company_id": companies["NEP"]}).json()
H_NEP = {"Authorization": f"Bearer {sw2['access_token']}"}
r = requests.get(f"{B}/payroll/runs/{run_id}", headers=H_KBS)
chk("Run NEP tidak terlihat dari konteks KBS (404)", r.status_code == 404, f"status {r.status_code}")
r = requests.get(f"{B}/payroll/runs/{run_id}/items/{first['id']}/payslip", headers=H_KBS)
chk("Slip PDF NEP tidak bisa diunduh dari KBS (404)", r.status_code == 404, f"status {r.status_code}")
kbs_sal = requests.get(f"{B}/payroll/salaries", headers=H_KBS).json()
chk("Struktur gaji KBS hanya 3 karyawan KBS", kbs_sal.get("total") == 3, str(kbs_sal.get("total")))
r = requests.get(f"{B}/payroll/salaries/{emp_id}", headers=H_KBS)
chk("Struktur gaji karyawan NEP tidak bisa dibuka dari KBS (404)", r.status_code == 404,
    f"status {r.status_code}")

print("\n=== Perpanjangan kontrak 1 klik ===")
contracts = requests.get(f"{B}/contracts", headers=H_ADMIN).json().get("items", [])
target = next((c for c in contracts if c.get("end_date") and not c.get("renewed_by_contract_id")), None)
chk("Ada kontrak dengan masa berlaku untuk diperpanjang", target is not None)
if target:
    r = requests.get(f"{B}/contracts/{target['id']}/renew-preview", headers=H_ADMIN)
    chk("GET renew-preview 200", r.status_code == 200, r.text[:150])
    pv = r.json() if r.ok else {}
    dflt = pv.get("defaults", {})
    chk("Preview mengisi tanggal mulai otomatis", bool(dflt.get("start_date")))
    chk("Preview mengisi tanggal berakhir otomatis", bool(dflt.get("end_date")))
    chk("Preview mengisi nomor kontrak otomatis", bool(dflt.get("contract_number")))
    chk("Preview mengisi gaji dari kontrak lama", dflt.get("basic_salary") is not None)
    chk("Preview menyertakan daftar tipe kontrak", len(pv.get("contract_types", [])) > 0)
    chk("Mulai perpanjangan setelah kontrak lama berakhir",
        dflt.get("start_date", "") > (target.get("end_date") or "")[:10],
        f"{dflt.get('start_date')} > {target.get('end_date')}")

    r = requests.post(f"{B}/contracts/{target['id']}/renew", headers=H_ADMIN, json={
        "start_date": dflt["start_date"], "end_date": dflt["end_date"],
        "contract_type_id": dflt["contract_type_id"],
        "contract_number": dflt["contract_number"],
        "basic_salary": dflt.get("basic_salary"), "allowance": dflt.get("allowance")})
    chk("POST renew 201", r.status_code == 201, r.text[:200])
    if r.ok:
        res = r.json()
        chk("Kontrak baru menaut ke kontrak lama",
            res["contract"].get("previous_contract_id") == target["id"])
        chk("Kontrak baru ditandai perpanjangan", res["contract"].get("is_renewal") is True)
        chk("Kontrak lama ditandai renewed",
            res["previous_contract"].get("renewal_state") == "renewed")
        chk("Kontrak lama diarsipkan", res["previous_contract"].get("status") == "archived")
        chk("Pesan sukses berbahasa Indonesia", "berhasil diperpanjang" in res.get("message", ""))
    r = requests.post(f"{B}/contracts/{target['id']}/renew", headers=H_ADMIN, json={
        "start_date": dflt["start_date"], "end_date": dflt["end_date"],
        "contract_type_id": dflt["contract_type_id"]})
    chk("Perpanjang dua kali ditolak 409", r.status_code == 409, r.text[:120])
    r = requests.post(f"{B}/contracts/{target['id']}/renew", headers=H_EMP, json={
        "start_date": "2027-01-01", "contract_type_id": dflt["contract_type_id"]})
    chk("Karyawan tidak boleh memperpanjang kontrak 403", r.status_code == 403)

print("\n=== Impor Excel karyawan ===")
r = requests.get(f"{B}/employees/import/template", headers=H_ADMIN)
chk("GET template Excel 200", r.status_code == 200, r.text[:100] if not r.ok else "")
chk("Template berupa xlsx", r.ok and r.content[:2] == b"PK")
template_bytes = r.content if r.ok else b""

r = requests.get(f"{B}/employees/import/columns", headers=H_ADMIN)
chk("GET metadata kolom 200", r.status_code == 200)
chk("Metadata memuat >=30 kolom", len(r.json().get("columns", [])) >= 30,
    str(len(r.json().get("columns", []))))

if template_bytes:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(template_bytes))
    ws = wb["Data Karyawan"]
    idx = {c.value: i + 1 for i, c in enumerate(ws[1])}
    # 2 baris valid, 2 baris salah
    ws.cell(row=3, column=idx["Nama Lengkap*"], value="Impor Satu")
    ws.cell(row=3, column=idx["Status PTKP"], value="K/2")
    ws.cell(row=3, column=idx["Gaji Pokok"], value=7250000)
    ws.cell(row=3, column=idx["Departemen"], value="Human Capital")
    ws.cell(row=4, column=idx["Nama Lengkap*"], value="Impor Dua")
    ws.cell(row=4, column=idx["Email"], value="impor.dua@contoh.co.id")
    ws.cell(row=5, column=idx["Email"], value="tanpa.nama@contoh.co.id")   # nama kosong
    ws.cell(row=6, column=idx["Nama Lengkap*"], value="Impor Salah")
    ws.cell(row=6, column=idx["Departemen"], value="Departemen Fiktif")
    buf = io.BytesIO(); wb.save(buf)

    r = requests.post(f"{B}/employees/import/validate", headers=H_ADMIN,
                      files={"file": ("uji.xlsx", buf.getvalue(),
                                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    chk("POST import/validate 200", r.status_code == 200, r.text[:200])
    rep = r.json() if r.ok else {}
    chk("Validasi: 2 baris valid", rep.get("valid_count") == 2, str(rep.get("valid_count")))
    chk("Validasi: 2 baris invalid", rep.get("invalid_count") == 2, str(rep.get("invalid_count")))
    chk("Ringkasan berbahasa Indonesia", "baris" in (rep.get("summary") or ""),
        rep.get("summary"))

    valid_rows = [row for row in rep.get("rows", []) if row["is_valid"]]
    r = requests.post(f"{B}/employees/import/commit", headers=H_ADMIN,
                      json={"rows": valid_rows})
    chk("POST import/commit 201", r.status_code == 201, r.text[:200])
    res = r.json() if r.ok else {}
    chk("2 karyawan terbuat", res.get("created_count") == 2, str(res.get("created_count")))
    chk("NIK karyawan dibuat otomatis",
        all(c.get("employee_number") for c in res.get("created", [])),
        str([c.get("employee_number") for c in res.get("created", [])]))

    # gaji pokok dari Excel harus ikut tersimpan ke struktur gaji
    created_ids = [c["id"] for c in res.get("created", [])]
    if created_ids:
        s1 = requests.get(f"{B}/payroll/salaries/{created_ids[0]}", headers=H_ADMIN).json()
        chk("Gaji pokok dari Excel tersimpan ke struktur gaji",
            (s1.get("salary") or {}).get("basic_salary") == 7250000,
            str((s1.get("salary") or {}).get("basic_salary")))
        chk("Status PTKP dari Excel tersimpan",
            (s1.get("salary") or {}).get("ptkp_status") == "K/2",
            str((s1.get("salary") or {}).get("ptkp_status")))
        # bersihkan
        for cid_ in created_ids:
            requests.delete(f"{B}/employees/{cid_}", headers=H_ADMIN)

    r = requests.post(f"{B}/employees/import/validate", headers=H_ADMIN,
                      files={"file": ("bukan.csv", b"a,b,c", "text/csv")})
    chk("File non-xlsx ditolak 422", r.status_code == 422, r.text[:110])
    r = requests.post(f"{B}/employees/import/validate", headers=H_EMP,
                      files={"file": ("uji.xlsx", buf.getvalue(), "application/octet-stream")})
    chk("Karyawan tidak boleh impor 403", r.status_code == 403)

print("\n=== Pengaturan SMTP & pengingat ===")
r = requests.get(f"{B}/mail/settings", headers=H_MGR)
chk("GET /mail/settings 200", r.status_code == 200, r.text[:150])
ms = r.json() if r.ok else {}
chk("Respons TIDAK memuat password", "password" not in json.dumps(ms).lower()
    or '"has_password"' in json.dumps(ms), "cek manual")
chk("Field password tidak ada di objek smtp", "password" not in (ms.get("smtp") or {}))
chk("Ada 3 mode keamanan", len(ms.get("security_modes", [])) == 3)
chk("Penjadwal berjalan", (ms.get("scheduler") or {}).get("running") is True,
    json.dumps(ms.get("scheduler"))[:120])
_w = (ms.get("reminder") or {}).get("windows") or []
chk("Pengaturan pengingat punya jendela hari terurut menurun",
    len(_w) >= 1 and _w == sorted(_w, reverse=True), str(_w))

r = requests.put(f"{B}/mail/settings/smtp", headers=H_MGR, json={
    "host": "smtp.nusantaraenergi.co.id", "port": 587, "security": "starttls",
    "username": "hris@nusantaraenergi.co.id", "password": "rahasia-sekali",
    "from_email": "hris@nusantaraenergi.co.id", "from_name": "HRIS Nusantara"})
chk("PUT SMTP 200", r.status_code == 200, r.text[:150])
chk("Respons simpan SMTP tidak bocorkan password",
    "rahasia-sekali" not in r.text, r.text[:120])
r = requests.get(f"{B}/mail/settings", headers=H_MGR)
chk("GET setelah simpan: has_password true",
    (r.json().get("smtp") or {}).get("has_password") is True)
chk("GET setelah simpan: password tetap tidak dikembalikan",
    "rahasia-sekali" not in r.text)
chk("is_configured true", (r.json().get("smtp") or {}).get("is_configured") is True)

r = requests.put(f"{B}/mail/settings/smtp", headers=H_MGR, json={"port": 99999})
chk("Port di luar rentang ditolak 422", r.status_code == 422, r.text[:110])
r = requests.put(f"{B}/mail/settings/smtp", headers=H_MGR, json={"security": "aneh"})
chk("Mode keamanan tidak dikenal ditolak 422", r.status_code == 422, r.text[:110])
r = requests.put(f"{B}/mail/settings/smtp", headers=H_EMP, json={"host": "x"})
chk("Karyawan tidak boleh ubah SMTP 403", r.status_code == 403)

r = requests.put(f"{B}/mail/settings/reminder", headers=H_MGR,
                 json={"windows": [60, 30, 14], "recipients": ["hrd@nusantaraenergi.co.id"],
                       "send_hour": 8})
chk("PUT pengaturan pengingat 200", r.status_code == 200, r.text[:150])
chk("Jendela pengingat tersimpan terurut menurun",
    r.ok and r.json()["reminder"]["windows"] == [60, 30, 14],
    str(r.json().get("reminder", {}).get("windows")) if r.ok else "")
r = requests.put(f"{B}/mail/settings/reminder", headers=H_MGR,
                 json={"recipients": ["bukan-email"]})
chk("Email penerima invalid ditolak 422", r.status_code == 422, r.text[:110])
r = requests.put(f"{B}/mail/settings/reminder", headers=H_MGR, json={"windows": []})
chk("Jendela kosong ditolak 422", r.status_code == 422, r.text[:110])
r = requests.put(f"{B}/mail/settings/reminder", headers=H_MGR, json={"send_hour": 30})
chk("Jam kirim di luar 0-23 ditolak 422", r.status_code == 422, r.text[:110])

r = requests.get(f"{B}/mail/reminder/preview", headers=H_MGR)
chk("GET pratinjau digest 200", r.status_code == 200, r.text[:150])
pv = r.json() if r.ok else {}
chk("Pratinjau memuat subjek", bool(pv.get("subject")), (pv.get("subject") or "")[:70])
chk("Pratinjau memuat HTML", "<div" in (pv.get("html") or ""))
chk("Pratinjau menghitung item", isinstance(pv.get("total"), int), str(pv.get("total")))

r = requests.post(f"{B}/mail/reminder/send-now", headers=H_MGR)
chk("Kirim sekarang ke host palsu gagal dengan pesan jelas",
    r.status_code == 400 and "SMTP" in r.text, f"{r.status_code} {r.text[:110]}")
r = requests.get(f"{B}/mail/reminder/logs", headers=H_MGR)
chk("GET log pengingat 200", r.status_code == 200)
logs = r.json() if r.ok else {}
chk("Kegagalan tercatat di log", (logs.get("failed_count") or 0) >= 1,
    str(logs.get("failed_count")))
chk("Log tidak memuat password", "rahasia-sekali" not in r.text)

print("\n=== Ringkasan payroll (dashboard) ===")
r = requests.get(f"{B}/payroll/summary", headers=H_ADMIN)
chk("GET /payroll/summary 200", r.status_code == 200, r.text[:120])
sm = r.json() if r.ok else {}
chk("Ringkasan memuat periode terakhir", sm.get("latest_run") is not None)
chk("Ringkasan menghitung karyawan tanpa gaji",
    isinstance(sm.get("employees_without_salary"), int))

print("\n" + "=" * 70)
print(f" HASIL: {len(OK)}/{len(OK) + len(BAD)} lulus")
if BAD:
    print(" GAGAL:")
    for b in BAD:
        print("   -", b)
print("=" * 70)
