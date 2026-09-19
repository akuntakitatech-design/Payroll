"""POC terisolasi untuk fitur baru: Payroll (BPJS + PPh 21 TER + rekonsiliasi
Desember), Impor Excel karyawan, Slip Gaji PDF, dan Email Reminder via SMTP.

Jalankan:  cd /app/backend && python test_core.py

Skrip ini TIDAK menyentuh FastAPI maupun MongoDB. Tujuannya membuktikan bahwa
inti perhitungan dan integrasi bekerja sebelum dibangun menjadi API + UI.
"""
from __future__ import annotations

import asyncio
import sys
import threading
import traceback
from datetime import datetime, timedelta, timezone

from app.core.excel import (
    build_employee_template,
    parse_employee_rows,
    validate_employee_rows,
)
from app.core.mailer import SmtpConfig, send_email
from app.core.payroll import (
    ComponentInput,
    EmployeeInput,
    StatutoryConfig,
    compute_december_recalc,
    compute_payslip,
    number_to_words_id,
    overtime_pay,
    summarize_run,
)
from app.core.pdf import build_payslip_pdf
from app.core.reminder_mail import build_reminder_digest, select_expiring
from app.core.ter import (
    PTKP_STATUSES,
    TER_TABLES,
    article_17_tax,
    ptkp_annual,
    ter_category,
    ter_rate,
    validate_tables,
)

PASS, FAIL = [], []


def check(name: str, condition: bool, detail: str = "") -> None:
    if condition:
        PASS.append(name)
        print(f"  [OK]   {name}" + (f" \u2014 {detail}" if detail else ""))
    else:
        FAIL.append(f"{name} :: {detail}")
        print(f"  [FAIL] {name} \u2014 {detail}")


def approx(a, b, tol=1.0) -> bool:
    return abs(float(a) - float(b)) <= tol


def rp(v) -> str:
    return f"Rp {float(v):,.0f}".replace(",", ".")


# =========================================================================
# (a) Integritas tabel TER
# =========================================================================
def test_ter_tables():
    print("\n--- (a) Integritas tabel TER PMK 168/2023 ---")
    try:
        validate_tables()
        check("Struktur tabel TER valid (44/40/41 lapisan, monoton, 0%..34%)", True)
    except Exception as exc:  # noqa: BLE001
        check("Struktur tabel TER valid", False, str(exc))

    check("Kategori A punya 44 lapisan", len(TER_TABLES["A"]) == 44, f"{len(TER_TABLES['A'])}")
    check("Kategori B punya 40 lapisan", len(TER_TABLES["B"]) == 40, f"{len(TER_TABLES['B'])}")
    check("Kategori C punya 41 lapisan", len(TER_TABLES["C"]) == 41, f"{len(TER_TABLES['C'])}")

    # pemetaan PTKP -> kategori
    for status, expected in [
        ("TK/0", "A"), ("TK/1", "A"), ("K/0", "A"),
        ("TK/2", "B"), ("TK/3", "B"), ("K/1", "B"), ("K/2", "B"),
        ("K/3", "C"),
    ]:
        check(f"PTKP {status} -> TER {expected}", ter_category(status) == expected,
              ter_category(status))

    # ambang bebas potong
    check("TER A bebas potong s.d. 5.400.000", ter_rate("TK/0", 5_400_000) == 0.0)
    check("TER A mulai kena di 5.400.001", ter_rate("TK/0", 5_400_001) == 0.0025)
    check("TER B bebas potong s.d. 6.200.000", ter_rate("K/1", 6_200_000) == 0.0)
    check("TER C bebas potong s.d. 6.600.000", ter_rate("K/3", 6_600_000) == 0.0)

    # spot-check terhadap dokumen resmi DJP
    official = [
        ("TK/0", 10_000_000, 0.0200, "A 9.650.001-10.050.000 = 2%"),
        ("TK/0", 12_000_000, 0.0400, "A 11.600.001-12.500.000 = 4%"),
        ("K/1", 15_000_000, 0.0600, "B 14.950.001-16.400.000 = 6%"),
        ("K/1", 8_000_000, 0.0100, "B 7.300.001-9.200.000 = 1%"),
        ("K/3", 20_000_000, 0.0800, "C 19.500.001-22.700.000 = 8%"),
        ("TK/0", 1_500_000_000, 0.3400, "A puncak >1,4M = 34%"),
        ("K/1", 1_500_000_000, 0.3400, "B puncak >1,405M = 34%"),
        ("K/3", 1_500_000_000, 0.3400, "C puncak >1,419M = 34%"),
    ]
    for status, gross, expected, label in official:
        got = ter_rate(status, gross)
        check(f"Tarif resmi: {label}", got == expected, f"dapat {got*100:.2f}%")

    # PTKP setahun
    for status, expected in [("TK/0", 54_000_000), ("TK/1", 58_500_000), ("K/0", 58_500_000),
                             ("K/1", 63_000_000), ("K/2", 67_500_000), ("K/3", 72_000_000)]:
        check(f"PTKP {status} = {rp(expected)}", ptkp_annual(status) == expected,
              rp(ptkp_annual(status)))

    # Pasal 17
    check("Pasal 17: 50jt -> 2,5jt", approx(article_17_tax(50_000_000), 2_500_000))
    check("Pasal 17: 60jt -> 3jt", approx(article_17_tax(60_000_000), 3_000_000))
    check("Pasal 17: 100jt -> 9jt", approx(article_17_tax(100_000_000), 9_000_000),
          rp(article_17_tax(100_000_000)))
    # 60jt x 5% = 3jt ; 190jt x 15% = 28,5jt ; 50jt x 25% = 12,5jt  -> 44jt
    check("Pasal 17: 300jt -> 44jt", approx(article_17_tax(300_000_000), 44_000_000),
          rp(article_17_tax(300_000_000)))
    check("Pasal 17: 250jt -> 31,5jt", approx(article_17_tax(250_000_000), 31_500_000),
          rp(article_17_tax(250_000_000)))
    check("Pasal 17: 0 -> 0", article_17_tax(0) == 0)

    # semua status dikenali
    check("Semua 8 status PTKP terpetakan",
          all(ter_category(s) in ("A", "B", "C") for s in PTKP_STATUSES))


# =========================================================================
# (b) Perhitungan slip gaji bulanan
# =========================================================================
def test_monthly_payslip():
    print("\n--- (b) Perhitungan slip gaji bulanan ---")
    cfg = StatutoryConfig()

    def reconcile(slip, label):
        t = slip["totals"]
        real_ded = sum(d["amount"] for d in slip["deductions"] if not d.get("informational"))
        check(f"{label}: bruto = jumlah penghasilan",
              approx(t["gross"], sum(e["amount"] for e in slip["earnings"])))
        check(f"{label}: total potongan konsisten", approx(t["total_deductions"], real_ded))
        check(f"{label}: netto = bruto - potongan",
              approx(t["net_pay"], t["gross"] - t["total_deductions"]),
              f"netto {rp(t['net_pay'])}")

    # --- 1. Di bawah ambang PPh 21 -> pajak nol
    low = compute_payslip(EmployeeInput(
        employee_id="e1", full_name="Karyawan Bawah Ambang",
        basic_salary=4_500_000, ptkp_status="TK/0"), cfg, month=3)
    check("Gaji 4,5jt TK/0: PPh 21 = 0", low["totals"]["pph21"] == 0,
          rp(low["totals"]["pph21"]))
    check("Gaji 4,5jt: BPJS pekerja terpotong",
          low["totals"]["bpjs_employee"] > 0, rp(low["totals"]["bpjs_employee"]))
    check("Gaji 4,5jt: BPJS Kesehatan 1% = 45.000",
          approx(low["bpjs"]["kesehatan"]["employee"], 45_000))
    check("Gaji 4,5jt: JHT 2% = 90.000", approx(low["bpjs"]["jht"]["employee"], 90_000))
    check("Gaji 4,5jt: JP 1% = 45.000", approx(low["bpjs"]["jp"]["employee"], 45_000))
    reconcile(low, "Bawah ambang")

    # --- 2. Menengah TK/0 dengan tunjangan
    mid = compute_payslip(EmployeeInput(
        employee_id="e2", full_name="Karyawan Menengah", basic_salary=9_000_000,
        ptkp_status="TK/0",
        components=[
            ComponentInput("TJ_JAB", "Tunjangan Jabatan", "earning", 1_500_000,
                           taxable=True, include_in_bpjs_base=True),
            ComponentInput("TJ_TRP", "Tunjangan Transport", "earning", 500_000,
                           taxable=True, prorate=True),
        ]), cfg, month=4)
    check("Menengah: bruto = 11.000.000", approx(mid["totals"]["gross"], 11_000_000),
          rp(mid["totals"]["gross"]))
    check("Menengah: kategori TER = A", mid["tax"]["ter_category"] == "A")
    check("Menengah: PPh 21 > 0", mid["totals"]["pph21"] > 0, rp(mid["totals"]["pph21"]))
    # bruto kena pajak = 11.000.000 + natura pemberi kerja (Kes 4% + JKK + JKM atas 10,5jt)
    expected_taxable = 11_000_000 + round(10_500_000 * (0.04 + 0.0089 + 0.003))
    check("Menengah: bruto kena pajak termasuk natura pemberi kerja",
          approx(mid["tax"]["taxable_gross"], expected_taxable, 5),
          f"{rp(mid['tax']['taxable_gross'])} vs {rp(expected_taxable)}")
    check("Menengah: PPh21 = tarif TER x bruto kena pajak",
          approx(mid["totals"]["pph21"],
                 mid["tax"]["taxable_gross"] * mid["tax"]["ter_rate"], 2))
    check("Menengah: dasar BPJS termasuk tunjangan jabatan (bukan transport)",
          approx(mid["bpjs"]["wage_base"], 10_500_000), rp(mid["bpjs"]["wage_base"]))
    reconcile(mid, "Menengah")

    # --- 3. K/3 -> kategori C
    k3 = compute_payslip(EmployeeInput(
        employee_id="e3", full_name="Karyawan K/3", basic_salary=20_000_000,
        ptkp_status="K/3"), cfg, month=5)
    check("K/3: kategori TER = C", k3["tax"]["ter_category"] == "C")
    check("K/3 gaji 20jt: dasar JP dibatasi 11.086.300",
          approx(k3["bpjs"]["jp"]["base"], 11_086_300), rp(k3["bpjs"]["jp"]["base"]))
    check("K/3 gaji 20jt: JHT tanpa batas (dasar 20jt)",
          approx(k3["bpjs"]["jht"]["base"], 20_000_000), rp(k3["bpjs"]["jht"]["base"]))
    reconcile(k3, "K/3")

    # --- 4. Gaji tinggi -> kedua batas BPJS aktif
    high = compute_payslip(EmployeeInput(
        employee_id="e4", full_name="Karyawan Tinggi", basic_salary=50_000_000,
        ptkp_status="TK/0"), cfg, month=6)
    check("Tinggi: dasar BPJS Kesehatan dibatasi 12.000.000",
          approx(high["bpjs"]["kesehatan"]["base"], 12_000_000),
          rp(high["bpjs"]["kesehatan"]["base"]))
    check("Tinggi: flag capped Kesehatan aktif", high["bpjs"]["kesehatan"]["capped"] is True)
    check("Tinggi: BPJS Kesehatan pekerja = 120.000",
          approx(high["bpjs"]["kesehatan"]["employee"], 120_000))
    check("Tinggi: dasar JP dibatasi 11.086.300",
          approx(high["bpjs"]["jp"]["base"], 11_086_300))
    check("Tinggi: JP pekerja = 110.863", approx(high["bpjs"]["jp"]["employee"], 110_863))
    check("Tinggi: JHT pekerja 2% x 50jt = 1.000.000",
          approx(high["bpjs"]["jht"]["employee"], 1_000_000))
    reconcile(high, "Tinggi")

    # --- 5. Lembur
    ot_hours = 10
    ot = compute_payslip(EmployeeInput(
        employee_id="e5", full_name="Karyawan Lembur", basic_salary=8_650_000,
        ptkp_status="TK/0", overtime_hours=ot_hours), cfg, month=7)
    hourly = 8_650_000 / 173  # = 50.000
    expected_ot = round(hourly * (1 * 1.5 + 9 * 2.0))
    check("Lembur: upah/jam = upah/173 = 50.000", approx(hourly, 50_000))
    check(f"Lembur {ot_hours} jam: 1,5x jam-1 + 2x sisanya = {rp(expected_ot)}",
          approx(overtime_pay(8_650_000, ot_hours, cfg), expected_ot))
    ot_row = next((e for e in ot["earnings"] if e["code"] == "OVERTIME"), None)
    check("Lembur muncul sebagai komponen penghasilan", ot_row is not None)
    check("Lembur: bruto = pokok + lembur",
          approx(ot["totals"]["gross"], 8_650_000 + expected_ot))
    check("Lembur: tidak menambah dasar BPJS",
          approx(ot["bpjs"]["wage_base"], 8_650_000), rp(ot["bpjs"]["wage_base"]))
    reconcile(ot, "Lembur")
    check("Lembur 0 jam -> 0", overtime_pay(8_650_000, 0, cfg) == 0)

    # --- 6. Ketidakhadiran tanpa upah (proration)
    absent = compute_payslip(EmployeeInput(
        employee_id="e6", full_name="Karyawan Absen", basic_salary=8_800_000,
        ptkp_status="TK/0", working_days=22, unpaid_days=2,
        components=[
            ComponentInput("TJ_MKN", "Tunjangan Makan", "earning", 660_000, prorate=True),
            ComponentInput("TJ_JAB", "Tunjangan Jabatan", "earning", 1_000_000, prorate=False),
        ]), cfg, month=8)
    check("Absen 2/22 hari: gaji pokok terpotong proporsional",
          approx(absent["basic_salary_paid"], 8_800_000 * 20 / 22),
          rp(absent["basic_salary_paid"]))
    meal = next(e for e in absent["earnings"] if e["code"] == "TJ_MKN")
    jab = next(e for e in absent["earnings"] if e["code"] == "TJ_JAB")
    check("Tunjangan prorate ikut terpotong", approx(meal["amount"], 660_000 * 20 / 22),
          rp(meal["amount"]))
    check("Tunjangan non-prorate tetap utuh", approx(jab["amount"], 1_000_000))
    check("Potongan tidak hadir tercatat informational",
          any(d["code"] == "ABSENCE" and d.get("informational") for d in absent["deductions"]))
    check("Absen: paid_days = 20", absent["attendance"]["paid_days"] == 20)
    reconcile(absent, "Absen")

    # --- 7. Tunjangan & potongan ad-hoc
    adhoc = compute_payslip(EmployeeInput(
        employee_id="e7", full_name="Karyawan Ad-hoc", basic_salary=10_000_000,
        ptkp_status="K/1",
        components=[
            ComponentInput("BONUS", "Bonus Proyek", "earning", 2_000_000, taxable=True),
            ComponentInput("NONTAX", "Reimburse Medis", "earning", 750_000, taxable=False),
            ComponentInput("KOP", "Potongan Koperasi", "deduction", 300_000),
            ComponentInput("PINJ", "Angsuran Pinjaman", "deduction", 1_200_000),
        ]), cfg, month=9)
    check("Ad-hoc: bruto = 12.750.000", approx(adhoc["totals"]["gross"], 12_750_000),
          rp(adhoc["totals"]["gross"]))
    check("Ad-hoc: komponen non-taxable tidak masuk bruto kena pajak",
          adhoc["tax"]["taxable_gross"] < adhoc["totals"]["gross"],
          rp(adhoc["tax"]["taxable_gross"]))
    check("Ad-hoc: potongan lain-lain = 1.500.000",
          approx(adhoc["totals"]["other_deductions"], 1_500_000),
          rp(adhoc["totals"]["other_deductions"]))
    check("Ad-hoc: kategori TER = B (K/1)", adhoc["tax"]["ter_category"] == "B")
    reconcile(adhoc, "Ad-hoc")

    # --- 8. Non-aktif BPJS
    no_bpjs = compute_payslip(EmployeeInput(
        employee_id="e8", full_name="Tanpa BPJS", basic_salary=10_000_000,
        ptkp_status="TK/0", bpjs_kesehatan_enrolled=False,
        bpjs_jht_enrolled=False, bpjs_jp_enrolled=False), cfg, month=10)
    check("Tanpa BPJS: potongan BPJS pekerja = 0",
          no_bpjs["totals"]["bpjs_employee"] == 0)
    check("Tanpa BPJS: netto = bruto - PPh21",
          approx(no_bpjs["totals"]["net_pay"],
                 no_bpjs["totals"]["gross"] - no_bpjs["totals"]["pph21"]))

    # --- 9. Surcharge tanpa NPWP (opsional)
    cfg_np = StatutoryConfig(non_npwp_surcharge=True)
    with_npwp = compute_payslip(EmployeeInput(
        employee_id="e9", full_name="Ada NPWP", basic_salary=12_000_000,
        ptkp_status="TK/0", has_npwp=True), cfg_np, month=4)
    without = compute_payslip(EmployeeInput(
        employee_id="e10", full_name="Tanpa NPWP", basic_salary=12_000_000,
        ptkp_status="TK/0", has_npwp=False), cfg_np, month=4)
    check("Tanpa NPWP: PPh21 20% lebih tinggi (bila kebijakan aktif)",
          approx(without["totals"]["pph21"], with_npwp["totals"]["pph21"] * 1.2, 2),
          f"{rp(without['totals']['pph21'])} vs {rp(with_npwp['totals']['pph21'])}")

    # --- 10. Ringkasan run
    summary = summarize_run([low, mid, k3, high, ot, absent, adhoc])
    check("Ringkasan run: jumlah karyawan = 7", summary["employee_count"] == 7)
    check("Ringkasan run: total netto = jumlah netto slip",
          approx(summary["net_pay"],
                 sum(s["totals"]["net_pay"] for s in [low, mid, k3, high, ot, absent, adhoc])),
          rp(summary["net_pay"]))
    check("Ringkasan run: biaya pemberi kerja > bruto",
          summary["employer_cost"] > summary["gross"])

    # --- 11. Angka ke kata
    check("Angka ke kata: 1.500.000",
          number_to_words_id(1_500_000) == "satu juta lima ratus ribu rupiah",
          number_to_words_id(1_500_000))
    check("Angka ke kata: 0", number_to_words_id(0) == "nol rupiah")

    return mid


# =========================================================================
# (c) Rekonsiliasi Desember (Pasal 17)
# =========================================================================
def test_december_recalc():
    print("\n--- (c) Rekonsiliasi PPh 21 Desember (Pasal 17) ---")
    cfg = StatutoryConfig()
    emp = EmployeeInput(employee_id="d1", full_name="Karyawan Setahun",
                        basic_salary=15_000_000, ptkp_status="TK/0")

    # Simulasi Januari-November memakai TER
    ytd_gross = ytd_tax = ytd_jhtjp = 0.0
    for m in range(1, 12):
        slip = compute_payslip(emp, cfg, month=m, year=2026)
        ytd_gross += slip["tax"]["taxable_gross"]
        ytd_tax += slip["totals"]["pph21"]
        ytd_jhtjp += slip["jht_jp_employee"]
    check("Jan-Nov terpotong TER (11 bulan)", ytd_tax > 0, f"total {rp(ytd_tax)}")

    dec = compute_payslip(emp, cfg, month=12, year=2026,
                          ytd_taxable_gross=ytd_gross, ytd_pph21_withheld=ytd_tax,
                          ytd_jht_jp_employee=ytd_jhtjp, ytd_months=11)
    ann = dec["tax"]["annual"]
    check("Desember memakai metode Pasal 17",
          dec["tax"]["method"] == "article_17_annual", dec["tax"]["method"])

    # verifikasi manual
    annual_gross = ytd_gross + dec["tax"]["taxable_gross"]
    biaya = min(annual_gross * 0.05, 500_000 * 12)
    jhtjp = ytd_jhtjp + dec["jht_jp_employee"]
    pkp_raw = max(0.0, annual_gross - biaya - jhtjp - 54_000_000)
    pkp = int(pkp_raw // 1000) * 1000
    expected_annual_tax = round(article_17_tax(pkp))

    check("Bruto setahun = akumulasi 12 bulan",
          approx(ann["annual_taxable_gross"], annual_gross, 2), rp(ann["annual_taxable_gross"]))
    check("Biaya jabatan 5% maks 6.000.000/tahun",
          approx(ann["biaya_jabatan"], biaya, 2), rp(ann["biaya_jabatan"]))
    check("Biaya jabatan tidak melebihi 6.000.000", ann["biaya_jabatan"] <= 6_000_000)
    check("PTKP TK/0 = 54.000.000", ann["ptkp_annual"] == 54_000_000)
    check("PKP dibulatkan ke bawah ke ribuan", ann["pkp"] % 1000 == 0, str(ann["pkp"]))
    check("PKP sesuai hitung manual", approx(ann["pkp"], pkp, 1000), rp(ann["pkp"]))
    check("PPh 21 setahun sesuai Pasal 17",
          approx(ann["annual_tax"], expected_annual_tax, 2), rp(ann["annual_tax"]))
    check("Sudah dipotong Jan-Nov tercatat",
          approx(ann["already_withheld"], ytd_tax, 2), rp(ann["already_withheld"]))
    check("PPh21 Desember = setahun - sudah dipotong",
          approx(dec["totals"]["pph21"], ann["annual_tax"] - ytd_tax, 2),
          rp(dec["totals"]["pph21"]))
    check("Total setahun = PPh terutang setahun",
          approx(ytd_tax + dec["totals"]["pph21"], ann["annual_tax"], 2))
    check("Desember: netto = bruto - potongan",
          approx(dec["totals"]["net_pay"],
                 dec["totals"]["gross"] - dec["totals"]["total_deductions"]))

    # Kasus kelebihan potong -> restitusi (negatif)
    refund = compute_december_recalc(
        monthly_taxable_gross=8_000_000, ytd_taxable_gross=88_000_000,
        ytd_pph21_withheld=9_000_000,  # sengaja dilebihkan
        monthly_jht_jp_employee=240_000, ytd_jht_jp_employee=2_640_000,
        ptkp_status="TK/0", months_worked=12, cfg=cfg)
    check("Kelebihan potong menghasilkan penyesuaian negatif",
          refund["pph21"] < 0, rp(refund["pph21"]))
    check("Flag is_refund aktif", refund["annual"]["is_refund"] is True)

    # Kasus PKP nol (gaji kecil setahun)
    small = compute_december_recalc(
        monthly_taxable_gross=4_000_000, ytd_taxable_gross=44_000_000,
        ytd_pph21_withheld=0, monthly_jht_jp_employee=120_000,
        ytd_jht_jp_employee=1_320_000, ptkp_status="TK/0", months_worked=12, cfg=cfg)
    check("Penghasilan di bawah PTKP -> PKP 0 dan pajak 0",
          small["annual"]["pkp"] == 0 and small["pph21"] == 0,
          f"pkp={small['annual']['pkp']} pajak={rp(small['pph21'])}")

    # K/3 punya PTKP lebih besar -> pajak setahun lebih kecil
    a = compute_december_recalc(10_000_000, 110_000_000, 0, 300_000, 3_300_000, "TK/0", 12, cfg)
    c = compute_december_recalc(10_000_000, 110_000_000, 0, 300_000, 3_300_000, "K/3", 12, cfg)
    check("K/3 bayar pajak setahun lebih kecil dari TK/0",
          c["annual"]["annual_tax"] < a["annual"]["annual_tax"],
          f"K/3 {rp(c['annual']['annual_tax'])} < TK/0 {rp(a['annual']['annual_tax'])}")

    return dec


# =========================================================================
# (d) Roundtrip Excel: template -> isi -> parse -> validasi
# =========================================================================
def test_excel_roundtrip():
    print("\n--- (d) Impor Excel karyawan (template -> parse -> validasi) ---")
    from io import BytesIO

    from openpyxl import load_workbook

    masters = {
        "departments": [
            {"id": "dep-1", "code": "HC", "name": "Human Capital"},
            {"id": "dep-2", "code": "OPS", "name": "Operations"},
        ],
        "branches": [{"id": "br-1", "code": "HO", "name": "Head Office Jakarta"}],
        "positions": [{"id": "pos-1", "code": "STF", "name": "Staff"}],
        "employment_statuses": [{"id": "es-1", "code": "PKWT", "name": "Kontrak (PKWT)"}],
        "work_locations": [], "divisions": [], "job_grades": [],
        "cost_centers": [], "projects": [],
    }

    template = build_employee_template(masters)
    check("Template Excel dihasilkan", len(template) > 3000, f"{len(template)} byte")
    check("Template adalah file xlsx (zip)", template[:2] == b"PK")

    wb = load_workbook(BytesIO(template))
    check("Template punya sheet 'Data Karyawan'", "Data Karyawan" in wb.sheetnames,
          str(wb.sheetnames))
    check("Template punya sheet 'Referensi'", "Referensi" in wb.sheetnames)
    ref = wb["Referensi"]
    ref_values = {str(c.value) for row in ref.iter_rows() for c in row if c.value}
    check("Sheet referensi memuat nilai master (Human Capital)",
          "Human Capital" in ref_values)
    check("Sheet referensi memuat status PTKP K/3", "K/3" in ref_values)

    # ---- isi data: 4 baris valid, 5 baris sengaja salah ----
    ws = wb["Data Karyawan"]
    headers = [c.value for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(headers)}

    def put(row, **kw):
        for label, value in kw.items():
            ws.cell(row=row, column=idx[label], value=value)

    put(3, **{"Nama Lengkap*": "Ahmad Fauzi", "Nomor KTP": "3201010101010001",
              "Status PTKP": "K/1", "Jenis Kelamin": "L", "Tanggal Lahir": "1990-05-17",
              "Email": "ahmad.fauzi@contoh.co.id", "Departemen": "Human Capital",
              "Cabang": "HO", "Jabatan (master)": "Staff", "Gaji Pokok": 8500000,
              "Tanggal Masuk": "2024-01-15", "Status Kepegawaian": "Kontrak (PKWT)"})
    put(4, **{"Nama Lengkap*": "Siti Aminah", "Nomor KTP": "3201010101010002",
              "Status PTKP": "TK/0", "Jenis Kelamin": "P", "Departemen": "Operations",
              "Gaji Pokok": "9.500.000", "Tanggal Masuk": "15/02/2024"})
    put(5, **{"Nama Lengkap*": "Budi Santoso", "Status PTKP": "k/3",
              "Jenis Kelamin": "l", "Status Pernikahan": "kawin"})
    put(6, **{"Nama Lengkap*": "Dewi Lestari", "Nomor KTP": "3201010101010003",
              "Agama": "islam", "Kota": "Bandung"})
    # baris salah
    put(7, **{"Nomor KTP": "3201010101010009", "Departemen": "Human Capital"})       # nama kosong
    put(8, **{"Nama Lengkap*": "Error Tanggal", "Tanggal Lahir": "17 Mei 1990"})      # tanggal
    put(9, **{"Nama Lengkap*": "Error Departemen", "Departemen": "Tidak Ada"})        # master
    put(10, **{"Nama Lengkap*": "Error KTP Duplikat", "Nomor KTP": "3201010101010001"})  # dup in-file
    put(11, **{"Nama Lengkap*": "Error Email", "Email": "bukan-email",
               "Status PTKP": "X/9", "Gaji Pokok": "abc"})                            # multi error

    filled = BytesIO()
    wb.save(filled)
    filled_bytes = filled.getvalue()

    rows = parse_employee_rows(filled_bytes)
    check("Parse menemukan 9 baris data (baris petunjuk diabaikan)",
          len(rows) == 9, f"{len(rows)} baris")

    report = validate_employee_rows(
        rows, masters,
        existing_employee_numbers={"NEP-0001"},
        existing_niks={"3201019999999999"},
    )
    check("Laporan: 4 baris valid", report["valid_count"] == 4, str(report["valid_count"]))
    check("Laporan: 5 baris invalid", report["invalid_count"] == 5, str(report["invalid_count"]))
    check("Laporan: can_commit true", report["can_commit"] is True)

    by_row = {r["excel_row"]: r for r in report["rows"]}
    check("Baris 7 ditolak: nama wajib diisi",
          not by_row[7]["is_valid"] and any("wajib diisi" in e for e in by_row[7]["errors"]),
          str(by_row[7]["errors"]))
    check("Baris 8 ditolak: format tanggal",
          any("tanggal" in e.lower() for e in by_row[8]["errors"]), str(by_row[8]["errors"]))
    check("Baris 9 ditolak: departemen tidak ada di master",
          any("tidak ditemukan di master" in e for e in by_row[9]["errors"]),
          str(by_row[9]["errors"]))
    check("Baris 10 ditolak: KTP duplikat dalam file",
          any("duplikat" in e for e in by_row[10]["errors"]), str(by_row[10]["errors"]))
    check("Baris 11 ditolak dengan beberapa error sekaligus",
          len(by_row[11]["errors"]) >= 3, str(by_row[11]["errors"]))
    check("Pesan error berbahasa Indonesia",
          all(any(w in e.lower() for w in
                  ["wajib", "tidak", "duplikat", "format", "valid", "harus", "sudah", "bukan"])
              for r in report["rows"] for e in r["errors"]))

    # normalisasi baris valid
    r3, r4, r5 = by_row[3], by_row[4], by_row[5]
    check("Baris 3: departemen dipetakan ke id master",
          r3["payload"]["department_id"] == "dep-1", str(r3["payload"].get("department_id")))
    check("Baris 3: cabang dikenali lewat KODE (HO)",
          r3["payload"]["branch_id"] == "br-1")
    check("Baris 3: gaji pokok numerik", r3["payload"]["basic_salary"] == 8_500_000)
    check("Baris 4: gaji '9.500.000' dinormalisasi",
          r4["payload"]["basic_salary"] == 9_500_000, str(r4["payload"].get("basic_salary")))
    check("Baris 4: tanggal 15/02/2024 -> 2024-02-15",
          r4["payload"]["join_date"] == "2024-02-15", str(r4["payload"].get("join_date")))
    check("Baris 5: PTKP 'k/3' dinormalisasi ke 'K/3'",
          r5["payload"]["ptkp_status"] == "K/3", str(r5["payload"].get("ptkp_status")))
    check("Baris 5: gender 'l' dinormalisasi ke 'L'", r5["payload"]["gender"] == "L")

    # NIK yang sudah ada di DB harus ditolak
    dup = validate_employee_rows(rows, masters, existing_niks={"3201010101010001"})
    check("KTP yang sudah ada di database ditolak",
          any("sudah terdaftar" in e for r in dup["rows"] for e in r["errors"]))

    # file rusak
    try:
        parse_employee_rows(b"ini bukan excel")
        check("File non-Excel ditolak dengan pesan jelas", False, "tidak ada error")
    except ValueError as exc:
        check("File non-Excel ditolak dengan pesan jelas", "Excel" in str(exc), str(exc)[:70])

    return report


# =========================================================================
# (e) Slip gaji PDF
# =========================================================================
def test_payslip_pdf(slip_monthly, slip_december):
    print("\n--- (e) Slip gaji PDF ---")
    company = {"name": "PT Nusantara Energi Prima", "address": "Menara Prima Lt. 18",
               "city": "Jakarta Selatan"}

    pdf = build_payslip_pdf(slip_monthly, company,
                            generated_at=datetime.now().strftime("%d/%m/%Y %H:%M"))
    check("PDF bulanan dihasilkan", isinstance(pdf, bytes) and len(pdf) > 3000,
          f"{len(pdf)} byte")
    check("PDF punya header %PDF", pdf[:4] == b"%PDF", str(pdf[:8]))
    check("PDF diakhiri EOF", b"%%EOF" in pdf[-1024:])

    pdf_dec = build_payslip_pdf(slip_december, company, generated_at="31/12/2026 17:00")
    check("PDF Desember (Pasal 17) dihasilkan", pdf_dec[:4] == b"%PDF" and len(pdf_dec) > 3000,
          f"{len(pdf_dec)} byte")
    check("PDF Desember berbeda dari PDF bulanan", pdf_dec != pdf)

    # slip minimal / gaji nol tidak boleh error
    minimal = compute_payslip(EmployeeInput(employee_id="z", full_name="Tanpa Gaji",
                                            basic_salary=0, ptkp_status="TK/0"))
    pdf_min = build_payslip_pdf(minimal, {})
    check("PDF tetap dihasilkan untuk slip tanpa penghasilan", pdf_min[:4] == b"%PDF")

    with open("/tmp/poc_slip_gaji.pdf", "wb") as fh:
        fh.write(pdf)
    print("         contoh PDF: /tmp/poc_slip_gaji.pdf")


# =========================================================================
# (f) + (g) Email reminder via SMTP + seleksi expiry
# =========================================================================
def test_expiry_selection():
    print("\n--- (g) Seleksi item mendekati masa berlaku ---")
    today = datetime(2026, 6, 1, tzinfo=timezone.utc)

    def d(days):
        return (today + timedelta(days=days)).date().isoformat()

    records = [
        {"id": "1", "title": "Kontrak A", "expiry_date": d(-10)},   # kedaluwarsa
        {"id": "2", "title": "Kontrak B", "expiry_date": d(5)},     # window 7
        {"id": "3", "title": "Kontrak C", "expiry_date": d(25)},    # window 30
        {"id": "4", "title": "Kontrak D", "expiry_date": d(45)},    # window 60
        {"id": "5", "title": "Kontrak E", "expiry_date": d(85)},    # window 90
        {"id": "6", "title": "Kontrak F", "expiry_date": d(200)},   # di luar horizon
        {"id": "7", "title": "Kontrak G", "expiry_date": None},     # tanpa masa berlaku
    ]

    sel = select_expiring(records, [90, 60, 30, 7], today=today)
    ids = [s["id"] for s in sel]
    check("Item di luar horizon 90 hari tidak dipilih", "6" not in ids, str(ids))
    check("Item tanpa masa berlaku tidak dipilih", "7" not in ids)
    check("Item kedaluwarsa tetap dipilih", "1" in ids)
    check("Total terpilih = 5", len(sel) == 5, str(len(sel)))
    check("Hasil urut dari paling mendesak", ids[0] == "1" and ids[-1] == "5", str(ids))
    check("Window terkecil dipilih untuk 5 hari",
          next(s["window"] for s in sel if s["id"] == "2") == 7)
    check("Window 30 untuk 25 hari",
          next(s["window"] for s in sel if s["id"] == "3") == 30)
    check("Window 60 untuk 45 hari",
          next(s["window"] for s in sel if s["id"] == "4") == 60)
    check("Urgency 'expired' untuk item lewat",
          next(s["urgency"] for s in sel if s["id"] == "1") == "expired")

    narrow = select_expiring(records, [7], today=today)
    check("Window sempit (7 hari) hanya ambil 2 item", len(narrow) == 2, str(len(narrow)))
    no_exp = select_expiring(records, [90], today=today, include_expired=False)
    check("include_expired=False mengecualikan yang kedaluwarsa",
          all(s["days_left"] >= 0 for s in no_exp))
    check("Daftar kosong aman", select_expiring([], [30]) == [])
    return today


def test_smtp_send(today):
    print("\n--- (f) Pengiriman email reminder via SMTP ---")
    try:
        from aiosmtpd.controller import Controller
    except ImportError as exc:
        check("aiosmtpd tersedia untuk POC", False, str(exc))
        return

    received = []

    class Handler:
        async def handle_DATA(self, server, session, envelope):  # noqa: N802
            received.append({
                "mail_from": envelope.mail_from,
                "rcpt_tos": list(envelope.rcpt_tos),
                "content": envelope.content.decode("utf8", errors="replace"),
            })
            return "250 Message accepted for delivery"

    # aiosmtpd tidak mendukung port=0, jadi pilih port bebas terlebih dahulu
    import socket

    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        port = probe.getsockname()[1]

    controller = Controller(Handler(), hostname="127.0.0.1", port=port)
    controller.start()
    print(f"         SMTP uji berjalan di 127.0.0.1:{port}")

    try:
        def dd(days):
            return (today + timedelta(days=days)).date().isoformat()

        digest = build_reminder_digest(
            company_name="PT Nusantara Energi Prima",
            contracts=[
                {"id": "c1", "title": "PKWT 12 Bulan \u2014 Rina Kusuma",
                 "subtitle": "NEP-0001 \u00b7 Staff Administrasi HR", "expiry_date": dd(-3)},
                {"id": "c2", "title": "PKWT 6 Bulan \u2014 Joko Purnomo",
                 "subtitle": "NEP-0004 \u00b7 Teknisi Mekanik", "expiry_date": dd(21)},
            ],
            certifications=[
                {"id": "s1", "title": "Sertifikat K3 Umum \u2014 Bambang Setiawan",
                 "subtitle": "NEP-0002 \u00b7 Site Supervisor", "expiry_date": dd(44)},
            ],
            documents=[
                {"id": "d1", "title": "Medical Check Up \u2014 Joko Purnomo",
                 "subtitle": "Hasil Medical Check Up", "expiry_date": dd(60)},
            ],
            windows=[90, 60, 30, 7],
            today=today,
            app_url="https://contoh.preview.emergentagent.com/expiry-calendar",
        )
        check("Digest menghitung 4 item", digest["total"] == 4, str(digest["total"]))
        check("Digest memisahkan per jenis",
              digest["counts"]["contracts"] == 2 and digest["counts"]["certifications"] == 1
              and digest["counts"]["documents"] == 1, str(digest["counts"]))
        check("Digest menandai 1 item kedaluwarsa", digest["counts"]["expired"] == 1)
        check("Subjek memuat nama perusahaan dan jumlah",
              "PT Nusantara Energi Prima" in digest["subject"] and "4 item" in digest["subject"],
              digest["subject"])
        check("Subjek menyebut item kedaluwarsa", "kedaluwarsa" in digest["subject"])
        check("Body HTML tidak memakai latar transparan",
              "background:#FFFFFF" in digest["html"])

        cfg = {
            "host": "127.0.0.1", "port": port, "security": "none",
            "username": "", "password": "",
            "from_email": "hris@nusantaraenergi.co.id", "from_name": "HRIS Nusantara",
        }
        result = send_email(
            cfg,
            to=["hrd@nusantaraenergi.co.id", "direktur@nusantaraenergi.co.id"],
            subject=digest["subject"],
            html_body=digest["html"],
        )
        check("Email terkirim tanpa error", result["sent"] is True)
        check("Server SMTP menerima 1 pesan", len(received) == 1, str(len(received)))

        if received:
            msg = received[0]
            body = msg["content"]
            check("MAIL FROM sesuai konfigurasi",
                  msg["mail_from"] == "hris@nusantaraenergi.co.id", msg["mail_from"])
            check("Kedua penerima terdaftar di RCPT TO",
                  set(msg["rcpt_tos"]) == {"hrd@nusantaraenergi.co.id",
                                           "direktur@nusantaraenergi.co.id"},
                  str(msg["rcpt_tos"]))
            check("Header From memuat nama pengirim", "HRIS Nusantara" in body)
            check("Subjek terkirim utuh (termasuk yang panjang)",
                  "perlu perhatian" in body.replace("=\r\n", "").replace("=\n", ""))
            check("Body memuat baris kontrak Rina Kusuma", "Rina Kusuma" in body)
            check("Body memuat sertifikasi K3", "K3 Umum" in body)
            check("Body memuat dokumen Medical Check Up", "Medical Check Up" in body)
            check("Email multipart (teks + HTML)", "multipart/alternative" in body)
            check("Ada versi teks biasa untuk klien non-HTML",
                  "text/plain" in body and "text/html" in body)
            check("Password tidak pernah muncul di pesan", "password" not in body.lower())

        # validasi konfigurasi
        from app.core.mailer import MailerError
        for bad, label in [
            ({"host": "", "port": 587, "from_email": "a@b.co"}, "host kosong"),
            ({"host": "x", "port": 0, "from_email": "a@b.co"}, "port tidak valid"),
            ({"host": "x", "port": 25, "security": "aneh", "from_email": "a@b.co"}, "mode keamanan"),
            ({"host": "x", "port": 25, "from_email": ""}, "pengirim kosong"),
        ]:
            try:
                SmtpConfig(bad).validate()
                check(f"Konfigurasi ditolak: {label}", False, "lolos padahal invalid")
            except MailerError as exc:
                check(f"Konfigurasi ditolak: {label}", True, str(exc)[:55])

        try:
            send_email(cfg, to=[], subject="x", html_body="y")
            check("Penerima kosong ditolak", False, "lolos")
        except MailerError as exc:
            check("Penerima kosong ditolak", "penerima" in str(exc).lower(), str(exc)[:55])

        try:
            send_email({**cfg, "port": 59999}, to=["a@b.co"], subject="x", html_body="y")
            check("Host/port salah menghasilkan error jelas", False, "lolos")
        except MailerError as exc:
            check("Host/port salah menghasilkan error jelas",
                  "SMTP" in str(exc), str(exc)[:60])

        check("SmtpConfig.masked() menyembunyikan password",
              "password" not in SmtpConfig({**cfg, "password": "rahasia"}).masked()
              and SmtpConfig({**cfg, "password": "rahasia"}).masked()["has_password"] is True)
    finally:
        controller.stop()


# =========================================================================
def main():
    print("=" * 74)
    print(" POC INTI \u2014 Payroll (TER/BPJS), Impor Excel, Slip PDF, Email Reminder")
    print("=" * 74)

    try:
        test_ter_tables()
        slip_monthly = test_monthly_payslip()
        slip_december = test_december_recalc()
        test_excel_roundtrip()
        test_payslip_pdf(slip_monthly, slip_december)
        today = test_expiry_selection()
        test_smtp_send(today)
    except Exception:  # noqa: BLE001
        traceback.print_exc()
        FAIL.append(f"Exception tak tertangani: {sys.exc_info()[1]}")

    print("\n" + "=" * 74)
    total = len(PASS) + len(FAIL)
    print(f" HASIL: {len(PASS)}/{total} pemeriksaan lulus")
    if FAIL:
        print(f"\n GAGAL ({len(FAIL)}):")
        for f in FAIL:
            print(f"   - {f}")
        print("=" * 74)
        return 1
    print(" SEMUA PEMERIKSAAN INTI LULUS \u2014 siap lanjut ke pembangunan API + UI.")
    print("=" * 74)
    return 0


if __name__ == "__main__":
    sys.exit(main())
