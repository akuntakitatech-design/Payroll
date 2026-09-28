"""Phase 2A CP3 - Impor Master Aset + Saldo Awal / Opening Existing Holding.

CARA MENJALANKAN (hanya MariaDB development LOKAL `hris_dev`):
    cd backend && bash tests/dev_env_run.sh python3 tests/test_asset_cp3.py

Memakai ulang guard + fixture + cleanup bertanda run dari test_asset_cp2 (company/user/role/karyawan/aset milik run
ini saja; cleanup memverifikasi residue = 0 dan orphan = 0). Role uji = CUSTOM role dengan izin eksplisit (tanpa nama
role sebagai logika). Termasuk uji performa 5.000 baris dan konkurensi nomor BAST-EXS.
"""
from __future__ import annotations

import asyncio
import base64
import io
import pathlib
import re
import sys
import time
import zlib

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))

import test_asset_cp2 as T  # noqa: E402  (guard hris_dev dijalankan saat import)
import test_asset_cp2_1 as T21  # noqa: E402

import httpx  # noqa: E402
import sqlalchemy as sa  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

import server  # noqa: E402
from app.core import db as dbm  # noqa: E402
from app.core.db import get_db, new_id, now  # noqa: E402
from app.routers.asset_import import MASTER_COLS, OPENING_COLS  # noqa: E402

check, RUN, MARK, YEAR, TODAY = T.check, T.RUN, T.MARK, T.YEAR, T.TODAY
EXS_RE = re.compile(rf"^BAST-EXS/{YEAR}/(\d{{6}})$")
AST_RE = re.compile(r"^AST-(\d{6})$")
BASE = ["asset:view", "asset_master:view", "asset_bast:view", "asset_import:view", "asset_opening:view"]
DRAFTER = BASE + ["asset_import:create", "asset_opening:create", "asset_opening:edit"]
FINALIZER = DRAFTER + ["asset_import:commit", "asset_opening:publish", "asset_return:view", "asset_return:create",
                       "asset_return:edit", "asset_return:publish", "asset_value:edit"]
CP3_ALL = {f"asset_import:{a}" for a in ("view", "create", "commit")} | {f"asset_opening:{a}" for a in ("view", "create", "edit", "publish")}


# ------------------------------------------------------------------ helpers
def xlsx(cols, rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append([label for _, label in cols])
    for r in rows:
        ws.append([r.get(k) for k, _ in cols])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def preview(env, h, itype, rows, name="uji.xlsx"):
    cols = MASTER_COLS if itype == "MASTER" else OPENING_COLS
    return await env.c.post("/api/asset-imports/preview", params={"import_type": itype}, headers=h,
                            files={"file": (name, xlsx(cols, rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


async def rows_of(env, h, bid, row_class=None):
    r = await env.c.get(f"/api/asset-imports/{bid}/rows", params={"limit": 200, **({"row_class": row_class} if row_class else {})}, headers=h)
    return r.json()["items"]


def msgs(rows):
    return " | ".join(m["message"] for r in rows for m in r["messages"])


async def count(env, coll, **flt):
    return await get_db()[coll].count_documents({"company_id": env.cid, **flt})


async def mk_asset(env, h, name, **kw):
    body = {"name": f"{name} {MARK}", "category_id": env.alat, "condition_id": env.cond["BAIK"], "project_id": kw.pop("project_id", env.p1), **kw}
    r = await env.c.post("/api/assets", json=body, headers=h)
    assert r.status_code == 201, r.text
    return r.json()


def pdf_text(data: bytes) -> str:
    out = []
    for m in re.finditer(rb"stream\r?\n(.*?)endstream", data, re.S):
        s = m.group(1).strip()
        try:
            s = zlib.decompress(base64.a85decode(s[:-2] if s.endswith(b"~>") else s))
        except Exception:  # noqa: BLE001
            try:
                s = zlib.decompress(s)
            except Exception:  # noqa: BLE001
                continue
        out += [x.decode("latin1") for x in re.findall(rb"\((.*?)\)\s*Tj", s)]
    return " ".join(out)


async def counter(env, key, period):
    row = await get_db().document_sequence_counters.find_one({"company_id": env.cid, "sequence_key": key, "period_key": period}, {"_id": 0})
    return int((row or {}).get("next_value") or 1)


async def setup3(env):
    await T.setup(env)
    env.dr_id, env.dr = await T._mk_user(env, await T21._perm_role(env, "kappa", DRAFTER), env.cid, "drafter3")
    env.fz_id, env.fz = await T._mk_user(env, await T21._perm_role(env, "lambda", FINALIZER), env.cid, "finalizer3")
    o = (await env.c.get("/api/assets/form-options", headers=env.ga)).json()
    env.o = o
    env.st_code = {x["system_state"]: x["code"] for x in o["statuses"] if x.get("is_default")}
    env.proj_code = {p["id"]: p["code"] for p in o["projects"]}


def mrow(i, **kw):
    return {"name": f"Impor {i} {MARK}", "category": "ALAT", "unit": "UNIT", "condition": "BAIK", "status": "READY_X", **kw}


# ------------------------------------------------------------------ tests
async def t_rbac(env):
    db = get_db()
    perms = {p["key"] for p in await db.permissions.find({}, {"_id": 0}).to_list(3000) if p["key"].split(":")[0] in ("asset_import", "asset_opening")}
    check("permission CP3 terdaftar lengkap (7 key)", perms == CP3_ALL, str(sorted(perms)))
    for rk, exp in (("ga_admin", CP3_ALL), ("ga_staff", CP3_ALL - {"asset_import:commit", "asset_opening:publish"})):
        have = {p["permission_key"] for p in await db.role_permissions.find({"role_key": rk}, {"_id": 0}).to_list(3000)}
        check(f"preset {rk}: izin CP3 sesuai matriks", {k for k in have if k in CP3_ALL} == exp)
    r = await env.c.get("/api/asset-openings", headers=env.emp)
    check("role tanpa asset_opening:view -> 403", r.status_code == 403, str(r.status_code))
    r = await env.c.get("/api/asset-imports/template", params={"import_type": "MASTER"}, headers=env.emp)
    check("role tanpa asset_import:view -> 403 (template)", r.status_code == 403, str(r.status_code))


async def t_master_import(env):
    h = env.fz
    r = await env.c.get("/api/asset-imports/template", params={"import_type": "MASTER"}, headers=h)
    wb = load_workbook(io.BytesIO(r.content))
    head = [c.value for c in wb.worksheets[0][1]]
    check("template master .xlsx: 200 + kolom lengkap + sheet Referensi", r.status_code == 200 and head == [lbl for _, lbl in MASTER_COLS]
          and "Referensi" in wb.sheetnames, str(head)[:120])
    ready = env.st_code["READY"]
    existing = await mk_asset(env, env.ga, "Eksisting", serial_number=f"SN-EX-{RUN}")
    await get_db().assets.update_one({"company_id": env.cid, "id": existing["id"]}, {"$set": {"legacy_code": f"INV-EX-{RUN}", "legacy_code_norm": f"INV-EX-{RUN}"}})
    await get_db().asset_units.insert_one({"id": new_id(), "company_id": env.cid, "status": "active", "code": f"AMB1{RUN}", "name": f"Ambigu {RUN}", "created_at": now(), "updated_at": now()})
    await get_db().asset_units.insert_one({"id": new_id(), "company_id": env.cid, "status": "active", "code": f"AMB2{RUN}", "name": f"Ambigu {RUN}", "created_at": now(), "updated_at": now()})
    bad = [mrow(1, status=ready, legacy_code=f"INV-A-{RUN}"),
           mrow(2, status=ready, name=None),
           mrow(3, status=ready, category="TIDAKADA"),
           mrow(4, status=ready, unit="XX"),
           mrow(5, status=ready, condition="ZZ"),
           mrow(6, status=env.st_code["IN_USE"]),
           mrow(7, status=ready, category="LAPTOP"),
           mrow(8, status=ready, serial_number=f"SN-DUP-{RUN}"), mrow(9, status=ready, serial_number=f"SN-DUP-{RUN}"),
           mrow(10, status=ready, serial_number=f"SN-EX-{RUN}"),
           mrow(11, status=ready, legacy_code=f"INV-EX-{RUN}"),
           mrow(12, status=ready, acquisition_value="lima juta"),
           mrow(13, status=ready, acquisition_year="20X4"),
           mrow(14, status=ready, project="PROYEK-FIKTIF"),
           mrow(15, status=ready, unit=f"Ambigu {RUN}"),
           mrow(16, status=ready, legacy_code=f"INV-D-{RUN}"), mrow(17, status=ready, legacy_code=f"INV-D-{RUN}"),
           mrow(18, status="STATUS-X")]
    before = await count(env, "assets")
    ev_before = await count(env, "asset_events")
    ctr_before = await counter(env, "ASSET_CODE", "ALL")
    r = await preview(env, h, "MASTER", bad)
    b = r.json()
    check("preview master: 201 + ringkasan per kelas", r.status_code == 201 and b["total_rows"] == 18 and b["error_rows"] >= 16, str(b)[:200])
    check("preview TIDAK menulis data bisnis (aset/event/counter kode tetap)",
          await count(env, "assets") == before and await count(env, "asset_events") == ev_before
          and await counter(env, "ASSET_CODE", "ALL") == ctr_before)
    err = await rows_of(env, h, b["id"], "ERROR")
    text = msgs(err)
    for label, needle in (("nama kosong", "Nama Aset wajib diisi"), ("kategori tidak ditemukan", "Kategori 'TIDAKADA' tidak ditemukan"),
                          ("satuan tidak ditemukan", "Satuan 'XX' tidak ditemukan"), ("kondisi tidak ditemukan", "Kondisi 'ZZ' tidak ditemukan"),
                          ("status IN_USE ditolak", "tidak boleh diimpor lewat master"), ("serial wajib kategori", "Serial Number wajib diisi untuk kategori Laptop"),
                          ("duplicate serial dalam file", "duplikat di file"), ("duplicate serial vs DB", f"Serial Number sudah digunakan oleh {existing['asset_code']}"),
                          ("kode lama duplicate vs DB", f"Kode Aset Lama sudah digunakan oleh {existing['asset_code']}"),
                          ("nilai bukan angka", "Nilai Perolehan harus berupa angka"), ("tahun tidak valid", "Tahun Perolehan '20X4' tidak valid"),
                          ("project tidak ditemukan", "Proyek 'PROYEK-FIKTIF' tidak ditemukan"), ("referensi ambigu", "ambigu"),
                          ("status tidak valid", "Status 'STATUS-X' tidak ditemukan"), ("kode lama duplikat di file", f"Kode Aset Lama 'INV-D-{RUN}' duplikat di file")):
        check(f"validasi ERROR: {label}", needle in text, text[:160] if needle not in text else "")
    r = await env.c.get(f"/api/asset-imports/{b['id']}/rows", params={"row_class": "VALID"}, headers=h)
    check("filter baris VALID berfungsi", r.status_code == 200 and all(x["row_class"] == "VALID" for x in r.json()["items"]))
    r = await env.c.post(f"/api/asset-imports/{b['id']}/commit", json={}, headers=h)
    check("commit ditolak bila ada ERROR (409, tidak ada baris dilewati)", r.status_code == 409 and await count(env, "assets") == before, str(r.status_code))
    # impor bersih
    good = [mrow(i, status=ready, legacy_code=f"INV-G{i}-{RUN}", serial_number=f"SN-G{i}-{RUN}", project=env.proj_code[env.p1],
                 acquisition_year=2021, acquisition_value="12500000") for i in range(1, 4)]
    good.append(mrow(4, status=env.st_code["MAINTENANCE"], project=env.proj_code[env.p1]))   # tanpa SN & kode lama -> WARNING
    r = await preview(env, env.dr, "MASTER", good)
    text = msgs(await rows_of(env, env.dr, r.json()["id"], "ERROR"))
    check("tanpa asset_value:edit: Nilai Perolehan di file -> ERROR jelas", "tidak memiliki izin mengisi Nilai Perolehan" in text, text[:120])
    r = await preview(env, h, "MASTER", good)
    b = r.json()
    check("preview bersih: 0 ERROR, 1 WARNING", b["error_rows"] == 0 and b["warning_rows"] == 1, str(b)[:160])
    r = await env.c.post(f"/api/asset-imports/{b['id']}/commit", json={}, headers=env.dr)
    check("tanpa asset_import:commit -> 403", r.status_code == 403, str(r.status_code))
    ctr = await counter(env, "ASSET_CODE", "ALL")
    t0 = time.perf_counter()
    r = await env.c.post(f"/api/asset-imports/{b['id']}/commit", json={}, headers=h)
    c = r.json()
    codes = c.get("new_asset_codes", [])
    nums = [int(AST_RE.match(x).group(1)) for x in codes if AST_RE.match(x)]
    check("commit master: 4 aset + nomor batch", r.status_code == 200 and c["created_count"] == 4 and c["batch_number"].startswith(f"IMP-AST/{YEAR}/"), str(c)[:200])
    check("AST sequence benar: AST-{SEQ:6} berurutan, counter maju tepat 4", len(nums) == 4 and nums == list(range(nums[0], nums[0] + 4))
          and nums[0] == ctr and await counter(env, "ASSET_CODE", "ALL") == ctr + 4, str(codes))
    rows = await get_db().assets.find({"company_id": env.cid, "import_batch_id": b["id"]}, {"_id": 0}).to_list(10)
    check("aset impor: legacy_code terpisah dari asset_code, source IMPORT, status sesuai",
          len(rows) == 4 and all(x["source"] == "IMPORT" and x["asset_code"] != x.get("legacy_code") for x in rows)
          and {x["lifecycle_state"] for x in rows} == {"READY", "MAINTENANCE"} and not any(x["lifecycle_state"] == "IN_USE" for x in rows))
    check("event CREATED per aset impor", await count(env, "asset_events", event_type="CREATED", asset_id={"$in": [x["id"] for x in rows]}) == 4)
    r = await env.c.get(f"/api/asset-imports/{b['id']}/result", headers=h)
    ws = load_workbook(io.BytesIO(r.content)).worksheets[0]
    data = [[c.value for c in row] for row in ws.iter_rows()]
    check("hasil impor Excel: mapping Kode Lama -> Asset Code -> Nama -> SN -> Status",
          r.status_code == 200 and data[0] == ["Kode Aset Lama", "Asset Code Sistem", "Nama Aset", "Serial Number", "Status"]
          and data[1][0] == f"INV-G1-{RUN}" and data[1][1] == codes[0] and len(data) == 5, str(data[:2]))
    r = await env.c.post(f"/api/asset-imports/{b['id']}/commit", json={}, headers=h)
    check("commit ulang batch -> 409", r.status_code == 409, str(r.status_code))
    r = await env.c.get("/api/assets", params={"q": f"INV-G2-{RUN}"}, headers=env.ga)
    check("Daftar Aset: pencarian by Kode Aset Lama", r.json()["total"] == 1)
    env.imported = {x["legacy_code"]: x for x in rows if x.get("legacy_code")}
    env.t_commit_small = time.perf_counter() - t0


async def t_opening_manual(env):
    a1, a2, a3 = [await mk_asset(env, env.ga, f"Op{i}") for i in (1, 2, 3)]
    body = {"employee_id": env.e1, "opening_date": TODAY, "items": [{"asset_id": a1["id"], "condition_id": env.cond["BAIK"], "accessories": "Charger"}]}
    r = await env.c.post("/api/asset-openings", json=body, headers=env.dr)
    d = r.json()
    check("opening manual: drafter Simpan Draft -> 201 DRAFT, project default assignment", r.status_code == 201 and d["doc_state"] == "DRAFT"
          and d["project_id"] == env.p1, str(r.status_code))
    check("draft tidak mengubah aset / tidak membuat holding", (await get_db().assets.find_one({"id": a1["id"]}, {"_id": 0}))["lifecycle_state"] == "READY"
          and await count(env, "asset_holdings", asset_id=a1["id"]) == 0)
    r = await env.c.post(f"/api/asset-openings/{d['id']}/publish", headers=env.dr)
    check("drafter tanpa publish: Publish -> 403", r.status_code == 403, str(r.status_code))
    r = await env.c.post("/api/asset-openings/save-and-publish", json={**body, "items": [{"asset_id": a2["id"], "condition_id": env.cond["BAIK"]}]}, headers=env.dr)
    check("drafter tanpa publish: Simpan & Publish -> 403 (tanpa data)", r.status_code == 403 and await count(env, "asset_opening_items", asset_id=a2["id"]) == 0)
    ctr = await counter(env, "BAST_EXISTING", str(YEAR))
    r = await env.c.post(f"/api/asset-openings/{d['id']}/publish", headers=env.fz)
    p = r.json()
    num = p.get("bast_number") or ""
    m = EXS_RE.match(num)
    check("custom role dengan publish: Publish draft -> BAST-EXS/{YYYY}/{SEQ:6}", r.status_code == 200 and p["doc_state"] == "PUBLISHED"
          and bool(m) and int(m.group(1)) == ctr, num)
    a = await get_db().assets.find_one({"id": a1["id"]}, {"_id": 0})
    hold = await get_db().asset_holdings.find_one({"company_id": env.cid, "asset_id": a1["id"]}, {"_id": 0})
    ev = await get_db().asset_events.find_one({"company_id": env.cid, "asset_id": a1["id"], "event_type": "OPENING_EXISTING"}, {"_id": 0})
    bast = await get_db().asset_basts.find_one({"company_id": env.cid, "id": p["bast_id"]}, {"_id": 0})
    check("publish: aset IN_USE + holding ACTIVE (opening_id, tanpa handover) + event OPENING_EXISTING",
          a["lifecycle_state"] == "IN_USE" and hold and hold["holding_status"] == "ACTIVE" and hold["opening_id"] == d["id"]
          and not hold.get("handover_id") and hold["active_lock"] == a1["id"] and ev and ev["bast_id"] == p["bast_id"])
    check("BAST EXISTING: source OPENING, snapshot immutable berisi aset", bast["bast_type"] == "EXISTING" and bast["source_type"] == "OPENING"
          and bast["snapshot"]["items"][0]["asset_code"] == a1["asset_code"])
    r = await env.c.post("/api/asset-openings", json={**body}, headers=env.dr)
    check("duplicate opening (aset sudah punya holding) ditolak", r.status_code == 409, str(r.status_code))
    r = await env.c.post("/api/asset-openings/save-and-publish", json={"employee_id": env.e2, "opening_date": TODAY, "notes": "SP",
                         "items": [{"asset_id": a2["id"], "condition_id": env.cond["BAIK"]}, {"asset_id": a3["id"], "condition_id": env.cond["LECET"]}]}, headers=env.fz)
    sp = r.json()
    m2 = EXS_RE.match(sp.get("bast_number") or "")
    check("Simpan & Publish baru: 1 BAST-EXS untuk 2 aset, nomor berikutnya", r.status_code == 201 and bool(m2) and int(m2.group(1)) == ctr + 1
          and len(sp["items"]) == 2, str(r.status_code))
    dmg = await mk_asset(env, env.ga, "Rusak")
    r = await env.c.post(f"/api/assets/{dmg['id']}/status", json={"status_id": env.st["DAMAGED"], "reason": "Uji CP3"}, headers=env.ga)
    assert r.status_code == 200, r.text
    r = await env.c.post("/api/asset-openings", json={**body, "items": [{"asset_id": dmg["id"], "condition_id": env.cond["BAIK"]}]}, headers=env.fz)
    check("aset DAMAGED ditolak untuk saldo awal", r.status_code == 409, str(r.status_code))
    # siklus CP2 setelah opening
    r = await env.c.post("/api/asset-returns/save-and-publish", json={"employee_id": env.e1, "return_date": TODAY,
                         "items": [{"holding_id": hold["id"], "condition_id": env.cond["BAIK"]}]}, headers=env.fz)
    a = await get_db().assets.find_one({"id": a1["id"]}, {"_id": 0})
    check("setelah opening: Pengembalian CP2 berjalan -> PENDING_INSPECTION", r.status_code == 201 and a["lifecycle_state"] == "PENDING_INSPECTION", str(r.status_code) + r.text[:120])
    r = await env.c.get(f"/api/asset-basts/{p['bast_id']}/pdf", headers=env.fz)
    txt = pdf_text(r.content)
    check("PDF existing: on-the-fly, judul Saldo Awal + pernyataan + WIB", r.status_code == 200 and r.headers["content-type"] == "application/pdf"
          and "SALDO AWAL" in txt and "penguasaan karyawan pada saat pencatatan awal sistem" in txt and "WIB" in txt, txt[:160])
    r = await env.c.get("/api/asset-basts", params={"bast_type": "EXISTING", "limit": 100}, headers=env.ga)
    check("Dokumen BAST: filter EXISTING", r.status_code == 200 and r.json()["total"] == 2 and all(x["bast_type"] == "EXISTING" for x in r.json()["items"]))
    env.op_assets = (a2, a3)


async def t_atomic(env):
    a, b = await mk_asset(env, env.ga, "AtomA"), await mk_asset(env, env.ga, "AtomB")
    hid = await T21.fake_active_holding(env, b["id"], env.e3)     # state inkonsisten: READY tetapi ada holding aktif
    ctr = await counter(env, "BAST_EXISTING", str(YEAR))
    n_op, n_bast = await count(env, "asset_openings"), await count(env, "asset_basts")
    body = {"employee_id": env.e1, "opening_date": TODAY, "items": [{"asset_id": a["id"], "condition_id": env.cond["BAIK"]}]}
    # validasi awal lolos untuk A; B ditambahkan langsung ke draft lalu Simpan & Publish -> gagal DI DALAM transaksi
    r = await env.c.post("/api/asset-openings", json=body, headers=env.fz)
    op = r.json()
    await get_db().asset_opening_items.insert_one({"id": new_id(), "company_id": env.cid, "status": "active", "opening_id": op["id"],
                                                   "asset_id": b["id"], "line_no": 2, "condition_id": env.cond["BAIK"], "created_at": now(), "updated_at": now()})
    r = await env.c.post(f"/api/asset-openings/{op['id']}/publish", headers=env.fz)
    aa = await get_db().assets.find_one({"id": a["id"]}, {"_id": 0})
    cur = await get_db().asset_openings.find_one({"id": op["id"]}, {"_id": 0})
    check("atomic rollback: publish gagal (409) -> aset lain tetap READY, tanpa holding/BAST, counter tetap, draft tetap DRAFT",
          r.status_code == 409 and aa["lifecycle_state"] == "READY" and await count(env, "asset_holdings", asset_id=a["id"]) == 0
          and await count(env, "asset_basts") == n_bast and await counter(env, "BAST_EXISTING", str(YEAR)) == ctr and cur["doc_state"] == "DRAFT",
          f"{r.status_code} {r.text[:120]}")
    r = await env.c.post(f"/api/asset-openings/{op['id']}/save-and-publish", json={**body, "notes": "UBAH", "items": [
        {"asset_id": a["id"], "condition_id": env.cond["BAIK"]}]}, headers=env.fz)
    # tanpa B -> lolos; buktikan jalur gagal dgn draft baru yang berisi aset berpemegang
    ok = r.status_code == 200
    op2 = (await env.c.post("/api/asset-openings", json={**body, "notes": "ASLI", "items": [{"asset_id": (await mk_asset(env, env.ga, "AtomC"))["id"],
                                                                                              "condition_id": env.cond["BAIK"]}]}, headers=env.fz)).json()
    await get_db().asset_holdings.delete_one({"id": hid})
    hid2 = await T21.fake_active_holding(env, op2["items"][0]["asset_id"], env.e3)
    r = await env.c.post(f"/api/asset-openings/{op2['id']}/save-and-publish", json={**body, "notes": "BERUBAH",
                         "items": [{"asset_id": op2["items"][0]["asset_id"], "condition_id": env.cond["LECET"]}]}, headers=env.fz)
    cur = await get_db().asset_openings.find_one({"id": op2["id"]}, {"_id": 0})
    item = await get_db().asset_opening_items.find_one({"opening_id": op2["id"]}, {"_id": 0})
    check("Simpan & Publish draft existing gagal -> draft kembali ke versi tersimpan (notes & kondisi lama)",
          ok and r.status_code == 409 and cur["notes"] == "ASLI" and item["condition_id"] == env.cond["BAIK"] and cur["doc_state"] == "DRAFT",
          f"{r.status_code} {cur.get('notes')}")
    await get_db().asset_holdings.delete_one({"id": hid2})
    check("tidak ada opening parsial tercatat", await count(env, "asset_openings") == n_op + 2)


async def t_bulk_opening(env):
    ready = env.st_code["READY"]
    rows = [mrow(i, status=ready, legacy_code=f"INV-B{i}-{RUN}", serial_number=f"SN-B{i}-{RUN}", project=env.proj_code[env.p1]) for i in range(1, 6)]
    b = (await preview(env, env.fz, "MASTER", rows)).json()
    c = (await env.c.post(f"/api/asset-imports/{b['id']}/commit", json={}, headers=env.fz)).json()
    codes = c["new_asset_codes"]
    emp = {x["id"]: x for x in await get_db().employees.find({"id": {"$in": [env.e1, env.e2, env.e3]}}, {"_id": 0}).to_list(5)}
    en1, en2, en3 = emp[env.e1]["employee_number"], emp[env.e2]["employee_number"], emp[env.e3]["employee_number"]
    r = await env.c.get("/api/asset-imports/template", params={"import_type": "OPENING"}, headers=env.fz)
    check("template saldo awal .xlsx tersedia", r.status_code == 200 and [c.value for c in load_workbook(io.BytesIO(r.content)).worksheets[0][1]]
          == [lbl for _, lbl in OPENING_COLS])
    bad = [{"employee_number": "TIDAK-ADA", "asset_code": codes[0], "opening_date": TODAY, "condition": "BAIK"},
           {"employee_number": en1, "asset_code": "AST-999999", "opening_date": TODAY, "condition": "BAIK"},
           {"employee_number": en1, "legacy_code": f"INV-B2-{RUN}", "opening_date": TODAY, "condition": "BAIK"},
           {"employee_number": en1, "serial_number": f"SN-B2-{RUN}", "opening_date": TODAY, "condition": "BAIK"},
           {"employee_number": en1, "opening_date": TODAY, "condition": "BAIK"},
           {"employee_number": en2, "asset_code": codes[3], "opening_date": "2099-01-01", "condition": "BAIK"},
           {"employee_number": en2, "asset_code": codes[4], "opening_date": "2024-02-30", "condition": "BAIK"},
           {"employee_number": en3, "asset_code": env.op_assets[0]["asset_code"], "opening_date": TODAY, "condition": "BAIK"}]
    b = (await preview(env, env.fz, "OPENING", bad)).json()
    text = msgs(await rows_of(env, env.fz, b["id"], "ERROR"))
    for label, needle in (("karyawan tidak ditemukan", "Nomor Karyawan 'TIDAK-ADA' tidak ditemukan"), ("aset tidak ditemukan", "AST-999999"),
                          ("aset ganda di file", "muncul lebih dari sekali"), ("identitas aset kosong", "Isi salah satu"),
                          ("tanggal masa depan", "masa depan"), ("tanggal tidak valid", "'2024-02-30' tidak valid"),
                          ("aset sudah dipegang", "berstatus Dipakai"), ("grup tanggal/proyek beda", "satu karyawan = satu dokumen")):
        check(f"bulk opening ERROR: {label}", needle in text, text[:160] if needle not in text else "")
    good = [{"employee_number": en1, "asset_code": codes[0], "opening_date": TODAY, "condition": "BAIK", "accessories": "Tas"},
            {"employee_number": en1, "legacy_code": f"INV-B2-{RUN}", "opening_date": TODAY, "condition": "LECET"},
            {"employee_number": en2, "serial_number": f"SN-B3-{RUN}", "opening_date": TODAY, "condition": "BAIK"}]
    b = (await preview(env, env.dr, "OPENING", good)).json()
    check("bulk opening preview bersih (identifikasi code/kode lama/serial)", b["error_rows"] == 0 and b["total_rows"] == 3, str(b)[:160])
    sm = b.get("summary") or {}
    check("preview saldo awal: ringkasan grouping per karyawan (2 karyawan, E1 = 2 aset)", sm.get("employee_groups") == 2
          and {g["employee_number"]: g["assets"] for g in sm.get("groups", [])} == {en1: 2, en2: 1}, str(sm)[:200])
    r = await env.c.post(f"/api/asset-imports/{b['id']}/commit", json={"mode": "publish"}, headers=env.dr)
    check("bulk opening: tanpa commit/publish -> 403", r.status_code == 403, str(r.status_code))
    ctr = await counter(env, "BAST_EXISTING", str(YEAR))
    r = await env.c.post(f"/api/asset-imports/{b['id']}/commit", json={"mode": "publish"}, headers=env.fz)
    c = r.json()
    nums = sorted(int(EXS_RE.match(x).group(1)) for x in c.get("bast_numbers", []) if EXS_RE.match(x))
    check("bulk opening publish: grouping per karyawan -> 2 opening / 2 BAST-EXS berurutan", r.status_code == 200 and c["opening_count"] == 2
          and nums == [ctr, ctr + 1], str(c)[:200])
    ops = await get_db().asset_openings.find({"company_id": env.cid, "import_batch_id": b["id"]}, {"_id": 0}).to_list(5)
    e1op = next(o for o in ops if o["employee_id"] == env.e1)
    check("karyawan E1: 2 aset dalam satu opening PUBLISHED", e1op["doc_state"] == "PUBLISHED"
          and await count(env, "asset_opening_items", opening_id=e1op["id"]) == 2
          and await count(env, "asset_holdings", opening_id=e1op["id"], holding_status="ACTIVE") == 2)
    r = await env.c.get(f"/api/asset-imports/{b['id']}/result", headers=env.fz)
    data = [[x.value for x in row] for row in load_workbook(io.BytesIO(r.content)).worksheets[0].iter_rows()]
    check("hasil impor saldo awal: nomor BAST-EXS per baris", r.status_code == 200 and all(EXS_RE.match(str(x[5])) for x in data[1:]) and len(data) == 4)
    b = (await preview(env, env.fz, "OPENING", [{"employee_number": en3, "asset_code": codes[4], "opening_date": TODAY,
                                                  "project": env.proj_code[env.p2], "condition": "BAIK"}])).json()
    r = await env.c.post(f"/api/asset-imports/{b['id']}/commit", json={"mode": "draft"}, headers=env.fz)
    op = await get_db().asset_openings.find_one({"company_id": env.cid, "import_batch_id": b["id"]}, {"_id": 0})
    a = await get_db().assets.find_one({"company_id": env.cid, "asset_code": codes[4]}, {"_id": 0})
    check("bulk opening mode draft: opening DRAFT, aset tetap READY", r.status_code == 200 and op["doc_state"] == "DRAFT" and a["lifecycle_state"] == "READY")
    r = await env.c.post(f"/api/asset-openings/{op['id']}/cancel", headers=env.dr)
    op2 = await get_db().asset_openings.find_one({"company_id": env.cid, "id": op["id"]}, {"_id": 0})
    a = await get_db().assets.find_one({"company_id": env.cid, "asset_code": codes[4]}, {"_id": 0})
    check("batalkan draft saldo awal (asset_opening:edit): CANCELLED, aset tetap READY, tanpa BAST",
          r.status_code == 200 and op2["doc_state"] == "CANCELLED" and a["lifecycle_state"] == "READY" and not op2.get("bast_id"), str(r.status_code))
    r = await env.c.post(f"/api/asset-openings/{op['id']}/publish", headers=env.fz)
    check("draft yang dibatalkan tidak dapat dipublish (409)", r.status_code == 409, str(r.status_code))


async def t_scope(env):
    r = await env.c.get("/api/asset-openings/employee-search", params={"q": RUN, "limit": 50}, headers=env.r1)
    ids = {x["id"] for x in r.json()["items"]}
    check("01I: pencarian karyawan saldo awal hanya dalam scope", env.e1 in ids and env.e3 not in ids and env.eb not in ids, str(len(ids)))
    a = await mk_asset(env, env.ga, "ScopeP1")
    r = await env.c.post("/api/asset-openings", json={"employee_id": env.e3, "opening_date": TODAY,
                         "items": [{"asset_id": a["id"], "condition_id": env.cond["BAIK"]}]}, headers=env.r1)
    check("01I: opening untuk karyawan di luar scope -> 404", r.status_code == 404, str(r.status_code))
    op = await get_db().asset_openings.find_one({"company_id": env.cid, "project_id": env.p2}, {"_id": 0})
    r = await env.c.get(f"/api/asset-openings/{op['id']}", headers=env.r1)
    check("01I: UUID opening di luar scope -> 404 generik", r.status_code == 404, str(r.status_code))
    r = await env.c.get(f"/api/asset-openings/{op['id']}", headers=env.adminb)
    check("tenant lain: UUID opening -> 404/403", r.status_code in (403, 404), str(r.status_code))
    b = (await preview(env, env.r1, "MASTER", [mrow(1, status=env.st_code["READY"], project=env.proj_code[env.p2]),
                                               mrow(2, status=env.st_code["READY"])])).json()
    text = msgs(await rows_of(env, env.r1, b["id"], "ERROR"))
    check("01I: impor master project di luar scope / kosong -> ERROR", "di luar cakupan data Anda" in text and "Proyek wajib diisi" in text, text[:160])
    emp3 = (await get_db().employees.find_one({"id": env.e3}, {"_id": 0}))["employee_number"]
    b2 = (await preview(env, env.r1, "OPENING", [{"employee_number": emp3, "asset_code": a["asset_code"], "opening_date": TODAY, "condition": "BAIK"}])).json()
    text = msgs(await rows_of(env, env.r1, b2["id"], "ERROR"))
    check("01I: impor saldo awal karyawan di luar scope -> ERROR", "di luar cakupan data Anda" in text, text[:160])
    other = await get_db().asset_import_batches.find_one({"company_id": env.cid, "uploaded_by": env.fz_id}, {"_id": 0})
    r = await env.c.get(f"/api/asset-imports/{other['id']}", headers=env.r1)
    lst = (await env.c.get("/api/asset-imports", params={"limit": 100}, headers=env.r1)).json()
    check("01I: restricted hanya melihat batch miliknya (batch lain -> 404)", r.status_code == 404
          and all(x["id"] in (b["id"], b2["id"]) for x in lst["items"]), str(r.status_code))


async def t_concurrency(env):
    assets = [await mk_asset(env, env.ga, f"Conc{i}") for i in range(8)]
    ctr = await counter(env, "BAST_EXISTING", str(YEAR))

    async def sp(a, emp):
        return await env.c.post("/api/asset-openings/save-and-publish", json={"employee_id": emp, "opening_date": TODAY,
                                "items": [{"asset_id": a["id"], "condition_id": env.cond["BAIK"]}]}, headers=env.fz)
    res = await asyncio.gather(*[sp(a, env.e1 if i % 2 else env.e2) for i, a in enumerate(assets)])
    bad = [f"{r.status_code}:{r.text[:80]}" for r in res if r.status_code != 201]
    if bad:
        print("  concurrency non-201:", bad)
    nums = sorted(int(EXS_RE.match(r.json()["bast_number"]).group(1)) for r in res if r.status_code == 201)
    check("konkurensi: 8 Simpan & Publish paralel -> 8 BAST-EXS unik & berurutan", len(nums) == 8 and nums == list(range(ctr, ctr + 8)), str(nums))
    x = await mk_asset(env, env.ga, "ConcSame")
    res = await asyncio.gather(sp(x, env.e1), sp(x, env.e2))
    codes = sorted(r.status_code for r in res)
    check("konkurensi: aset sama di 2 opening paralel -> tepat 1 berhasil, 1 ditolak",
          codes[0] == 201 and codes[1] in (409, 422) and await count(env, "asset_holdings", asset_id=x["id"], holding_status="ACTIVE") == 1, str(codes))
    dup = await get_db().asset_holdings.find({"company_id": env.cid, "holding_status": "ACTIVE"}, {"_id": 0}).to_list(20000)
    per = {}
    for h in dup:
        per[h["asset_id"]] = per.get(h["asset_id"], 0) + 1
    in_use = await get_db().assets.find({"company_id": env.cid, "lifecycle_state": "IN_USE"}, {"_id": 0}).to_list(20000)
    held = set(per)
    check("invariant: maks. 1 holding aktif per aset & setiap IN_USE punya holding", max(per.values() or [1]) == 1
          and all(a["id"] in held for a in in_use))


async def t_perf(env):
    ready = env.st_code["READY"]
    rows = [mrow(i, status=ready, legacy_code=f"P{RUN}-{i:05d}", serial_number=f"PSN{RUN}-{i:05d}", project=env.proj_code[env.p1],
                 acquisition_year=2020, acquisition_value="1000000") for i in range(1, 5001)]
    ctr = await counter(env, "ASSET_CODE", "ALL")
    t0 = time.perf_counter()
    r = await preview(env, env.fz, "MASTER", rows, "perf5000.xlsx")
    t_prev = time.perf_counter() - t0
    b = r.json()
    check("performa: preview 5.000 baris (201, 0 ERROR)", r.status_code == 201 and b["total_rows"] == 5000 and b["error_rows"] == 0, str(b)[:160])
    t1 = time.perf_counter()
    r = await env.c.post(f"/api/asset-imports/{b['id']}/commit", json={}, headers=env.fz)
    t_commit = time.perf_counter() - t1
    c = r.json()
    n = await count(env, "assets", import_batch_id=b["id"])
    distinct = len({x["asset_code"] for x in await get_db().assets.find({"company_id": env.cid, "import_batch_id": b["id"]}, {"_id": 0}).to_list(6000)})
    check("performa: commit 5.000 aset satu transaksi, kode unik, counter +5000", r.status_code == 200 and c["created_count"] == 5000 and n == 5000
          and distinct == 5000 and await counter(env, "ASSET_CODE", "ALL") == ctr + 5000, str(r.status_code))
    env.perf = {"preview_s": round(t_prev, 2), "commit_s": round(t_commit, 2)}
    check(f"performa: preview {t_prev:.1f}s + commit {t_commit:.1f}s (< 120s total)", t_prev + t_commit < 120)


async def cp3_orphans() -> int:
    async with dbm.get_engine().connect() as conn:
        total = 0
        for q in ("SELECT COUNT(*) FROM asset_opening_items i LEFT JOIN asset_openings o ON o.id = i.opening_id WHERE o.id IS NULL",
                  "SELECT COUNT(*) FROM asset_import_rows r LEFT JOIN asset_import_batches b ON b.id = r.batch_id WHERE b.id IS NULL",
                  "SELECT COUNT(*) FROM asset_holdings h LEFT JOIN asset_openings o ON o.id = h.opening_id WHERE h.opening_id IS NOT NULL AND o.id IS NULL"):
            total += (await conn.execute(sa.text(q))).scalar() or 0
        return total


async def main():
    env = T.Env()
    env.perf = {}
    rep = {"residue": -1, "orphan": -1, "detail": ["cleanup tidak berjalan"]}
    async with dbm.get_engine().connect() as conn:
        dbname = (await conn.execute(sa.text("SELECT DATABASE()"))).scalar()
    if dbname != "hris_dev":
        raise SystemExit(f"DITOLAK: database aktual {dbname!r} bukan hris_dev.")
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=server.app), base_url="http://t03", timeout=300) as c:
        env.c = c
        try:
            await setup3(env)
            for t in (t_rbac, t_master_import, t_opening_manual, t_atomic, t_bulk_opening, t_scope, t_concurrency, t_perf):
                try:
                    await t(env)
                except Exception as exc:  # noqa: BLE001
                    import traceback
                    traceback.print_exc()
                    check(f"{t.__name__} (exception)", False, repr(exc)[:200])
        finally:
            rep = await T.cleanup(env)
    orphan3 = await cp3_orphans()
    await dbm.close_db()
    check("fixture cleanup: residue 0 (hanya data bertanda run ini)", rep["residue"] == 0, str(rep["detail"]))
    check("fixture cleanup: orphan 0 (CP2 + CP3)", rep["orphan"] == 0 and orphan3 == 0, f"{rep['detail']} cp3={orphan3}")
    failed = [r for r in T.RESULTS if not r[1]]
    print(f"\n== CP3 targeted: {len(T.RESULTS) - len(failed)}/{len(T.RESULTS)} PASS ==")
    print(f"Performa 5.000 baris: {env.perf}")
    print(f"Fixture cleanup: {'PASS' if rep['residue'] == 0 and rep['orphan'] == 0 and orphan3 == 0 else 'FAIL'} — residue {rep['residue']}, "
          f"orphan {rep['orphan'] + orphan3} (marker {MARK})")
    return failed


def test_asset_cp3():
    assert not asyncio.run(main())


if __name__ == "__main__":
    sys.exit(1 if asyncio.run(main()) else 0)
