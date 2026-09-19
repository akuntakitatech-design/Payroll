"""Impor massal data karyawan dari Excel (openpyxl).

Alur: unduh template -> isi -> unggah -> validasi (dry-run) -> commit.
Validasi mengembalikan pesan per baris dalam Bahasa Indonesia sehingga HR
bisa memperbaiki file sebelum data masuk ke database.
"""
from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .ter import PTKP_STATUSES

SHEET_DATA = "Data Karyawan"
SHEET_REF = "Referensi"

GENDERS = ["L", "P"]
MARITAL = ["belum_kawin", "kawin", "duda", "janda"]
RELIGIONS = ["islam", "kristen", "katolik", "hindu", "buddha", "konghucu", "lainnya"]

# (kolom excel, field db, wajib?, keterangan, master lookup key)
COLUMNS: List[Tuple[str, str, bool, str, Optional[str]]] = [
    ("Nama Lengkap*", "full_name", True, "Minimal 2 karakter", None),
    ("NIK Karyawan", "employee_number", False, "Kosongkan agar dibuat otomatis", None),
    ("Nomor KTP", "nik", False, "16 digit angka", None),
    ("NPWP", "npwp", False, "Opsional", None),
    ("Status PTKP", "ptkp_status", False, "TK/0 s.d. K/3 — dipakai untuk PPh 21", None),
    ("Jenis Kelamin", "gender", False, "L atau P", None),
    ("Tempat Lahir", "birth_place", False, "", None),
    ("Tanggal Lahir", "birth_date", False, "Format YYYY-MM-DD", None),
    ("Status Pernikahan", "marital_status", False, "belum_kawin/kawin/duda/janda", None),
    ("Agama", "religion", False, "islam/kristen/katolik/hindu/buddha/konghucu/lainnya", None),
    ("Pendidikan", "education", False, "", None),
    ("Email", "email", False, "Format email valid", None),
    ("Telepon", "phone", False, "", None),
    ("Alamat", "address", False, "", None),
    ("Kota", "city", False, "", None),
    ("Jabatan (teks)", "job_title", False, "Nama jabatan bebas", None),
    ("Tanggal Masuk", "join_date", False, "Format YYYY-MM-DD", None),
    ("Status Kepegawaian", "employment_status_id", False, "Harus ada di master", "employment_statuses"),
    ("Cabang", "branch_id", False, "Harus ada di master", "branches"),
    ("Lokasi Kerja", "work_location_id", False, "Harus ada di master", "work_locations"),
    ("Departemen", "department_id", False, "Harus ada di master", "departments"),
    ("Divisi", "division_id", False, "Harus ada di master", "divisions"),
    ("Jabatan (master)", "position_id", False, "Harus ada di master", "positions"),
    ("Grade / Level", "job_grade_id", False, "Harus ada di master", "job_grades"),
    ("Cost Center", "cost_center_id", False, "Harus ada di master", "cost_centers"),
    ("Proyek", "project_id", False, "Harus ada di master", "projects"),
    ("Gaji Pokok", "basic_salary", False, "Angka tanpa titik, mis. 8000000", None),
    ("Nama Bank", "bank_name", False, "", None),
    ("No. Rekening", "bank_account_number", False, "", None),
    ("Nama Pemilik Rekening", "bank_account_name", False, "", None),
    ("No. BPJS Kesehatan", "bpjs_kesehatan_number", False, "", None),
    ("No. BPJS Ketenagakerjaan", "bpjs_tk_number", False, "", None),
    ("Catatan", "notes", False, "", None),
]

MASTER_LABELS = {
    "employment_statuses": "Status Kepegawaian",
    "branches": "Cabang",
    "work_locations": "Lokasi Kerja",
    "departments": "Departemen",
    "divisions": "Divisi",
    "positions": "Jabatan",
    "job_grades": "Grade / Level",
    "cost_centers": "Cost Center",
    "projects": "Proyek",
}

HEADER_FILL = PatternFill("solid", fgColor="0F5C4F")
REQ_FILL = PatternFill("solid", fgColor="B45309")
HINT_FILL = PatternFill("solid", fgColor="F1F5F9")


def build_employee_template(masters: Optional[Dict[str, List[Dict[str, Any]]]] = None) -> bytes:
    """Buat file template Excel berisi header, baris petunjuk dan sheet referensi."""
    masters = masters or {}
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_DATA

    for idx, (label, _field, required, hint, _lookup) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=label)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = REQ_FILL if required else HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        hint_cell = ws.cell(row=2, column=idx, value=hint)
        hint_cell.font = Font(italic=True, size=8, color="475569")
        hint_cell.fill = HINT_FILL
        hint_cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(idx)].width = max(16, min(30, len(label) + 6))

    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 28
    ws.freeze_panes = "A3"

    # ---- sheet referensi ----
    ref = wb.create_sheet(SHEET_REF)
    ref["A1"] = "Nilai yang valid untuk kolom bertipe master data"
    ref["A1"].font = Font(bold=True, size=11, color="0F5C4F")
    ref["A2"] = "Tulis PERSIS seperti daftar di bawah (boleh kode atau nama)."
    ref["A2"].font = Font(italic=True, size=9, color="475569")

    col = 1
    blocks = [("Status PTKP", PTKP_STATUSES), ("Jenis Kelamin", GENDERS),
              ("Status Pernikahan", MARITAL), ("Agama", RELIGIONS)]
    for key, label in MASTER_LABELS.items():
        values = [str(m.get("name") or m.get("code") or "") for m in (masters.get(key) or [])]
        blocks.append((label, [v for v in values if v]))

    for label, values in blocks:
        head = ref.cell(row=4, column=col, value=label)
        head.font = Font(bold=True, size=10, color="FFFFFF")
        head.fill = HEADER_FILL
        head.alignment = Alignment(horizontal="center")
        ref.column_dimensions[get_column_letter(col)].width = max(18, len(label) + 4)
        for i, val in enumerate(values, start=5):
            ref.cell(row=i, column=col, value=val)
        col += 1

    # dropdown untuk PTKP / gender / marital / agama
    for label, values, target_field in [
        ("Status PTKP", PTKP_STATUSES, "ptkp_status"),
        ("Jenis Kelamin", GENDERS, "gender"),
        ("Status Pernikahan", MARITAL, "marital_status"),
        ("Agama", RELIGIONS, "religion"),
    ]:
        col_idx = next((i for i, c in enumerate(COLUMNS, start=1) if c[1] == target_field), None)
        if not col_idx:
            continue
        dv = DataValidation(
            type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
            showDropDown=False,
        )
        dv.error = f"Nilai tidak valid untuk {label}."
        dv.errorTitle = "Input tidak valid"
        ws.add_data_validation(dv)
        letter = get_column_letter(col_idx)
        dv.add(f"{letter}3:{letter}1000")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _clean(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _parse_date(value: Any) -> Tuple[Optional[str], Optional[str]]:
    text = _clean(value)
    if not text:
        return None, None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat(), None
        except ValueError:
            continue
    return None, f"format tanggal '{text}' tidak dikenali (gunakan YYYY-MM-DD)"


def _parse_number(value: Any) -> Tuple[Optional[float], Optional[str]]:
    text = _clean(value)
    if not text:
        return None, None
    normalized = re.sub(r"[^\d,.\-]", "", text).replace(".", "").replace(",", ".")
    try:
        num = float(normalized)
    except ValueError:
        return None, f"nilai '{text}' bukan angka yang valid"
    if num < 0:
        return None, "nilai tidak boleh negatif"
    return num, None


def parse_employee_rows(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Baca file Excel dan kembalikan baris mentah (belum divalidasi)."""
    try:
        wb = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"File tidak dapat dibaca sebagai Excel (.xlsx): {exc}") from exc

    ws = wb[SHEET_DATA] if SHEET_DATA in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("File Excel kosong.")

    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    expected = [c[0] for c in COLUMNS]
    if header[: len(expected)] != expected:
        # toleransi: cocokkan berdasarkan nama kolom yang ada
        mapping = {}
        for idx, name in enumerate(header):
            match = next((c for c in COLUMNS if c[0] == name), None)
            if match:
                mapping[match[1]] = idx
        if "full_name" not in mapping:
            raise ValueError(
                "Header file tidak sesuai template. Unduh ulang template dan jangan "
                "mengubah baris judul kolom."
            )
    else:
        mapping = {c[1]: i for i, c in enumerate(COLUMNS)}

    out = []
    for excel_row, raw in enumerate(rows[1:], start=2):
        values = {field: (raw[i] if i < len(raw) else None) for field, i in mapping.items()}
        # baris petunjuk dari template
        if excel_row == 2 and _clean(values.get("full_name")) in (None, "Minimal 2 karakter"):
            continue
        if all(_clean(v) is None for v in values.values()):
            continue
        out.append({"excel_row": excel_row, "values": values})
    return out


def validate_employee_rows(
    raw_rows: List[Dict[str, Any]],
    masters: Dict[str, List[Dict[str, Any]]],
    existing_employee_numbers: Optional[set] = None,
    existing_niks: Optional[set] = None,
) -> Dict[str, Any]:
    """Validasi baris dan kembalikan laporan siap tampil di UI.

    `masters` = {collection_name: [{id, code, name}, ...]}
    """
    existing_employee_numbers = {
        str(x).strip().lower() for x in (existing_employee_numbers or set()) if x
    }
    existing_niks = {str(x).strip().lower() for x in (existing_niks or set()) if x}

    lookups: Dict[str, Dict[str, str]] = {}
    for key in MASTER_LABELS:
        table = {}
        for item in masters.get(key) or []:
            for candidate in (item.get("name"), item.get("code"), item.get("id")):
                if candidate:
                    table[str(candidate).strip().lower()] = item["id"]
        lookups[key] = table

    seen_numbers: Dict[str, int] = {}
    seen_niks: Dict[str, int] = {}
    results: List[Dict[str, Any]] = []

    for row in raw_rows:
        excel_row = row["excel_row"]
        values = row["values"]
        errors: List[str] = []
        warnings: List[str] = []
        payload: Dict[str, Any] = {}

        for label, field, required, _hint, lookup in COLUMNS:
            raw = values.get(field)
            text = _clean(raw)

            if required and not text:
                errors.append(f"{label.rstrip('*')} wajib diisi")
                continue
            if not text:
                continue

            if field == "full_name":
                if len(text) < 2:
                    errors.append("Nama Lengkap minimal 2 karakter")
                else:
                    payload[field] = text
            elif field in ("birth_date", "join_date"):
                parsed, err = _parse_date(raw)
                if err:
                    errors.append(f"{label}: {err}")
                else:
                    payload[field] = parsed
            elif field == "basic_salary":
                parsed, err = _parse_number(raw)
                if err:
                    errors.append(f"{label}: {err}")
                else:
                    payload[field] = parsed
            elif field == "email":
                if not re.match(r"^[^@\s]+@[^@\s]+\.[A-Za-z]{2,}$", text):
                    errors.append(f"Email '{text}' tidak valid")
                else:
                    payload[field] = text.lower()
            elif field == "nik":
                digits = re.sub(r"\D", "", text)
                if len(digits) != 16:
                    errors.append(f"Nomor KTP harus 16 digit (ditemukan {len(digits)} digit)")
                elif digits.lower() in existing_niks:
                    errors.append(f"Nomor KTP {digits} sudah terdaftar di sistem")
                elif digits.lower() in seen_niks:
                    errors.append(f"Nomor KTP {digits} duplikat dengan baris {seen_niks[digits.lower()]}")
                else:
                    seen_niks[digits.lower()] = excel_row
                    payload[field] = digits
            elif field == "employee_number":
                low = text.lower()
                if low in existing_employee_numbers:
                    errors.append(f"NIK Karyawan '{text}' sudah dipakai di sistem")
                elif low in seen_numbers:
                    errors.append(f"NIK Karyawan '{text}' duplikat dengan baris {seen_numbers[low]}")
                else:
                    seen_numbers[low] = excel_row
                    payload[field] = text
            elif field == "ptkp_status":
                normalized = text.upper().replace(" ", "")
                if normalized not in PTKP_STATUSES:
                    errors.append(
                        f"Status PTKP '{text}' tidak dikenal (pilih {', '.join(PTKP_STATUSES)})"
                    )
                else:
                    payload[field] = normalized
            elif field == "gender":
                normalized = text.upper()[:1]
                if normalized not in GENDERS:
                    errors.append(f"Jenis Kelamin '{text}' tidak valid (gunakan L atau P)")
                else:
                    payload[field] = normalized
            elif field == "marital_status":
                normalized = text.lower().replace(" ", "_")
                if normalized not in MARITAL:
                    errors.append(f"Status Pernikahan '{text}' tidak valid")
                else:
                    payload[field] = normalized
            elif field == "religion":
                normalized = text.lower()
                if normalized not in RELIGIONS:
                    warnings.append(f"Agama '{text}' di luar daftar standar, disimpan apa adanya")
                payload[field] = normalized
            elif lookup:
                resolved = lookups.get(lookup, {}).get(text.lower())
                if not resolved:
                    errors.append(
                        f"{MASTER_LABELS[lookup]} '{text}' tidak ditemukan di master data"
                    )
                else:
                    payload[field] = resolved
            else:
                payload[field] = text

        results.append(
            {
                "excel_row": excel_row,
                "full_name": _clean(values.get("full_name")),
                "employee_number": _clean(values.get("employee_number")),
                "is_valid": not errors,
                "errors": errors,
                "warnings": warnings,
                "payload": payload,
            }
        )

    valid = [r for r in results if r["is_valid"]]
    invalid = [r for r in results if not r["is_valid"]]
    return {
        "total_rows": len(results),
        "valid_count": len(valid),
        "invalid_count": len(invalid),
        "can_commit": bool(valid),
        "rows": results,
        "summary": (
            f"{len(valid)} baris siap diimpor, {len(invalid)} baris perlu diperbaiki."
            if invalid
            else f"Semua {len(valid)} baris valid dan siap diimpor."
        ),
    }


def template_columns() -> List[Dict[str, Any]]:
    """Metadata kolom untuk ditampilkan di UI."""
    return [
        {"label": label, "field": field, "required": required, "hint": hint,
         "master": MASTER_LABELS.get(lookup) if lookup else None}
        for label, field, required, hint, lookup in COLUMNS
    ]
