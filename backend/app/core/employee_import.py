"""Upgrade 01E - Engine Migrasi / Impor Data Karyawan (Excel multi-sheet).

Alur: Upload -> Analyze -> Validate -> Preview (READ-ONLY terhadap data bisnis; hanya menulis metadata
batch + hasil analisis) -> Commit eksplisit (transaksi DB terpisah PER KARYAWAN).

Aturan bisnis utama:
- nomor karyawan = business key lintas sheet; import TIDAK pernah mengubah nomor karyawan;
- sel Excel kosong TIDAK menimpa nilai DB yang sudah terisi (blank-preserve);
- NIK KTP milik karyawan lain / duplikat antar karyawan di file = CONFLICT (tidak pernah merge otomatis);
- Status Karyawan memakai helper 01B (`history_doc`, `legacy_for`) + riwayat source=IMPORT;
- project/penempatan memakai helper & aturan 01D (assignment + riwayat + sinkron legacy `employees.project_id`);
  tanggal mulai penempatan TIDAK dikarang (kosong = tidak diketahui, diberi warning);
- keluarga dicocokkan aman (NIK, lalu hubungan+nama+tgl lahir); identitas ragu = warning review, tanpa merge/hapus;
- nilai sensitif (NIK/rekening/NPWP/BPJS) selalu dimasking di hasil analisis, pesan error, audit, dan log.
"""
from __future__ import annotations

import hashlib
import json
import logging
import re
from datetime import date, timedelta
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

from fastapi import HTTPException, status as http_status
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import assignment as asg
from . import employee_status as emp_status
from .audit import build_audit_entry
from .db import NO_ID, audit_fields, new_id, now, transaction
from .excel import _clean, _parse_date
from .employee_import_mapping import (
    BOOL_FALSE, BOOL_TRUE, DATA_SHEETS, DIGIT_RULES, ENUMS, LIMITS, MAPPING_VERSION, MASTER_SOURCES, SHEETS,
    SOFT_ENUMS, Col, employee_key, norm_header,
)
from .sensitive import mask_value
from .tenancy import get_tenant_db

logger = logging.getLogger(__name__)

NEW, UPDATE, UNCHANGED, CONFLICT, ERROR = "NEW", "UPDATE", "UNCHANGED", "CONFLICT", "ERROR"
CLASSES = (NEW, UPDATE, UNCHANGED, CONFLICT, ERROR)
FAMILY_TABLE = "employee_family_members"
LONG_TEXT_FIELDS = {"address", "domicile_address", "notes"}  # kolom TEXT (tanpa batas 255)


class ImportRowError(Exception):
    """Kegagalan terkontrol saat commit satu karyawan (pesan aman ditampilkan, tanpa nilai sensitif)."""


# ============================================================ template
HEADER_FILL = PatternFill("solid", fgColor="0F5C4F")
REQ_FILL = PatternFill("solid", fgColor="B45309")
HINT_FILL = PatternFill("solid", fgColor="F1F5F9")

PETUNJUK_LINES = [
    "PETUNJUK MIGRASI / IMPOR DATA KARYAWAN (Upgrade 01E - mapping draft " + MAPPING_VERSION + ")",
    "",
    "1. Jangan mengubah nama sheet dan judul kolom (baris 1). Baris 2 = petunjuk kolom. Data mulai baris 3.",
    "2. Nomor Karyawan adalah kunci lintas sheet: tulis PERSIS sama di DATA_KARYAWAN, KEPEGAWAIAN, BANK_BPJS, KELUARGA.",
    "3. Nomor Karyawan yang sudah ada = UPDATE; yang belum ada = karyawan BARU (Nama Lengkap wajib).",
    "4. Sel KOSONG tidak akan menghapus data yang sudah tersimpan (data lama dipertahankan).",
    "5. NIK KTP yang sudah dipakai karyawan lain akan ditandai CONFLICT dan tidak disimpan.",
    "6. Perubahan Project = Pindah Penempatan: isi Tanggal Mulai Penempatan (tidak boleh di masa depan).",
    "7. Perubahan Status Karyawan (Aktif/Standby/Tidak Aktif) wajib disertai Tanggal Efektif Status.",
    "8. Status Kepegawaian (PKWT/PKWTT) BERBEDA dengan Status Karyawan.",
    "9. Kolom NIK/Rekening/NPWP/BPJS: format sel sebagai TEKS agar angka tidak berubah.",
    "10. KELUARGA boleh banyak baris per karyawan. Baris yang sama tidak akan digandakan saat impor ulang.",
    "11. Nilai master (Jabatan, Departemen, Project, dll.) harus sesuai sheet MASTER_REFERENCE (kode atau nama).",
    "12. Proses: Unggah -> Analisis & Preview (belum mengubah data) -> Commit (disimpan per karyawan).",
]


def build_template(masters: Dict[str, List[Dict[str, Any]]], company: Dict[str, Any]) -> bytes:
    wb = Workbook()
    first = True
    for spec in SHEETS:
        ws = wb.active if first else wb.create_sheet(spec.name)
        first = False
        ws.title = spec.name
        if spec.name == "PETUNJUK":
            for i, line in enumerate(PETUNJUK_LINES, start=1):
                ws.cell(row=i, column=1, value=line).font = Font(bold=(i == 1), size=11 if i == 1 else 10)
            ws.cell(row=len(PETUNJUK_LINES) + 2, column=1,
                    value=f"Perusahaan: {company.get('name') or '-'} (kode {company.get('code') or '-'})")
            ws.column_dimensions["A"].width = 120
            continue
        if spec.name == "MASTER_REFERENCE":
            _write_reference(ws, masters)
            continue
        for idx, col in enumerate(spec.columns, start=1):
            required = col.required or col.required_new
            head = ws.cell(row=spec.header_row, column=idx, value=col.headers[0])
            head.font = Font(bold=True, color="FFFFFF", size=10)
            head.fill = REQ_FILL if required else HEADER_FILL
            head.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            hint = ws.cell(row=spec.header_row + 1, column=idx, value=col.hint or ("Sensitif" if col.sensitive else ""))
            hint.font = Font(italic=True, size=8, color="475569")
            hint.fill = HINT_FILL
            hint.alignment = Alignment(wrap_text=True, vertical="top")
            letter = get_column_letter(idx)
            ws.column_dimensions[letter].width = max(16, min(34, len(col.headers[0]) + 6))
            if col.type in ("digits", "text", "phone") or col.key == "employee_number":
                ws.column_dimensions[letter].number_format = "@"
            values = _dropdown_values(col)
            if values:
                dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{letter}{spec.data_start_row}:{letter}6000")
        ws.row_dimensions[spec.header_row].height = 32
        ws.freeze_panes = f"B{spec.data_start_row}"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _dropdown_values(col: Col) -> List[str]:
    return {"gender": ["L", "P"], "marital": ["Belum Menikah", "Menikah", "Cerai"],
            "relationship": ["Suami", "Istri", "Anak", "Ayah", "Ibu", "Saudara", "Lainnya"]}.get(
        col.validator or "", ["Ya", "Tidak"] if col.type == "bool" else [])


def _write_reference(ws, masters: Dict[str, List[Dict[str, Any]]]) -> None:
    ws["A1"] = "Nilai valid untuk kolom master (tulis kode ATAU nama persis seperti di bawah)"
    ws["A1"].font = Font(bold=True, size=11, color="0F5C4F")
    blocks: List[Tuple[str, List[str]]] = []
    for key, (_table, label) in MASTER_SOURCES.items():
        rows = masters.get(key) or []
        blocks.append((label, [f"{r.get('code') or ''} | {r.get('name') or ''}".strip(" |") for r in rows]))
    blocks += [("Jenis Kelamin", ["L", "P"]), ("Status Pernikahan", ["Belum Menikah", "Menikah", "Cerai"]),
               ("Agama", sorted({v for v in ENUMS["religion"].values()})),
               ("Pendidikan", ["SD", "SMP", "SMA", "D3", "S1", "S2", "S3"]),
               ("Hubungan Keluarga", ["Suami", "Istri", "Anak", "Ayah", "Ibu", "Saudara", "Lainnya"])]
    for c, (label, values) in enumerate(blocks, start=1):
        head = ws.cell(row=3, column=c, value=label)
        head.font = Font(bold=True, size=10, color="FFFFFF")
        head.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(c)].width = max(20, len(label) + 4)
        for i, v in enumerate(values, start=4):
            ws.cell(row=i, column=c, value=v)


# ============================================================ parse (upload)
def parse_workbook(raw: bytes) -> Tuple[Dict[str, List[Dict[str, Any]]], List[str], List[str]]:
    """Baca workbook -> {sheet: [{row, values{key: raw}}]}, header_errors, header_warnings. Tidak menyimpan file."""
    try:
        wb = load_workbook(BytesIO(raw), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001
        raise HTTPException(http_status.HTTP_422_UNPROCESSABLE_ENTITY,
                            "File tidak dapat dibaca sebagai Excel (.xlsx).") from None
    errors: List[str] = []
    warnings: List[str] = []
    out: Dict[str, List[Dict[str, Any]]] = {}
    names = {n.strip().upper(): n for n in wb.sheetnames}
    try:
        for spec in DATA_SHEETS:
            real = names.get(spec.name)
            if not real:
                if spec.required:
                    errors.append(f"Sheet wajib '{spec.name}' tidak ditemukan.")
                out[spec.name] = []
                continue
            ws = wb[real]
            rows = ws.iter_rows(min_row=spec.header_row, values_only=True)
            header = next(rows, None) or ()
            index: Dict[str, int] = {norm_header(h): i for i, h in enumerate(header) if norm_header(h)}
            mapping: Dict[str, int] = {}
            known = set()
            for col in spec.columns:
                pos = next((index[norm_header(h)] for h in col.headers if norm_header(h) in index), None)
                known.update(norm_header(h) for h in col.headers)
                if pos is None:
                    if col.required or col.required_new:
                        errors.append(f"Sheet '{spec.name}': kolom wajib '{col.headers[0]}' tidak ditemukan di baris header.")
                    continue
                mapping[col.key] = pos
            unknown = [h for h in index if h not in known]
            if unknown:
                warnings.append(f"Sheet '{spec.name}': {len(unknown)} kolom tidak dikenal diabaikan ({', '.join(unknown[:5])}).")
            data: List[Dict[str, Any]] = []
            hint_no = spec.columns[0].hint if spec.columns else None
            for offset, raw_row in enumerate(rows, start=spec.header_row + 1):
                values = {k: (raw_row[i] if i < len(raw_row) else None) for k, i in mapping.items()}
                if all(_clean(v) is None for v in values.values()):
                    continue
                # baris petunjuk (sebelum data_start_row) dilewati HANYA bila memang berisi petunjuk/kosong kunci
                if offset < spec.data_start_row and _clean(values.get("employee_number")) in (None, hint_no):
                    continue
                data.append({"row": offset, "values": values})
            out[spec.name] = data
    finally:
        wb.close()
    return out, errors, warnings


# ============================================================ context (read-only)
async def load_context(company_id: str) -> Dict[str, Any]:
    db = get_tenant_db(company_id)
    masters: Dict[str, List[Dict[str, Any]]] = {}
    lookups: Dict[str, Dict[str, Dict[str, Any]]] = {}
    by_id: Dict[str, Dict[str, Dict[str, Any]]] = {}
    for key, (table, _label) in MASTER_SOURCES.items():
        flt: Dict[str, Any] = {"company_id": company_id, "status": {"$ne": "deleted"}}
        rows = await db[table].find(flt, NO_ID).to_list(20000)
        masters[key] = sorted(rows, key=lambda r: str(r.get("name") or ""))
        by_id[key] = {r["id"]: r for r in rows}
        table_map: Dict[str, Dict[str, Any]] = {}
        for r in rows:
            for cand in (r.get("code"), r.get("name")):
                if cand:
                    table_map.setdefault(str(cand).strip().lower(), r)
        lookups[key] = table_map
    employees = await db.employees.find({"company_id": company_id}, NO_ID).to_list(200000)
    emp_by_key: Dict[str, Dict[str, Any]] = {}
    nik_owner: Dict[str, Dict[str, Any]] = {}
    for e in employees:
        k = employee_key(e.get("employee_number"))
        if k and (k not in emp_by_key or emp_by_key[k].get("status") == "deleted"):
            emp_by_key[k] = e
        if e.get("nik") and e.get("status") != "deleted":
            nik_owner[str(e["nik"]).strip()] = e
    company = await db.companies.find_one({"id": company_id}, NO_ID) or {}
    return {"company_id": company_id, "masters": masters, "lookups": lookups, "by_id": by_id,
            "emp_by_key": emp_by_key, "nik_owner": nik_owner, "company": company}


async def _load_related(company_id: str, emp_ids: List[str]) -> Tuple[Dict[str, Dict[str, Any]], Dict[str, List[Dict[str, Any]]]]:
    db = get_tenant_db(company_id)
    active: Dict[str, Dict[str, Any]] = {}
    family: Dict[str, List[Dict[str, Any]]] = {}
    for i in range(0, len(emp_ids), 500):
        chunk = emp_ids[i:i + 500]
        for a in await db[asg.TABLE].find({"company_id": company_id, "employee_id": {"$in": chunk},
                                           "assignment_status": asg.ACTIVE}, NO_ID).to_list(100000):
            active[a["employee_id"]] = a
        for f in await db[FAMILY_TABLE].find({"company_id": company_id, "employee_id": {"$in": chunk},
                                              "status": {"$ne": "deleted"}}, NO_ID).to_list(100000):
            family.setdefault(f["employee_id"], []).append(f)
    return active, family


# ============================================================ normalisasi sel
def _norm_text(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().lower()


def normalize_cell(col: Col, raw: Any, ctx: Dict[str, Any], today: str) -> Tuple[Any, Optional[str], Optional[str]]:
    """-> (nilai_normal, error, warning). Pesan untuk kolom sensitif TIDAK memuat nilainya."""
    text = _clean(raw)
    if text is None:
        return None, None, None
    shown = "" if col.sensitive else f" '{text[:60]}'"
    t = col.type
    if t in ("text", "name") and col.key not in LONG_TEXT_FIELDS and len(text) > 255:
        return None, f"{col.title} maksimal 255 karakter", None
    if t == "name":
        return (text, None, None) if len(text) >= 2 else (None, f"{col.title} minimal 2 karakter", None)
    if t in ("date", "date_past"):
        val, err = _parse_date(raw)
        if err:
            return None, f"{col.title}: format tanggal{shown} tidak dikenali (gunakan YYYY-MM-DD)", None
        if t == "date_past" and val > today:
            return None, f"{col.title} tidak boleh melebihi hari ini", None
        return val, None, None
    if t == "digits":
        digits = re.sub(r"[\s.\-/]", "", text)
        lo, hi = DIGIT_RULES.get(col.validator or "", (1, 64))
        if not digits.isdigit() or not (lo <= len(digits) <= hi):
            need = f"{lo} digit" if lo == hi else f"{lo}-{hi} digit"
            return None, f"{col.title} tidak valid (harus angka {need})", None
        return digits, None, None
    if t == "email":
        if not re.match(r"^[^@\s]+@[^@\s]+\.[A-Za-z]{2,}$", text):
            return None, f"Email{shown} tidak valid", None
        return text.lower(), None, None
    if t == "phone":
        cleaned = re.sub(r"[^\d+]", "", text)
        if len(re.sub(r"\D", "", cleaned)) < 6:
            return None, f"{col.title}{shown} tidak valid", None
        return cleaned, None, None
    if t == "enum":
        table = ENUMS.get(col.validator or "", {})
        val = table.get(_norm_text(text))
        if val:
            return val, None, None
        if col.validator in SOFT_ENUMS:
            return text, None, f"{col.title}{shown} di luar daftar standar, disimpan apa adanya"
        return None, f"{col.title}{shown} tidak dikenal", None
    if t == "bool":
        low = _norm_text(text)
        if low in BOOL_TRUE:
            return True, None, None
        if low in BOOL_FALSE:
            return False, None, None
        return None, f"{col.title}{shown} harus Ya/Tidak", None
    if t == "master":
        row = ctx["lookups"].get(col.master or "", {}).get(_norm_text(text))
        label = MASTER_SOURCES.get(col.master or "", ("", col.title))[1]
        if not row:
            return None, f"{label}{shown} tidak ditemukan di master perusahaan aktif", None
        if col.master == "employee_business_statuses" and not row.get("is_active"):
            return None, f"{label}{shown} sedang nonaktif", None
        return row["id"], None, None
    return text, None, None


def _mask(field: str, value: Any, sensitive: bool) -> Any:
    return mask_value(value) if sensitive and value not in (None, "") else value


# ============================================================ analyze + validate
def _err(level: str, message: str, sheet: Optional[str] = None, row: Optional[int] = None,
         field: Optional[str] = None) -> Dict[str, Any]:
    return {"level": level, "message": message, "sheet": sheet, "row": row, "field": field}


async def analyze(company_id: str, parsed: Dict[str, List[Dict[str, Any]]], ctx: Optional[Dict[str, Any]] = None
                  ) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """Analisis READ-ONLY. Return (entities, ctx). Entity memuat plan (nilai asli, TIDAK disimpan ke DB)
    dan `changes` versi masking (disimpan untuk preview)."""
    ctx = ctx or await load_context(company_id)
    today = emp_status.today_local()
    entities: Dict[str, Dict[str, Any]] = {}
    order: List[str] = []

    def entity(key: str, raw_number: str) -> Dict[str, Any]:
        if key not in entities:
            entities[key] = {"key": key, "employee_number": raw_number, "rows": {}, "single": {}, "family_rows": [],
                             "errors": [], "warnings": [], "cell_fields": {}}
            order.append(key)
        return entities[key]

    col_by_sheet = {s.name: s for s in DATA_SHEETS}
    for spec in DATA_SHEETS:
        seen_rows: Dict[str, int] = {}
        for item in parsed.get(spec.name, []):
            row, values = item["row"], item["values"]
            number = _clean(values.get("employee_number"))
            if not number:
                ent = entity(f"#{spec.name}!{row}", f"(kosong) {spec.name} baris {row}")
                ent["rows"].setdefault(spec.name, []).append(row)
                ent["errors"].append(_err(ERROR, "Nomor Karyawan wajib diisi", spec.name, row))
                continue
            key = employee_key(number)
            ent = entity(key, number)
            ent["rows"].setdefault(spec.name, []).append(row)
            if spec.kind == "single" and key in seen_rows:
                ent["errors"].append(_err(ERROR, f"Nomor Karyawan duplikat di sheet {spec.name} (baris {seen_rows[key]} dan {row})",
                                          spec.name, row))
                continue
            seen_rows[key] = row
            normalized: Dict[str, Any] = {}
            row_errors = 0
            for col in spec.columns:
                if col.key == "employee_number":
                    continue
                val, err, warn = normalize_cell(col, values.get(col.key), ctx, today)
                if err:
                    row_errors += 1
                    ent["errors"].append(_err(ERROR, err, spec.name, row, col.title))
                elif col.required and val is None:
                    row_errors += 1
                    ent["errors"].append(_err(ERROR, f"{col.title} wajib diisi", spec.name, row, col.title))
                if warn:
                    ent["warnings"].append(_err("WARNING", warn, spec.name, row, col.title))
                if val is not None:
                    normalized[col.key] = val
            if spec.kind == "single":
                ent["single"][spec.name] = normalized
            else:
                ent["family_rows"].append({"row": row, "values": normalized, "valid": row_errors == 0})

    emp_ids = [ctx["emp_by_key"][k]["id"] for k in order if k in ctx["emp_by_key"]]
    active_map, family_map = await _load_related(company_id, emp_ids)

    # NIK duplikat antar karyawan dalam file
    nik_in_file: Dict[str, List[str]] = {}
    for k in order:
        nik = entities[k]["single"].get("DATA_KARYAWAN", {}).get("nik")
        if nik:
            nik_in_file.setdefault(nik, []).append(k)

    results = []
    for seq, k in enumerate(order, start=1):
        ent = entities[k]
        _plan_entity(ent, ctx, active_map, family_map, nik_in_file, today, col_by_sheet)
        ent["seq"] = seq
        results.append(ent)
    return results, ctx


def _label_of(ctx: Dict[str, Any], col: Col, value: Any) -> Any:
    if col.type == "master" and value:
        row = ctx["by_id"].get(col.master or "", {}).get(value)
        return (row or {}).get("name") or (row or {}).get("code") or value
    return value


def _same(a: Any, b: Any) -> bool:
    return _norm_text(a) == _norm_text(b)


def _plan_entity(ent: Dict[str, Any], ctx: Dict[str, Any], active_map, family_map, nik_in_file, today, col_by_sheet) -> None:
    key = ent["key"]
    emp = ctx["emp_by_key"].get(key)
    errors, warnings = ent["errors"], ent["warnings"]
    data_row = ent["single"].get("DATA_KARYAWAN")
    kep = ent["single"].get("KEPEGAWAIAN", {})
    bank = ent["single"].get("BANK_BPJS", {})
    ent["employee_id"] = emp["id"] if emp else None
    ent["full_name"] = (data_row or {}).get("full_name") or (emp or {}).get("full_name")
    ent["emp_updated_at"] = str(emp.get("updated_at")) if emp else None
    changes: List[Dict[str, Any]] = []
    plan: Dict[str, Any] = {"fields": {}, "placement": None, "mirror": {}, "status": None, "family_insert": [],
                            "family_update": []}
    ent["plan"] = plan
    ent["changes"] = changes

    if key.startswith("#"):
        ent["row_class"] = ERROR
        ent["active_assignment_id"] = None
        ent["change_signature"] = hashlib.sha256(key.encode()).hexdigest()
        return
    # ---- identitas & konflik
    if emp and emp.get("status") == "deleted":
        errors.append(_err(CONFLICT, "Nomor karyawan ini milik data karyawan yang sudah dihapus. Gunakan nomor lain / tangani manual."))
    elif emp and emp.get("status") == "archived":
        errors.append(_err(CONFLICT, "Karyawan sedang diarsipkan. Pulihkan dari arsip terlebih dahulu sebelum diimpor."))
    if not emp:
        if data_row is None:
            errors.append(_err(ERROR, "Nomor karyawan tidak ada di sheet DATA_KARYAWAN maupun di data karyawan existing (orphan)."))
        elif not data_row.get("full_name") and not any(e["message"].startswith("Nama") for e in errors):
            errors.append(_err(ERROR, "Nama Lengkap wajib diisi untuk karyawan baru", "DATA_KARYAWAN"))
    company_val = (data_row or {}).get("company")
    if company_val:
        comp = ctx["company"]
        if _norm_text(company_val) not in {_norm_text(comp.get("code")), _norm_text(comp.get("name"))}:
            errors.append(_err(ERROR, f"Perusahaan '{company_val[:60]}' tidak sesuai dengan perusahaan aktif", "DATA_KARYAWAN"))
    nik = (data_row or {}).get("nik")
    if nik:
        owner = ctx["nik_owner"].get(nik)
        if owner and (not emp or owner["id"] != emp["id"]):
            errors.append(_err(CONFLICT, f"NIK KTP ({mask_value(nik)}) sudah dipakai karyawan lain "
                                         f"(Nomor {owner.get('employee_number')}). Tidak digabung otomatis.", "DATA_KARYAWAN"))
        if len(nik_in_file.get(nik, [])) > 1:
            others = [x for x in nik_in_file[nik] if x != key]
            errors.append(_err(CONFLICT, f"NIK KTP ({mask_value(nik)}) juga dipakai nomor karyawan lain di file ini "
                                         f"({', '.join(others[:3])}).", "DATA_KARYAWAN"))
        if emp and emp.get("nik") and str(emp["nik"]).strip() != nik:
            errors.append(_err(CONFLICT, "NIK KTP di Excel berbeda dengan NIK tersimpan. Perubahan identitas harus lewat profil karyawan.",
                               "DATA_KARYAWAN"))

    # ---- field karyawan (blank-preserve: hanya nilai terisi yang dibandingkan)
    incoming: Dict[str, Tuple[Any, Col, str]] = {}
    for sheet_name, values in (("DATA_KARYAWAN", data_row or {}), ("KEPEGAWAIAN", kep), ("BANK_BPJS", bank)):
        spec = col_by_sheet[sheet_name]
        for col in spec.columns:
            if col.target in ("employee", "placement") and col.key in values:
                incoming[col.key] = (values[col.key], col, sheet_name)
    meta_start = kep.get("placement_start_date")

    # penempatan: project berubah -> Pindah (01D). Lokasi dari project tetap bila sel lokasi kosong.
    new_project = incoming.get("project_id", (None,))[0]
    projects = ctx["by_id"]["projects"]
    proj_row = projects.get(new_project) if new_project else None
    cur_project = (emp or {}).get("project_id")
    project_changes = bool(new_project and new_project != cur_project)
    if project_changes and proj_row and proj_row.get("work_location_id") and "work_location_id" not in incoming:
        wl_col = next(c for c in col_by_sheet["KEPEGAWAIAN"].columns if c.key == "work_location_id")
        if (emp or {}).get("work_location_id") != proj_row["work_location_id"]:
            incoming["work_location_id"] = (proj_row["work_location_id"], wl_col, "KEPEGAWAIAN")
            warnings.append(_err("WARNING", "Lokasi kerja diisi otomatis dari lokasi tetap project baru", "KEPEGAWAIAN"))

    merged = {f: (incoming[f][0] if f in incoming else (emp or {}).get(f)) for f in asg.PLACEMENT_FIELDS}
    if proj_row and proj_row.get("work_location_id") and merged.get("work_location_id") \
            and merged["work_location_id"] != proj_row["work_location_id"]:
        errors.append(_err(ERROR, "Lokasi kerja / site tidak sesuai dengan lokasi yang ditetapkan pada project", "KEPEGAWAIAN"))
    div = ctx["by_id"]["divisions"].get(merged.get("division_id") or "")
    if div and div.get("department_id") and merged.get("department_id") and div["department_id"] != merged["department_id"]:
        errors.append(_err(ERROR, "Divisi bukan bagian dari departemen yang dipilih", "KEPEGAWAIAN"))

    for field, (val, col, sheet_name) in incoming.items():
        old = (emp or {}).get(field)
        if emp and _same(old, val):
            continue
        plan["fields"][field] = val
        changes.append({"field": field, "label": col.title, "sheet": sheet_name,
                        "old": _mask(field, _label_of(ctx, col, old), col.sensitive) if emp else None,
                        "new": _mask(field, _label_of(ctx, col, val), col.sensitive)})

    active = active_map.get(emp["id"]) if emp else None
    ent["active_assignment_id"] = (active or {}).get("id")
    if emp:
        plan["fields"].pop("project_id", None)
        if project_changes:
            if active:
                if not meta_start:
                    errors.append(_err(ERROR, "Tanggal Mulai Penempatan wajib diisi karena project berubah (Pindah Penempatan)",
                                       "KEPEGAWAIAN"))
                elif active.get("start_date") and meta_start < active["start_date"]:
                    errors.append(_err(ERROR, f"Tanggal Mulai Penempatan tidak boleh sebelum mulai penempatan saat ini "
                                              f"({active['start_date']})", "KEPEGAWAIAN"))
                plan["placement"] = {"action": "transfer", "start": meta_start, "project_id": new_project}
            else:
                plan["placement"] = {"action": "assign", "start": meta_start, "project_id": new_project}
                if not meta_start:
                    warnings.append(_err("WARNING", "Tanggal mulai penempatan tidak diketahui; disimpan kosong (tidak dikarang)",
                                         "KEPEGAWAIAN"))
        elif active:
            plan["mirror"] = {f: plan["fields"][f] for f in asg.MIRROR_FIELDS if f in plan["fields"]}
    elif new_project:
        start = meta_start
        if not start and kep.get("join_date") and kep["join_date"] <= today:
            start = kep["join_date"]
            warnings.append(_err("WARNING", "Tanggal mulai penempatan diambil dari Tanggal Masuk (aturan 01D)", "KEPEGAWAIAN"))
        elif not start:
            warnings.append(_err("WARNING", "Tanggal mulai penempatan tidak diketahui; disimpan kosong (tidak dikarang)",
                                 "KEPEGAWAIAN"))
        plan["placement"] = {"action": "assign", "start": start, "project_id": new_project}

    # ---- Status Karyawan (01B) - terpisah dari Status Kepegawaian
    st_id = kep.get("employee_status")
    statuses = ctx["by_id"]["employee_business_statuses"]
    if emp:
        if st_id and st_id != emp.get("current_employee_status_id"):
            eff = kep.get("status_effective_date")
            if not eff:
                errors.append(_err(ERROR, "Tanggal Efektif Status wajib diisi karena Status Karyawan berubah", "KEPEGAWAIAN"))
            plan["status"] = {"status_id": st_id, "effective_date": eff,
                              "reason": (kep.get("status_reason") or "Perubahan status dari impor Excel")[:255]}
            prev = statuses.get(emp.get("current_employee_status_id") or "")
            changes.append({"field": "current_employee_status_id", "label": "Status Karyawan", "sheet": "KEPEGAWAIAN",
                            "old": (prev or {}).get("name"), "new": statuses[st_id].get("name")})
    else:
        plan["status"] = {"status_id": st_id, "effective_date": kep.get("status_effective_date") or today,
                          "reason": (kep.get("status_reason") or "Status awal dari impor Excel")[:255]}
        if st_id:
            changes.append({"field": "current_employee_status_id", "label": "Status Karyawan", "sheet": "KEPEGAWAIAN",
                            "old": None, "new": statuses[st_id].get("name")})
    if plan["placement"]:
        pc = next(c for c in col_by_sheet["KEPEGAWAIAN"].columns if c.key == "project_id")
        changes.append({"field": "assignment", "label": "Penempatan (01D)", "sheet": "KEPEGAWAIAN",
                        "old": _label_of(ctx, pc, cur_project) if emp else None,
                        "new": f"{_label_of(ctx, pc, new_project)} ({plan['placement']['action']}, mulai "
                               f"{plan['placement']['start'] or 'tidak diketahui'})"})

    # ---- keluarga (matching aman, rerun tidak menggandakan)
    _plan_family(ent, emp, family_map.get(emp["id"], []) if emp else [], changes, plan)

    # ---- klasifikasi
    levels = {e["level"] for e in errors}
    if ERROR in levels:
        ent["row_class"] = ERROR
    elif CONFLICT in levels:
        ent["row_class"] = CONFLICT
    elif not emp:
        ent["row_class"] = NEW
    elif changes:
        ent["row_class"] = UPDATE
    else:
        ent["row_class"] = UNCHANGED
    sig_src = json.dumps({"c": ent["row_class"], "ch": changes, "u": ent["emp_updated_at"],
                          "a": ent["active_assignment_id"], "e": [e["message"] for e in errors]},
                         sort_keys=True, default=str)
    ent["change_signature"] = hashlib.sha256(sig_src.encode()).hexdigest()


FAMILY_COMPARE = ("relationship", "full_name", "nik", "birth_place", "birth_date", "gender", "occupation", "phone",
                  "is_emergency_contact")


def _plan_family(ent, emp, existing: List[Dict[str, Any]], changes, plan) -> None:
    warnings = ent["warnings"]
    seen = set()
    claimed = set()
    for fr in ent["family_rows"]:
        if not fr["valid"]:
            continue
        v, row = fr["values"], fr["row"]
        name_n = _norm_text(v.get("full_name"))
        dedupe = ("nik", v["nik"]) if v.get("nik") else ("id", v.get("relationship"), name_n, v.get("birth_date"))
        if dedupe in seen:
            warnings.append(_err("WARNING", "Baris keluarga duplikat di file, diabaikan", "KELUARGA", row))
            continue
        seen.add(dedupe)
        match = None
        if v.get("nik"):
            match = next((f for f in existing if f.get("nik") and str(f["nik"]) == v["nik"]), None)
        if not match:
            cands = [f for f in existing if f.get("relationship") == v.get("relationship")
                     and _norm_text(f.get("full_name")) == name_n
                     and (not v.get("birth_date") or not f.get("birth_date") or f.get("birth_date") == v.get("birth_date"))]
            if len(cands) > 1:
                warnings.append(_err("WARNING", "Identitas anggota keluarga ambigu (lebih dari satu data cocok); perlu review manual, "
                                                "tidak diubah", "KELUARGA", row))
                continue
            if cands:
                c = cands[0]
                if v.get("nik") and c.get("nik") and str(c["nik"]) != v["nik"]:
                    warnings.append(_err("WARNING", "Nama & hubungan cocok tetapi NIK berbeda; perlu review manual, tidak diubah",
                                         "KELUARGA", row))
                    continue
                match = c
            elif any(_norm_text(f.get("full_name")) == name_n for f in existing):
                warnings.append(_err("WARNING", "Nama sama dengan anggota keluarga lain (hubungan/tanggal lahir berbeda); "
                                                "perlu review manual, tidak ditambahkan", "KELUARGA", row))
                continue
        label = f"Keluarga: {v.get('full_name')} ({v.get('relationship')})"
        if match:
            if match["id"] in claimed:
                warnings.append(_err("WARNING", "Beberapa baris cocok ke anggota keluarga yang sama; baris ini diabaikan",
                                     "KELUARGA", row))
                continue
            claimed.add(match["id"])
            patch = {f: v[f] for f in FAMILY_COMPARE if f in v and not _same(match.get(f), v[f])}
            if patch:
                plan["family_update"].append({"id": match["id"], "patch": patch, "name": v.get("full_name")})
                changes.append({"field": "family", "label": label, "sheet": "KELUARGA", "old": "ada",
                                "new": "diperbarui: " + ", ".join(sorted(patch))})
        else:
            plan["family_insert"].append(dict(v))
            changes.append({"field": "family", "label": label, "sheet": "KELUARGA", "old": None,
                            "new": "ditambahkan" + (f" (NIK {mask_value(v['nik'])})" if v.get("nik") else "")})


def summarize(entities: List[Dict[str, Any]]) -> Dict[str, int]:
    out = {c: 0 for c in CLASSES}
    for e in entities:
        out[e["row_class"]] += 1
    return out


def row_doc(company_id: str, batch_id: str, ent: Dict[str, Any], user_id: Optional[str]) -> Dict[str, Any]:
    """Baris hasil analisis untuk disimpan: TANPA nilai sensitif penuh (changes sudah dimasking)."""
    commit_status = "PENDING" if ent["row_class"] in (NEW, UPDATE) else "SKIPPED"
    return {"id": new_id(), "company_id": company_id, "status": "active", "batch_id": batch_id,
            "employee_number": ent["employee_number"][:64], "employee_id": ent.get("employee_id"),
            "full_name": (ent.get("full_name") or "")[:255] or None, "row_class": ent["row_class"],
            "commit_status": commit_status, "source_rows": ent["rows"], "changes": ent["changes"],
            "errors": ent["errors"], "warnings": ent["warnings"], "change_signature": ent["change_signature"],
            "seq": ent["seq"], **audit_fields(user_id)}


def file_hash(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


# ============================================================ commit (per karyawan, 1 transaksi)
def _legacy_sync(doc: Optional[Dict[str, Any]]) -> Dict[str, Any]:
    """Sama dengan bridge 01D: employees.* mengikuti assignment aktif."""
    if not doc:
        return {"project_id": None}
    return {f: doc.get(f) for f in asg.PLACEMENT_FIELDS}


def _transfer_end_date(old_start: Optional[str], new_start: str) -> str:
    """Aturan 01D: assignment lama berakhir sehari sebelum mulai baru (atau = mulai baru bila jatuh sebelum mulai lama)."""
    prev_day = (date.fromisoformat(new_start) - timedelta(days=1)).isoformat()
    return prev_day if (not old_start or prev_day >= old_start) else new_start


async def apply_entity(auth, ent: Dict[str, Any], ctx: Dict[str, Any], batch: Dict[str, Any]) -> Dict[str, Any]:
    cid, uid = auth.company_id, auth.user_id
    actor = auth.user.get("full_name") if auth.user else None
    bn = batch.get("batch_number")
    note = f"Impor Excel batch {bn}"
    plan = ent["plan"]
    statuses = ctx["by_id"]["employee_business_statuses"]
    audits: List[Dict[str, Any]] = []
    ts = now()
    async with transaction() as tx:
        if ent["row_class"] == NEW:
            number = ent["employee_number"]
            if await tx.select_one_for_update("employees", {"company_id": cid, "employee_number": number}):
                raise ImportRowError("Nomor karyawan sudah dipakai (dibuat proses lain setelah analisis).")
            nik = plan["fields"].get("nik")
            if nik and await tx.select_one_for_update("employees", {"company_id": cid, "nik": nik, "status": {"$ne": "deleted"}}):
                raise ImportRowError("NIK KTP sudah dipakai karyawan lain (berubah setelah analisis).")
            st_plan = plan["status"] or {}
            st = statuses.get(st_plan.get("status_id") or "") or await emp_status.default_status(cid, "ACTIVE")
            legacy = emp_status.legacy_for(st["system_category"]) if st else "active"
            doc = {"id": new_id(), "company_id": cid, "status": legacy, "employee_number": number,
                   **plan["fields"], "current_employee_status_id": (st or {}).get("id"), **audit_fields(uid)}
            doc.pop("employee_status", None)
            await tx.insert("employees", doc)
            if st:
                await tx.insert(emp_status.HISTORY, emp_status.history_doc(
                    cid, doc["id"], st, None, effective_date=st_plan.get("effective_date"), reason=st_plan.get("reason"),
                    notes=note, source="IMPORT", actor_id=uid, actor_name=actor, legacy_before=None, legacy_after=legacy))
            audits.append(build_audit_entry(auth, "import_create", "employee", doc["id"], doc.get("full_name"),
                                            after={k: v for k, v in doc.items() if k in plan["fields"]}, notes=note))
            if plan["placement"]:
                adoc = asg.new_assignment_doc(cid, doc["id"], doc, plan["placement"]["start"], "IMPORT",
                                              "Penempatan awal dari impor Excel", note, uid)
                await tx.insert(asg.TABLE, adoc)
                audits.append(build_audit_entry(auth, "assignment_create", "employee", doc["id"], doc.get("full_name"),
                                                after={"assignment_id": adoc["id"], "project_id": adoc.get("project_id"),
                                                       "mulai": adoc.get("start_date")}, notes=f"Penempatan awal via {note}"))
            employee_id, name = doc["id"], doc.get("full_name")
        else:
            employee_id = ent["employee_id"]
            locked = await tx.select_one_for_update("employees", {"company_id": cid, "id": employee_id})
            if not locked or str(locked.get("updated_at")) != ent["emp_updated_at"]:
                raise ImportRowError("Data karyawan berubah setelah analisis. Unggah & analisis ulang.")
            if locked.get("status") in emp_status.TECHNICAL_LEGACY:
                raise ImportRowError("Karyawan diarsipkan/dihapus; tidak diubah.")
            name = locked.get("full_name")
            patch = dict(plan["fields"])
            if patch:
                audits.append(build_audit_entry(auth, "import_update", "employee", employee_id, name,
                                                before={k: locked.get(k) for k in patch}, after=patch, notes=note))
            pl = plan["placement"]
            if pl:
                current = await tx.select_one_for_update(asg.TABLE, {"company_id": cid, "employee_id": employee_id,
                                                                    "assignment_status": asg.ACTIVE})
                if (current or {}).get("id") != ent.get("active_assignment_id"):
                    raise ImportRowError("Penempatan karyawan berubah setelah analisis. Analisis ulang.")
                data = {f: (patch.get(f) if patch.get(f) else locked.get(f)) for f in asg.PLACEMENT_FIELDS}
                data["project_id"] = ent_project = pl["project_id"]
                prev_id = None
                if current:
                    end_date = _transfer_end_date(current.get("start_date"), pl["start"])
                    await tx.update(asg.TABLE, {"company_id": cid, "id": current["id"]},
                                    {"assignment_status": asg.ENDED, "end_date": end_date,
                                     "end_reason": f"Pindah penempatan: {note}", "ended_by": uid, "updated_at": ts,
                                     "updated_by": uid})
                    prev_id = current["id"]
                adoc = asg.new_assignment_doc(cid, employee_id, data, pl["start"], "IMPORT",
                                              "Pindah penempatan dari impor Excel" if current else "Penempatan dari impor Excel",
                                              note, uid, previous_assignment_id=prev_id)
                await tx.insert(asg.TABLE, adoc)
                patch.update(_legacy_sync(adoc))
                audits.append(build_audit_entry(
                    auth, "assignment_transfer" if current else "assignment_create", "employee", employee_id, name,
                    before={"assignment_id": prev_id, "project_id": (current or {}).get("project_id")} if current else None,
                    after={"assignment_id": adoc["id"], "project_id": ent_project, "mulai": adoc.get("start_date")},
                    notes=f"Penempatan via {note}"))
            elif plan["mirror"] and ent.get("active_assignment_id"):
                await tx.update(asg.TABLE, {"company_id": cid, "id": ent["active_assignment_id"]},
                                {**plan["mirror"], "updated_at": ts, "updated_by": uid})
            stp = plan["status"]
            if stp:
                new_st = statuses[stp["status_id"]]
                prev = statuses.get(locked.get("current_employee_status_id") or "")
                legacy_after = emp_status.legacy_for(new_st["system_category"])
                await tx.insert(emp_status.HISTORY, emp_status.history_doc(
                    cid, employee_id, new_st, prev, effective_date=stp["effective_date"], reason=stp["reason"], notes=note,
                    source="IMPORT", actor_id=uid, actor_name=actor, legacy_before=locked.get("status"),
                    legacy_after=legacy_after))
                patch.update({"current_employee_status_id": new_st["id"], "status": legacy_after})
                audits.append(build_audit_entry(
                    auth, "status_change", "employee", employee_id, name,
                    before={"status_karyawan": (prev or {}).get("name"), "status_legacy": locked.get("status")},
                    after={"status_karyawan": new_st.get("name"), "status_legacy": legacy_after,
                           "tanggal_efektif": stp["effective_date"]},
                    notes=f"{(prev or {}).get('name') or '-'} -> {new_st.get('name')} (efektif {stp['effective_date']}) via {note}"))
            if patch:
                await tx.update("employees", {"company_id": cid, "id": employee_id},
                                {**patch, "updated_at": ts, "updated_by": uid})
        for f in plan["family_insert"]:
            fdoc = {"id": new_id(), "company_id": cid, "status": "active", "employee_id": employee_id,
                    "is_emergency_contact": bool(f.get("is_emergency_contact")),
                    **{k: v for k, v in f.items() if k != "is_emergency_contact"}, **audit_fields(uid)}
            await tx.insert(FAMILY_TABLE, fdoc)
            audits.append(build_audit_entry(auth, "family_create", "employee", employee_id, name,
                                            after={k: fdoc.get(k) for k in FAMILY_COMPARE}, notes=note))
        for fu in plan["family_update"]:
            await tx.update(FAMILY_TABLE, {"company_id": cid, "id": fu["id"], "employee_id": employee_id},
                            {**fu["patch"], "updated_at": ts, "updated_by": uid})
            audits.append(build_audit_entry(auth, "family_update", "employee", employee_id, name,
                                            after=fu["patch"], notes=note))
        for a in audits:
            await tx.insert("audit_logs", a)
    return {"employee_id": employee_id, "full_name": name}


def check_limits(raw: bytes, filename: Optional[str]) -> None:
    if not raw:
        raise HTTPException(http_status.HTTP_422_UNPROCESSABLE_ENTITY, "File yang diunggah kosong.")
    if len(raw) > LIMITS["max_file_mb"] * 1024 * 1024:
        raise HTTPException(http_status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                            f"Ukuran file melebihi {LIMITS['max_file_mb']} MB.")
    if not (filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(http_status.HTTP_422_UNPROCESSABLE_ENTITY, "Format file harus .xlsx.")


def validate_parsed(parsed: Dict[str, List[Dict[str, Any]]], header_errors: List[str]) -> None:
    if header_errors:
        raise HTTPException(http_status.HTTP_422_UNPROCESSABLE_ENTITY,
                            {"message": "Struktur file tidak sesuai template impor.", "header_errors": header_errors})
    total_keys = {employee_key(_clean(r["values"].get("employee_number")))
                  for s in ("DATA_KARYAWAN", "KEPEGAWAIAN", "BANK_BPJS", "KELUARGA") for r in parsed.get(s, [])}
    if not any(parsed.get(s) for s in parsed):
        raise HTTPException(http_status.HTTP_422_UNPROCESSABLE_ENTITY, "Tidak ada baris data yang terbaca (data mulai baris 3).")
    if len(total_keys) > LIMITS["max_employees"]:
        raise HTTPException(http_status.HTTP_422_UNPROCESSABLE_ENTITY,
                            f"Maksimal {LIMITS['max_employees']} karyawan per impor (file berisi {len(total_keys)}).")


def build_validation_result(batch: Dict[str, Any], rows: List[Dict[str, Any]]) -> bytes:
    """Excel hasil validasi/error. Hanya memakai data tersimpan (sudah dimasking) - tanpa nilai sensitif penuh."""
    wb = Workbook()
    ws = wb.active
    ws.title = "HASIL_VALIDASI"
    head = ["Sheet", "Row", "Employee Number", "Field", "Status", "Error/Warning", "Message"]
    ws.append(head)
    for i in range(1, len(head) + 1):
        ws.cell(row=1, column=i).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=1, column=i).fill = HEADER_FILL
    for r in rows:
        items = [(e, e.get("level") or ERROR) for e in (r.get("errors") or [])] + \
                [(w, "WARNING") for w in (r.get("warnings") or [])]
        if r.get("commit_status") == "FAILED" and r.get("commit_error"):
            items.append(({"message": r["commit_error"], "sheet": None, "row": None, "field": None}, "COMMIT_FAILED"))
        for item, level in items:
            src = r.get("source_rows") or {}
            sheet = item.get("sheet") or ", ".join(src.keys())
            row_no = item.get("row") or ", ".join(str(x) for v in src.values() for x in v)
            ws.append([sheet, str(row_no), r.get("employee_number"), item.get("field") or "", r.get("row_class"), level,
                       item.get("message")])
    for letter, width in zip("ABCDEFG", (18, 10, 22, 26, 12, 16, 90)):
        ws.column_dimensions[letter].width = width
    info = wb.create_sheet("RINGKASAN")
    for k in ("batch_number", "file_name", "batch_status", "total_employees", "count_new", "count_update",
              "count_unchanged", "count_conflict", "count_error", "committed_count", "failed_count"):
        info.append([k, str(batch.get(k) if batch.get(k) is not None else "")])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
