"""Helper Excel untuk Time Management: template, pembacaan mentah, dan ekspor.

Alur import selalu: Upload -> Mapping Kolom -> Preview -> Validasi -> Simpan.
Pembacaan file TIDAK pernah langsung menulis ke database.
"""
from __future__ import annotations

from datetime import date, datetime, time
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

from fastapi import HTTPException, status
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="0F5132")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet_bytes(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _style_header(ws, columns: List[str]) -> None:
    for idx, title in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = max(14, min(34, len(title) + 6))
    ws.freeze_panes = "A2"


# --------------------------------------------------------------------------
# Template standar
# --------------------------------------------------------------------------
ATTENDANCE_TEMPLATE_COLUMNS = [
    "Nomor Karyawan", "Nama", "Tanggal", "Jam Masuk", "Jam Pulang", "Kode Shift", "Catatan",
]
SCHEDULE_TEMPLATE_COLUMNS = [
    "Nomor Karyawan", "Nama", "Tanggal", "Kode Shift", "Kode Lokasi Kerja", "OFF", "Catatan",
]


def build_attendance_template(shift_codes: Optional[List[str]] = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Import Absensi"
    _style_header(ws, ATTENDANCE_TEMPLATE_COLUMNS)
    ws.append(["NEP-0001", "Contoh Karyawan", "2026-09-01", "08:05", "17:10", (shift_codes or ["OFFICE"])[0], "contoh baris"])
    info = wb.create_sheet("Petunjuk")
    info["A1"] = "Petunjuk Import Absensi"
    info["A1"].font = Font(bold=True, size=13)
    for i, line in enumerate([
        "1. Kolom boleh berbeda dari template ini; saat mengunggah Anda dapat memetakan kolom sendiri.",
        "2. Tanggal memakai format YYYY-MM-DD (contoh 2026-09-01).",
        "3. Jam memakai format 24 jam HH:MM (contoh 08:05, 23:00).",
        "4. Kode Shift bersifat opsional; bila kosong sistem memakai jadwal karyawan.",
        "5. Untuk shift malam, Jam Pulang boleh lebih kecil dari Jam Masuk (lintas hari).",
        "6. Baris yang sudah ada absensinya akan DILEWATI, kecuali Anda memilih opsi perbarui.",
        "7. Label status dari mesin absensi TIDAK dipakai; status dihitung ulang oleh sistem.",
    ], start=3):
        info[f"A{i}"] = line
    if shift_codes:
        info["A11"] = "Kode Shift tersedia: " + ", ".join(shift_codes)
    return _sheet_bytes(wb)


def build_schedule_template(
    shift_codes: Optional[List[str]] = None, location_codes: Optional[List[str]] = None
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Import Jadwal"
    _style_header(ws, SCHEDULE_TEMPLATE_COLUMNS)
    ws.append([
        "NEP-0001", "Contoh Karyawan", "2026-09-01",
        (shift_codes or ["OFFICE"])[0], (location_codes or ["LOC-HO"])[0], "TIDAK", "contoh baris",
    ])
    info = wb.create_sheet("Petunjuk")
    info["A1"] = "Petunjuk Import Jadwal"
    info["A1"].font = Font(bold=True, size=13)
    for i, line in enumerate([
        "1. Kolom boleh berbeda; pemetaan kolom dilakukan saat unggah.",
        "2. Tanggal memakai format YYYY-MM-DD.",
        "3. Kolom OFF diisi YA bila tanggal tersebut hari OFF (Kode Shift boleh kosong).",
        "4. Kode Lokasi Kerja mengikuti Master Lokasi Kerja perusahaan Anda.",
        "5. Jadwal yang sudah ada akan dilewati, kecuali Anda memilih opsi perbarui.",
    ], start=3):
        info[f"A{i}"] = line
    if shift_codes:
        info["A9"] = "Kode Shift tersedia: " + ", ".join(shift_codes)
    if location_codes:
        info["A10"] = "Kode Lokasi Kerja tersedia: " + ", ".join(location_codes)
    return _sheet_bytes(wb)


# --------------------------------------------------------------------------
# Pembacaan mentah (untuk mapping + preview)
# --------------------------------------------------------------------------
def _cell_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, datetime):
        if value.hour or value.minute or value.second:
            return value.strftime("%Y-%m-%d %H:%M")
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def read_rows(file_bytes: bytes, max_rows: int = 5000) -> Tuple[List[str], List[Dict[str, Any]]]:
    """Baca sheet pertama menjadi (daftar header, daftar baris dict)."""
    try:
        wb = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_ENTITY,
            "Berkas tidak dapat dibaca. Pastikan berformat Excel (.xlsx) dan tidak rusak.",
        )
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Berkas kosong.")

    headers: List[str] = []
    for idx, raw in enumerate(header_row or []):
        text = _cell_text(raw) or f"Kolom {idx + 1}"
        while text in headers:
            text = f"{text} ({idx + 1})"
        headers.append(text)

    data: List[Dict[str, Any]] = []
    for row_no, raw_row in enumerate(rows, start=2):
        if raw_row is None:
            continue
        values = {headers[i]: _cell_text(v) for i, v in enumerate(raw_row) if i < len(headers)}
        if not any(v for v in values.values()):
            continue
        values["__row__"] = row_no
        data.append(values)
        if len(data) >= max_rows:
            break
    wb.close()
    if not data:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_ENTITY,
            "Tidak ada baris data yang dapat dibaca setelah baris judul.",
        )
    return headers, data


def guess_mapping(headers: List[str], targets: Dict[str, List[str]]) -> Dict[str, Optional[str]]:
    """Tebak pemetaan kolom otomatis; HR tetap dapat menggantinya."""
    lowered = {h.lower().strip(): h for h in headers}
    mapping: Dict[str, Optional[str]] = {}
    for field, aliases in targets.items():
        found = None
        for alias in aliases:
            a = alias.lower()
            if a in lowered:
                found = lowered[a]
                break
        if not found:
            for alias in aliases:
                a = alias.lower()
                for key, original in lowered.items():
                    if a and (a in key or key in a):
                        found = original
                        break
                if found:
                    break
        mapping[field] = found
    return mapping


ATTENDANCE_FIELD_ALIASES: Dict[str, List[str]] = {
    "employee_number": ["nomor karyawan", "nomor induk", "nik", "pin", "employee number", "employee id", "id karyawan", "no karyawan"],
    "employee_name": ["nama", "nama karyawan", "name", "employee name"],
    "work_date": ["tanggal", "date", "tgl", "tanggal absen"],
    "check_in": ["jam masuk", "masuk", "in", "check in", "check-in", "clock in", "scan masuk"],
    "check_out": ["jam pulang", "pulang", "out", "check out", "check-out", "clock out", "scan pulang"],
    "shift_code": ["kode shift", "shift", "shift code"],
    "note": ["catatan", "keterangan", "note", "remark"],
}

SCHEDULE_FIELD_ALIASES: Dict[str, List[str]] = {
    "employee_number": ["nomor karyawan", "nomor induk", "nik", "pin", "employee number", "id karyawan", "no karyawan"],
    "employee_name": ["nama", "nama karyawan", "name"],
    "work_date": ["tanggal", "date", "tgl"],
    "shift_code": ["kode shift", "shift", "shift code"],
    "location_code": ["kode lokasi kerja", "lokasi kerja", "lokasi", "location", "work location"],
    "is_day_off": ["off", "hari off", "libur", "day off"],
    "note": ["catatan", "keterangan", "note"],
}

ATTENDANCE_REQUIRED_FIELDS = ["employee_number", "work_date"]
SCHEDULE_REQUIRED_FIELDS = ["employee_number", "work_date"]

ATTENDANCE_FIELD_LABELS = {
    "employee_number": "Nomor Karyawan / NIK",
    "employee_name": "Nama",
    "work_date": "Tanggal",
    "check_in": "Jam Masuk",
    "check_out": "Jam Pulang",
    "shift_code": "Kode Shift",
    "note": "Catatan",
}
SCHEDULE_FIELD_LABELS = {
    "employee_number": "Nomor Karyawan / NIK",
    "employee_name": "Nama",
    "work_date": "Tanggal",
    "shift_code": "Kode Shift",
    "location_code": "Kode Lokasi Kerja",
    "is_day_off": "OFF",
    "note": "Catatan",
}


def truthy(value: Any) -> bool:
    return str(value or "").strip().lower() in ("ya", "yes", "y", "1", "true", "off", "x", "v", "benar")


# --------------------------------------------------------------------------
# Ekspor
# --------------------------------------------------------------------------
def export_sheet(title: str, columns: List[Dict[str, str]], rows: List[Dict[str, Any]],
                 meta: Optional[List[Tuple[str, str]]] = None) -> bytes:
    """Bangun berkas Excel dengan judul kolom Bahasa Indonesia."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    _style_header(ws, [c["label"] for c in columns])
    for row in rows:
        ws.append([row.get(c["key"]) for c in columns])
    if meta:
        info = wb.create_sheet("Informasi")
        info["A1"] = "Informasi Laporan"
        info["A1"].font = Font(bold=True, size=13)
        for i, (label, value) in enumerate(meta, start=3):
            info[f"A{i}"] = label
            info[f"B{i}"] = value
        info.column_dimensions["A"].width = 30
        info.column_dimensions["B"].width = 50
    return _sheet_bytes(wb)
