"""Phase 2A CP3 - Impor Master Aset & Impor Saldo Awal (Opening Existing Holding) massal.

Alur wajib: Download Template -> Upload Excel -> Preview & Validasi (per baris VALID / WARNING / ERROR) -> Commit.
- Preview TIDAK menulis data bisnis (hanya batch + baris audit `asset_import_batches` / `asset_import_rows`).
- Commit hanya bila 0 ERROR (tidak ada baris yang dilewati diam-diam); validasi diulang terhadap DB terkini saat
  commit, lalu seluruh baris ditulis dalam SATU transaksi (gagal -> rollback semua).
- Import Master: Asset Code sistem SELALU dari sequence ASSET_CODE (AST-{SEQ:6}, alokasi massal satu kunci counter);
  Kode Aset Lama disimpan terpisah (`legacy_code`, unik bila diisi). Status IN_USE / PENDING_INSPECTION DITOLAK
  (IN_USE wajib punya holding -> gunakan Saldo Awal).
- Impor Saldo Awal: identifikasi aset Asset Code -> Kode Lama -> Serial (harus tepat satu); baris dikelompokkan per
  karyawan -> satu opening / satu BAST-EXS; publish memakai `finalize_openings` (logika yang sama dengan form manual).
- 01I: seluruh lookup karyawan/project/aset di SQL; restricted hanya melihat batch miliknya; UUID lain -> 404 generik.
- Parsing/validasi batch di backend (bukan ribuan request dari browser). File tidak disimpan; audit tanpa isi file.
"""
from __future__ import annotations

import hashlib
import io
from datetime import date
from typing import Any, Dict, List, Optional, Tuple

import sqlalchemy as sa
from fastapi import APIRouter, Body, Depends, File, HTTPException, Query, Response, UploadFile, status
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from ..core import asset_service as svc
from ..core import data_scope as dscope
from ..core import doc_sequence
from ..core.audit import log_action
from ..core.db import NO_ID, get_db, get_table, new_id, now
from ..core.deps import AuthContext, require_permission
from .asset_lifecycle import DRAFT, _atomic, _not_found, _page, _stamp
from .asset_opening import finalize_openings, lock_sequences, prepare_opening

router = APIRouter(prefix="/asset-imports", tags=["Manajemen Aset - Impor"])

MASTER, OPENING = "MASTER", "OPENING"
VALID, WARNING, ERROR = "VALID", "WARNING", "ERROR"
PREVIEWED, COMMITTED, CANCELLED = "PREVIEWED", "COMMITTED", "CANCELLED"
MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10000
CHUNK = 500
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

MASTER_COLS = [("name", "Nama Aset*"), ("category", "Kategori*"), ("unit", "Satuan*"),
               ("legacy_code", "Kode Aset Lama / Nomor Inventaris"), ("serial_number", "Serial Number"), ("brand", "Merk"),
               ("model", "Model / Tipe"), ("acquisition_year", "Tahun Perolehan"), ("acquisition_value", "Nilai Perolehan"),
               ("condition", "Kondisi*"), ("status", "Status Awal*"), ("project", "Proyek"), ("location", "Lokasi"),
               ("notes", "Catatan")]
OPENING_COLS = [("employee_number", "Nomor Karyawan*"), ("asset_code", "Asset Code Sistem"), ("legacy_code", "Kode Aset Lama"),
                ("serial_number", "Serial Number"), ("opening_date", "Tanggal Mulai / Saldo Awal* (YYYY-MM-DD)"),
                ("project", "Proyek"), ("condition", "Kondisi*"), ("accessories", "Kelengkapan"), ("notes", "Catatan")]
COLS = {MASTER: MASTER_COLS, OPENING: OPENING_COLS}


# ------------------------------------------------------------------ util
def _norm_header(v: Any) -> str:
    return " ".join(str(v or "").replace("*", "").split()).strip().lower()


def _cell(v: Any) -> Optional[str]:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if hasattr(v, "isoformat"):
        return v.isoformat()[:10]
    s = " ".join(str(v).split()).strip()
    return s or None


def _key(v: Optional[str]) -> str:
    return " ".join((v or "").split()).strip().upper()


def _msg(level: str, field: str, text: str) -> Dict[str, str]:
    return {"level": level, "field": field, "message": text}


def _classify(msgs: List[Dict[str, str]]) -> str:
    if any(m["level"] == ERROR for m in msgs):
        return ERROR
    return WARNING if msgs else VALID


def _in_chunks(values):
    values = sorted({v for v in values if v})
    for i in range(0, len(values), CHUNK):
        yield values[i:i + CHUNK]


async def _parse(file: UploadFile, itype: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    raw = await file.read()
    if not raw:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "File kosong.")
    if len(raw) > MAX_BYTES:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "Ukuran file maksimal 10 MB.")
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Format file harus .xlsx (gunakan template yang disediakan).")
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "File Excel tidak dapat dibaca. Pastikan memakai template .xlsx.") from None
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = next(it, None) or []
    wanted = {_norm_header(label): key for key, label in COLS[itype]}
    idx = {wanted[_norm_header(h)]: i for i, h in enumerate(header) if _norm_header(h) in wanted}
    required = [key for key, label in COLS[itype] if label.endswith("*") or "*" in label]
    missing = [label for key, label in COLS[itype] if key in required and key not in idx]
    if missing:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            f"Kolom wajib tidak ditemukan: {', '.join(missing)}. Gunakan template {itype.lower()} terbaru.")
    rows = []
    for n, values in enumerate(it, start=2):
        rec = {key: _cell(values[i]) if i < len(values) else None for key, i in idx.items()}
        if not any(rec.values()):
            continue
        rows.append({"row_no": n, "raw": rec})
        if len(rows) > MAX_ROWS:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"Maksimal {MAX_ROWS} baris per file.")
    wb.close()
    if not rows:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Tidak ada baris data pada file.")
    return rows, {"file_name": (file.filename or "")[:255], "file_hash": hashlib.sha256(raw).hexdigest(), "file_size": len(raw)}


class _Ref:
    """Pencocokan referensi master by kode ATAU nama (case-insensitive); nama ganda -> ambigu."""

    def __init__(self, rows: List[Dict[str, Any]], label: str):
        self.label, self.by_code, self.by_name = label, {}, {}
        for r in rows:
            if r.get("code"):
                self.by_code.setdefault(_key(r["code"]), []).append(r)
            self.by_name.setdefault(_key(r.get("name")), []).append(r)

    def match(self, value: Optional[str], field: str, msgs: List[Dict[str, str]], suffix: str = "") -> Optional[Dict[str, Any]]:
        if not value:
            return None
        hits = self.by_code.get(_key(value)) or self.by_name.get(_key(value)) or []
        if len(hits) == 1:
            return hits[0]
        if not hits:
            msgs.append(_msg(ERROR, field, f"{self.label} '{value}' tidak ditemukan{suffix}."))
        else:
            msgs.append(_msg(ERROR, field, f"{self.label} '{value}' ambigu (cocok dengan {len(hits)} data). Gunakan kode."))
        return None


async def _refs(ctx: AuthContext, scope) -> Dict[str, _Ref]:
    db, cid = get_db(), ctx.company_id
    act = {"company_id": cid, "status": "active"}
    pflt = dscope.with_project_scope(dict(act), scope, "id")
    return {"category": _Ref(await db.asset_categories.find(act, NO_ID).to_list(2000), "Kategori"),
            "unit": _Ref(await db.asset_units.find(act, NO_ID).to_list(2000), "Satuan"),
            "condition": _Ref(await db.asset_conditions.find(act, NO_ID).to_list(2000), "Kondisi"),
            "status": _Ref(await db.asset_statuses.find(act, NO_ID).to_list(2000), "Status"),
            "project": _Ref(await db.projects.find(pflt, NO_ID).to_list(5000), "Proyek"),
            "location": _Ref(await db.work_locations.find(act, NO_ID).to_list(5000), "Lokasi")}


async def _existing(cid: str, col: str, values, scope=None, extra: Optional[Dict[str, Any]] = None) -> Dict[str, List[Dict[str, Any]]]:
    out: Dict[str, List[Dict[str, Any]]] = {}
    for chunk in _in_chunks(values):
        flt = {"company_id": cid, col: {"$in": chunk}, "status": {"$ne": "deleted"}, **(extra or {})}
        if scope is not None:
            flt = dscope.with_project_scope(flt, scope)
        for r in await get_db().assets.find(flt, NO_ID).to_list(len(chunk) * 2):
            out.setdefault(r[col], []).append(r)
    return out


# ------------------------------------------------------------------ validasi MASTER
async def validate_master(ctx: AuthContext, rows: List[Dict[str, Any]]) -> None:
    cid, scope = ctx.company_id, await dscope.get_scope(ctx)
    refs = await _refs(ctx, scope)
    can_value = ctx.has_permission("asset_value", "edit")
    this_year = date.today().year
    for r in rows:
        raw, msgs, d = r["raw"], [], {}
        d["name"] = svc.clean_text(raw.get("name"))
        if not d["name"]:
            msgs.append(_msg(ERROR, "name", "Nama Aset wajib diisi."))
        cat = refs["category"].match(raw.get("category"), "category", msgs)
        if not raw.get("category"):
            msgs.append(_msg(ERROR, "category", "Kategori wajib diisi."))
        unit = refs["unit"].match(raw.get("unit"), "unit", msgs)
        if not raw.get("unit"):
            msgs.append(_msg(ERROR, "unit", "Satuan wajib diisi."))
        cond = refs["condition"].match(raw.get("condition"), "condition", msgs)
        if not raw.get("condition"):
            msgs.append(_msg(ERROR, "condition", "Kondisi wajib diisi."))
        st = refs["status"].match(raw.get("status"), "status", msgs)
        if not raw.get("status"):
            msgs.append(_msg(ERROR, "status", "Status Awal wajib diisi."))
        elif st and st.get("system_state") not in svc.MANUAL_STATES:
            label = svc.LIFECYCLE_STATES.get(st.get("system_state"), st.get("name"))
            msgs.append(_msg(ERROR, "status", f"Status '{label}' tidak boleh diimpor lewat master. Impor sebagai Siap Pakai, "
                                              "lalu catat pemegangnya lewat Saldo Awal."))
            st = None
        d["serial_number"] = svc.clean_text(raw.get("serial_number"))
        d["serial_number_norm"] = svc.norm_serial(d["serial_number"])
        if cat and cat.get("serial_number_required") and not d["serial_number_norm"]:
            msgs.append(_msg(ERROR, "serial_number", f"Serial Number wajib diisi untuk kategori {cat.get('name')}."))
        d["legacy_code"] = svc.clean_text(raw.get("legacy_code"), 64)
        d["legacy_code_norm"] = svc.norm_code(d["legacy_code"])
        d["brand"], d["model"] = svc.clean_text(raw.get("brand")), svc.clean_text(raw.get("model"))
        d["notes"] = svc.clean_text(raw.get("notes"), 5000)
        d["acquisition_year"] = None
        if raw.get("acquisition_year"):
            try:
                y = int(str(raw["acquisition_year"]).strip()[:4]) if str(raw["acquisition_year"]).strip()[:4].isdigit() else None
            except ValueError:
                y = None
            if not y or y < 1900 or y > this_year or len(str(raw["acquisition_year"]).strip()) != 4:
                msgs.append(_msg(ERROR, "acquisition_year", f"Tahun Perolehan '{raw['acquisition_year']}' tidak valid (1900-{this_year})."))
            else:
                d["acquisition_year"] = y
        d["acquisition_value"] = None
        if raw.get("acquisition_value"):
            try:
                v = svc.parse_value(raw["acquisition_value"])
                d["acquisition_value"] = str(v) if v is not None else None
            except HTTPException:
                msgs.append(_msg(ERROR, "acquisition_value", "Nilai Perolehan harus berupa angka (contoh: 15000000)."))
            if d["acquisition_value"] and not can_value:
                msgs.append(_msg(ERROR, "acquisition_value", "Anda tidak memiliki izin mengisi Nilai Perolehan. Kosongkan kolom ini."))
        proj = refs["project"].match(raw.get("project"), "project", msgs, " atau di luar cakupan data Anda")
        if not raw.get("project") and scope.restricted:
            msgs.append(_msg(ERROR, "project", "Proyek wajib diisi (akses Anda dibatasi per proyek)."))
        loc = refs["location"].match(raw.get("location"), "location", msgs)
        if not d["serial_number_norm"] and not d["legacy_code_norm"] and not any(m["field"] == "serial_number" for m in msgs):
            msgs.append(_msg(WARNING, "serial_number", "Tanpa Serial Number & Kode Aset Lama: saldo awal massal hanya bisa "
                                                       "mengenali aset ini lewat Asset Code sistem."))
        d.update({"category_id": (cat or {}).get("id"), "unit_id": (unit or {}).get("id"), "condition_id": (cond or {}).get("id"),
                  "status_id": (st or {}).get("id"), "lifecycle_state": (st or {}).get("system_state"),
                  "project_id": (proj or {}).get("id"), "work_location_id": (loc or {}).get("id")})
        r["data"], r["messages"] = d, msgs
    # duplikat dalam file + terhadap database (batch query)
    for field, col, label in (("serial_number_norm", "serial_number_norm", "Serial Number"),
                              ("legacy_code_norm", "legacy_code_norm", "Kode Aset Lama")):
        seen: Dict[str, List[Dict[str, Any]]] = {}
        for r in rows:
            if r["data"].get(field):
                seen.setdefault(r["data"][field], []).append(r)
        for val, group in seen.items():
            if len(group) > 1:
                nos = ", ".join(str(g["row_no"]) for g in group)
                for g in group:
                    g["messages"].append(_msg(ERROR, field.replace("_norm", ""), f"{label} '{val}' duplikat di file (baris {nos})."))
        db_hits = await _existing(cid, col, list(seen))
        for val, hits in db_hits.items():
            for g in seen.get(val, []):
                g["messages"].append(_msg(ERROR, field.replace("_norm", ""), f"{label} sudah digunakan oleh {hits[0].get('asset_code')}."))
    for r in rows:
        r["row_class"] = _classify(r["messages"])


# ------------------------------------------------------------------ validasi OPENING
async def validate_opening_rows(ctx: AuthContext, rows: List[Dict[str, Any]]) -> None:
    cid, scope = ctx.company_id, await dscope.get_scope(ctx)
    refs = await _refs(ctx, scope)
    db = get_db()
    today = date.today().isoformat()
    emps: Dict[str, List[Dict[str, Any]]] = {}
    for chunk in _in_chunks(r["raw"].get("employee_number") for r in rows):
        flt = dscope.with_scope({"company_id": cid, "status": "active", "employee_number": {"$in": chunk}}, scope)
        for e in await db.employees.find(flt, NO_ID).to_list(len(chunk) * 2):
            emps.setdefault(e["employee_number"], []).append(e)
    by_code = await _existing(cid, "asset_code_norm", [svc.norm_code(r["raw"].get("asset_code")) for r in rows], scope)
    by_legacy = await _existing(cid, "legacy_code_norm", [svc.norm_code(r["raw"].get("legacy_code")) for r in rows], scope)
    by_serial = await _existing(cid, "serial_number_norm", [svc.norm_serial(r["raw"].get("serial_number")) for r in rows], scope)
    emp_ids = sorted({e["id"] for v in emps.values() for e in v})
    default_proj: Dict[str, Optional[str]] = {}
    for chunk in _in_chunks(emp_ids):
        for a in await db.employee_assignments.find({"company_id": cid, "employee_id": {"$in": chunk}, "assignment_status": "ACTIVE"},
                                                     NO_ID).to_list(len(chunk) * 3):
            if a.get("project_id") and scope.allows_project(a["project_id"]):
                default_proj.setdefault(a["employee_id"], a["project_id"])
    for r in rows:
        raw, msgs, d = r["raw"], [], {}
        en = raw.get("employee_number")
        emp = None
        if not en:
            msgs.append(_msg(ERROR, "employee_number", "Nomor Karyawan wajib diisi."))
        else:
            hits = emps.get(en, [])
            if len(hits) == 1:
                emp = hits[0]
            elif not hits:
                msgs.append(_msg(ERROR, "employee_number", f"Nomor Karyawan '{en}' tidak ditemukan, tidak aktif, atau di luar cakupan data Anda."))
            else:
                msgs.append(_msg(ERROR, "employee_number", f"Nomor Karyawan '{en}' ambigu (dipakai {len(hits)} karyawan)."))
        asset, via = None, None
        for field, pool, norm in (("asset_code", by_code, svc.norm_code), ("legacy_code", by_legacy, svc.norm_code),
                                  ("serial_number", by_serial, svc.norm_serial)):
            if raw.get(field):
                hits = pool.get(norm(raw[field]), [])
                label = {"asset_code": "Asset Code", "legacy_code": "Kode Aset Lama", "serial_number": "Serial Number"}[field]
                if len(hits) == 1:
                    asset, via = hits[0], field
                elif not hits:
                    msgs.append(_msg(ERROR, field, f"Aset dengan {label} '{raw[field]}' tidak ditemukan atau di luar cakupan data Anda."))
                else:
                    msgs.append(_msg(ERROR, field, f"{label} '{raw[field]}' cocok dengan {len(hits)} aset (ambigu)."))
                break
        else:
            msgs.append(_msg(ERROR, "asset_code", "Isi salah satu: Asset Code Sistem, Kode Aset Lama, atau Serial Number."))
        if asset and asset.get("lifecycle_state") != svc.READY:
            msgs.append(_msg(ERROR, via, f"Aset {asset.get('asset_code')} berstatus {svc.LIFECYCLE_STATES.get(asset.get('lifecycle_state'))}; "
                                         "saldo awal hanya untuk aset Siap Pakai tanpa pemegang aktif."))
        od = raw.get("opening_date")
        d["opening_date"] = None
        if not od:
            msgs.append(_msg(ERROR, "opening_date", "Tanggal Mulai / Saldo Awal wajib diisi."))
        else:
            try:
                d["opening_date"] = date.fromisoformat(str(od)[:10]).isoformat()
                if d["opening_date"] > today:
                    msgs.append(_msg(ERROR, "opening_date", "Tanggal saldo awal tidak boleh di masa depan."))
            except ValueError:
                msgs.append(_msg(ERROR, "opening_date", f"Tanggal '{od}' tidak valid (format YYYY-MM-DD)."))
        proj = refs["project"].match(raw.get("project"), "project", msgs, " atau di luar cakupan data Anda")
        pid = (proj or {}).get("id") or (default_proj.get(emp["id"]) if emp and not raw.get("project") else None)
        if not pid and scope.restricted and not any(m["field"] == "project" for m in msgs):
            msgs.append(_msg(ERROR, "project", "Proyek wajib diisi (akses Anda dibatasi per proyek)."))
        cond = refs["condition"].match(raw.get("condition"), "condition", msgs)
        if not raw.get("condition"):
            msgs.append(_msg(ERROR, "condition", "Kondisi saat saldo awal wajib diisi."))
        d.update({"employee_id": (emp or {}).get("id"), "employee_name": (emp or {}).get("full_name"), "employee_number": en,
                  "asset_id": (asset or {}).get("id"), "asset_code": (asset or {}).get("asset_code"),
                  "asset_name": (asset or {}).get("name"), "project_id": pid, "condition_id": (cond or {}).get("id"),
                  "accessories": svc.clean_text(raw.get("accessories"), 2000), "item_notes": svc.clean_text(raw.get("notes"), 2000)})
        r["data"], r["messages"] = d, msgs
    # aset sudah punya pemegang aktif / ganda dalam file / sudah di draft lain; konsistensi grup per karyawan
    aids = [r["data"]["asset_id"] for r in rows if r["data"].get("asset_id")]
    held, drafts = set(), set()
    for chunk in _in_chunks(aids):
        held |= {h["asset_id"] for h in await db.asset_holdings.find({"company_id": cid, "active_lock": {"$in": chunk}}, NO_ID).to_list(len(chunk))}
        items = await db.asset_opening_items.find({"company_id": cid, "asset_id": {"$in": chunk}}, NO_ID).to_list(len(chunk) * 3)
        if items:
            live = {o["id"] for o in await db.asset_openings.find({"company_id": cid, "doc_state": DRAFT,
                                                                   "id": {"$in": sorted({i["opening_id"] for i in items})}}, NO_ID).to_list(len(items))}
            drafts |= {i["asset_id"] for i in items if i["opening_id"] in live}
    seen: Dict[str, List[Dict[str, Any]]] = {}
    groups: Dict[str, List[Dict[str, Any]]] = {}
    for r in rows:
        d = r["data"]
        if d.get("asset_id"):
            seen.setdefault(d["asset_id"], []).append(r)
            if d["asset_id"] in held:
                r["messages"].append(_msg(ERROR, "asset_code", f"Aset {d['asset_code']} sudah memiliki pemegang aktif."))
            elif d["asset_id"] in drafts:
                r["messages"].append(_msg(WARNING, "asset_code", f"Aset {d['asset_code']} juga ada di draft saldo awal lain."))
        if d.get("employee_id"):
            groups.setdefault(d["employee_id"], []).append(r)
    for group in seen.values():
        if len(group) > 1:
            nos = ", ".join(str(g["row_no"]) for g in group)
            for g in group:
                g["messages"].append(_msg(ERROR, "asset_code", f"Aset {g['data']['asset_code']} muncul lebih dari sekali di file (baris {nos})."))
    for group in groups.values():
        if len({(g["data"].get("opening_date"), g["data"].get("project_id")) for g in group}) > 1:
            for g in group:
                g["messages"].append(_msg(ERROR, "opening_date", "Baris milik karyawan yang sama harus memiliki Tanggal & Proyek "
                                                                 "yang sama (satu karyawan = satu dokumen saldo awal)."))
    for r in rows:
        r["row_class"] = _classify(r["messages"])


VALIDATORS = {MASTER: validate_master, OPENING: validate_opening_rows}


# ------------------------------------------------------------------ batch helpers
def _counts(rows) -> Dict[str, int]:
    return {"total_rows": len(rows), "valid_rows": sum(r["row_class"] == VALID for r in rows),
            "warning_rows": sum(r["row_class"] == WARNING for r in rows), "error_rows": sum(r["row_class"] == ERROR for r in rows)}


def _group_summary(rows) -> Dict[str, Any]:
    """Ringkasan pengelompokan saldo awal per karyawan (1 karyawan = 1 dokumen BAST-EXS) untuk layar preview."""
    groups: Dict[str, Dict[str, Any]] = {}
    for r in rows:
        d = r.get("data") or {}
        if not d.get("employee_id"):
            continue
        g = groups.setdefault(d["employee_id"], {"employee_number": d.get("employee_number"), "employee_name": d.get("employee_name"),
                                                 "opening_date": d.get("opening_date"), "assets": 0, "error_rows": 0})
        g["assets"] += 1
        g["error_rows"] += r.get("row_class") == ERROR
    items = sorted(groups.values(), key=lambda g: g.get("employee_number") or "")
    return {"employee_groups": len(items), "groups": items[:500]}


async def _visible_batch(ctx: AuthContext, batch_id: str) -> Dict[str, Any]:
    b = await get_db().asset_import_batches.find_one({"company_id": ctx.company_id, "id": batch_id}, NO_ID)
    scope = await dscope.get_scope(ctx)
    if not b or (scope.restricted and b.get("uploaded_by") != ctx.user_id):
        raise _not_found()
    return b


async def _batch_out(ctx: AuthContext, b: Dict[str, Any]) -> Dict[str, Any]:
    names = {}
    ids = [i for i in (b.get("uploaded_by"), b.get("committed_by")) if i]
    if ids:
        names = {u["id"]: u.get("full_name") or u.get("email") for u in await get_db().users.find({"id": {"$in": ids}}, NO_ID).to_list(5)}
    return {**{k: b.get(k) for k in ("id", "batch_number", "import_type", "batch_state", "file_name", "total_rows", "valid_rows",
                                    "warning_rows", "error_rows", "created_count", "opening_count", "commit_mode",
                                    "uploaded_at", "committed_at", "summary")},
            "uploaded_by_name": names.get(b.get("uploaded_by")), "committed_by_name": names.get(b.get("committed_by")),
            "can_commit": b.get("batch_state") == PREVIEWED and not b.get("error_rows")}


async def _load_rows(cid: str, batch_id: str) -> List[Dict[str, Any]]:
    return await get_db().asset_import_rows.find({"company_id": cid, "batch_id": batch_id}, NO_ID).sort([("row_no", 1)]).to_list(MAX_ROWS + 10)


# ------------------------------------------------------------------ template
@router.get("/template")
async def download_template(import_type: str = Query(MASTER), ctx: AuthContext = Depends(require_permission("asset_import", "view"))):
    itype = import_type.upper()
    if itype not in COLS:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Jenis template tidak dikenal.")
    scope = await dscope.get_scope(ctx)
    db, cid = get_db(), ctx.company_id
    wb = Workbook()
    ws = wb.active
    ws.title = "Saldo Awal" if itype == OPENING else "Master Aset"
    ws.append([label for _, label in COLS[itype]])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEF3")
        ws.column_dimensions[c.column_letter].width = max(16, len(str(c.value)) + 2)
    ref = wb.create_sheet("Referensi")
    ref.append(["Jenis", "Kode", "Nama", "Keterangan"])
    act = {"company_id": cid, "status": "active"}
    for label, coll, extra in (("Kategori", "asset_categories", lambda r: "Serial wajib" if r.get("serial_number_required") else ""),
                               ("Satuan", "asset_units", lambda r: ""), ("Kondisi", "asset_conditions", lambda r: ""),
                               ("Lokasi", "work_locations", lambda r: "")):
        for r in await db[coll].find(act, NO_ID).sort([("name", 1)]).to_list(2000):
            ref.append([label, r.get("code"), r.get("name"), extra(r)])
    if itype == MASTER:
        for r in await db.asset_statuses.find(act, NO_ID).sort([("sort_order", 1)]).to_list(200):
            ok = r.get("system_state") in svc.MANUAL_STATES
            ref.append(["Status Awal", r.get("code"), r.get("name"), "Boleh" if ok else "TIDAK boleh lewat impor master"])
    for r in await db.projects.find(dscope.with_project_scope(dict(act), scope, "id"), NO_ID).sort([("name", 1)]).to_list(5000):
        ref.append(["Proyek", r.get("code"), r.get("name"), ""])
    guide = wb.create_sheet("Petunjuk")
    notes = ["Kolom bertanda * wajib diisi. Referensi dicocokkan dengan Kode atau Nama pada sheet Referensi.",
             "Asset Code sistem TIDAK diisi di impor master - sistem membuat AST-xxxxxx otomatis.",
             "Aset yang sedang dipegang karyawan: impor master sebagai Siap Pakai, lalu gunakan Impor Saldo Awal."]
    if itype == OPENING:
        notes = ["Identifikasi aset: isi Asset Code Sistem (disarankan), atau Kode Aset Lama, atau Serial Number.",
                 "Baris dengan Nomor Karyawan yang sama digabung menjadi satu dokumen BAST-EXS (tanggal & proyek harus sama).",
                 "Proyek kosong = memakai proyek assignment aktif karyawan. Tanggal format YYYY-MM-DD."]
    for line in notes:
        guide.append([line])
    buf = io.BytesIO()
    wb.save(buf)
    fname = f"template-{'saldo-awal' if itype == OPENING else 'master'}-aset.xlsx"
    return Response(buf.getvalue(), media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{fname}"'})


# ------------------------------------------------------------------ preview
@router.post("/preview", status_code=201)
async def preview(import_type: str = Query(...), file: UploadFile = File(...),
                  ctx: AuthContext = Depends(require_permission("asset_import", "create"))):
    """Upload + validasi batch. HANYA menulis batch & baris audit - tidak ada aset/holding/BAST yang dibuat."""
    itype = import_type.upper()
    if itype not in COLS:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Jenis impor tidak dikenal.")
    if itype == OPENING and not ctx.has_permission("asset_opening", "create"):
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Anda tidak memiliki hak akses untuk membuat saldo awal.")
    rows, meta = await _parse(file, itype)
    await VALIDATORS[itype](ctx, rows)
    counts = _counts(rows)
    prepared = await doc_sequence.prepare(ctx.company_id, "ASSET_IMPORT", actor_id=ctx.user_id)
    batch = {**_stamp(ctx.user_id), "company_id": ctx.company_id, "import_type": itype, "batch_state": PREVIEWED,
             **meta, **counts, "created_count": 0, "opening_count": 0, "uploaded_by": ctx.user_id, "uploaded_at": now(),
             "summary": _group_summary(rows) if itype == OPENING else None}
    async with _atomic("Batch impor bentrok. Silakan ulangi.") as tx:
        async def _taken(c):
            return bool(await tx.select_one("asset_import_batches", {"company_id": ctx.company_id, "batch_number": c}))
        batch["batch_number"] = await doc_sequence.allocate_in_tx(tx, prepared, exists=_taken)
        await tx.insert("asset_import_batches", dict(batch))
        await tx.insert_many("asset_import_rows", [
            {**_stamp(ctx.user_id), "company_id": ctx.company_id, "batch_id": batch["id"], "row_no": r["row_no"],
             "row_class": r["row_class"], "data": {**r["data"], "raw": r["raw"]}, "messages": r["messages"]} for r in rows])
    await log_action(ctx, "preview", "asset_import", batch["id"], batch["batch_number"],
                     after={"type": itype, "file_name": meta["file_name"], **counts}, module="asset")
    return await _batch_out(ctx, batch)


@router.get("")
async def list_batches(import_type: Optional[str] = None, page: int = Query(1, ge=1), limit: int = Query(20, ge=1, le=100),
                       ctx: AuthContext = Depends(require_permission("asset_import", "view"))):
    flt: Dict[str, Any] = {"company_id": ctx.company_id}
    if import_type:
        flt["import_type"] = import_type.upper()
    if (await dscope.get_scope(ctx)).restricted:
        flt["uploaded_by"] = ctx.user_id
    db = get_db()
    total = await db.asset_import_batches.count_documents(flt)
    rows = await db.asset_import_batches.find(flt, NO_ID).sort([("created_at", -1), ("id", 1)]).skip((page - 1) * limit).limit(limit).to_list(limit)
    return _page(total, page, limit, [await _batch_out(ctx, b) for b in rows])


@router.get("/{batch_id}")
async def get_batch(batch_id: str, ctx: AuthContext = Depends(require_permission("asset_import", "view"))):
    return await _batch_out(ctx, await _visible_batch(ctx, batch_id))


@router.get("/{batch_id}/rows")
async def batch_rows(batch_id: str, row_class: Optional[str] = None, page: int = Query(1, ge=1), limit: int = Query(50, ge=1, le=200),
                     ctx: AuthContext = Depends(require_permission("asset_import", "view"))):
    await _visible_batch(ctx, batch_id)
    flt: Dict[str, Any] = {"company_id": ctx.company_id, "batch_id": batch_id}
    if row_class:
        flt["row_class"] = row_class.upper()
    db = get_db()
    total = await db.asset_import_rows.count_documents(flt)
    rows = await db.asset_import_rows.find(flt, NO_ID).sort([("row_no", 1)]).skip((page - 1) * limit).limit(limit).to_list(limit)
    return _page(total, page, limit, [{k: r.get(k) for k in ("id", "row_no", "row_class", "messages", "asset_id", "opening_id")}
                                      | {"data": r.get("data") or {}} for r in rows])


@router.post("/{batch_id}/cancel")
async def cancel_batch(batch_id: str, ctx: AuthContext = Depends(require_permission("asset_import", "create"))):
    b = await _visible_batch(ctx, batch_id)
    async with _atomic("Batch sedang diproses.") as tx:
        cur = await tx.select_one_for_update("asset_import_batches", {"company_id": ctx.company_id, "id": b["id"]})
        if not cur or cur.get("batch_state") != PREVIEWED:
            raise HTTPException(status.HTTP_409_CONFLICT, "Hanya batch berstatus preview yang dapat dibatalkan.")
        await tx.update("asset_import_batches", {"company_id": ctx.company_id, "id": b["id"]},
                        {"batch_state": CANCELLED, **_stamp(ctx.user_id, False)})
    await log_action(ctx, "cancel", "asset_import", b["id"], b["batch_number"], module="asset")
    return {"id": b["id"], "batch_state": CANCELLED}


# ------------------------------------------------------------------ commit
async def _revalidate(ctx: AuthContext, b: Dict[str, Any], stored: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    rows = [{"row_no": r["row_no"], "raw": (r.get("data") or {}).get("raw") or {}, "id": r["id"]} for r in stored]
    await VALIDATORS[b["import_type"]](ctx, rows)
    return rows


async def _bulk_row_update(tx, rows: List[Dict[str, Any]], field: str) -> None:
    t = get_table("asset_import_rows")
    stmt = sa.update(t).where(t.c.id == sa.bindparam("_rid")).values({field: sa.bindparam("_val")})
    params = [{"_rid": r["id"], "_val": r[field]} for r in rows if r.get(field)]
    for i in range(0, len(params), 1000):
        await tx.conn.execute(stmt, params[i:i + 1000])


@router.post("/{batch_id}/commit")
async def commit_batch(batch_id: str, payload: Dict[str, Any] = Body(default={}),
                       ctx: AuthContext = Depends(require_permission("asset_import", "commit"))):
    b = await _visible_batch(ctx, batch_id)
    if b.get("batch_state") != PREVIEWED:
        raise HTTPException(status.HTTP_409_CONFLICT, "Batch ini sudah diproses atau dibatalkan.")
    mode = str((payload or {}).get("mode") or "draft").lower()
    if b["import_type"] == OPENING:
        if not ctx.has_permission("asset_opening", "create"):
            raise HTTPException(status.HTTP_403_FORBIDDEN, "Anda tidak memiliki hak akses untuk membuat saldo awal.")
        if mode == "publish" and not ctx.has_permission("asset_opening", "publish"):
            raise HTTPException(status.HTTP_403_FORBIDDEN, "Anda tidak memiliki hak akses untuk menerbitkan saldo awal.")
        if mode not in ("draft", "publish"):
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Mode commit tidak dikenal (draft / publish).")
    stored = await _load_rows(ctx.company_id, batch_id)
    rows = await _revalidate(ctx, b, stored)
    counts = _counts(rows)
    if counts["error_rows"]:
        # data berubah sejak preview -> simpan hasil validasi terbaru, commit ditolak (tidak ada baris yang dilewati)
        async with _atomic("Batch sedang diproses.") as tx:
            t = get_table("asset_import_rows")
            stmt = sa.update(t).where(t.c.id == sa.bindparam("_rid")).values(row_class=sa.bindparam("_c"), messages=sa.bindparam("_m"))
            await tx.conn.execute(stmt, [{"_rid": r["id"], "_c": r["row_class"], "_m": r["messages"]} for r in rows])
            await tx.update("asset_import_batches", {"company_id": ctx.company_id, "id": batch_id},
                            {**counts, **_stamp(ctx.user_id, False),
                             **({"summary": _group_summary(rows)} if b["import_type"] == OPENING else {})})
        raise HTTPException(status.HTTP_409_CONFLICT, f"Commit dibatalkan: {counts['error_rows']} baris ERROR. Perbaiki file lalu upload ulang.")
    if b["import_type"] == MASTER:
        result = await _commit_master(ctx, b, rows)
    else:
        result = await _commit_opening(ctx, b, rows, mode)
    await log_action(ctx, "commit", "asset_import", batch_id, b["batch_number"],
                     after={"type": b["import_type"], "file_name": b.get("file_name"), "rows": len(rows), **result["audit"]}, module="asset")
    out = await _batch_out(ctx, await _visible_batch(ctx, batch_id))
    out["new_asset_codes"] = result.get("codes", [])[:200]
    out["bast_numbers"] = result.get("basts", [])[:200]
    return out


async def _lock_batch(tx, ctx: AuthContext, batch_id: str) -> None:
    cur = await tx.select_one_for_update("asset_import_batches", {"company_id": ctx.company_id, "id": batch_id})
    if not cur or cur.get("batch_state") != PREVIEWED:
        raise HTTPException(status.HTTP_409_CONFLICT, "Batch ini sudah diproses atau dibatalkan.")


async def _commit_master(ctx: AuthContext, b: Dict[str, Any], rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    cid, actor = ctx.company_id, ctx.user_id
    prepared = await doc_sequence.prepare(cid, svc.ASSET_CODE_SEQ, actor_id=actor)
    ts = now()
    async with _atomic("Konflik impor bersamaan (kode/serial sudah dipakai). Seluruh impor dibatalkan.") as tx:
        await _lock_batch(tx, ctx, b["id"])

        async def _taken(cands):
            t = get_table("assets")
            norms = [svc.norm_code(c) for c in cands]
            res = await tx.conn.execute(sa.select(t.c.asset_code_norm).where(t.c.company_id == cid, t.c.asset_code_norm.in_(norms)))
            used = {x[0] for x in res.fetchall()}
            return {c for c in cands if svc.norm_code(c) in used}
        codes = await doc_sequence.allocate_many_in_tx(tx, prepared, len(rows), taken=_taken)
        assets, events = [], []
        for r, code in zip(rows, codes):
            d = r["data"]
            a = {"id": new_id(), "company_id": cid, "status": "active", "asset_code": code, "asset_code_norm": svc.norm_code(code),
                 "name": d["name"], "category_id": d["category_id"], "unit_id": d["unit_id"], "brand": d.get("brand"),
                 "model": d.get("model"), "serial_number": d.get("serial_number"), "serial_number_norm": d.get("serial_number_norm"),
                 "legacy_code": d.get("legacy_code"), "legacy_code_norm": d.get("legacy_code_norm"),
                 "acquisition_date": None, "acquisition_year": d.get("acquisition_year"), "acquisition_value": d.get("acquisition_value"),
                 "project_id": d.get("project_id"), "work_location_id": d.get("work_location_id"), "condition_id": d["condition_id"],
                 "status_id": d["status_id"], "lifecycle_state": d["lifecycle_state"], "notes": d.get("notes"), "source": "IMPORT",
                 "import_batch_id": b["id"], "row_version": 1, "created_at": ts, "updated_at": ts, "created_by": actor, "updated_by": actor}
            assets.append(a)
            r["asset_id"] = a["id"]
            events.append({"id": new_id(), "company_id": cid, "status": "active", "asset_id": a["id"], "event_type": "CREATED",
                           "event_at": ts, "actor_user_id": actor, "to_state": a["lifecycle_state"], "project_id": a["project_id"],
                           "work_location_id": a["work_location_id"], "condition_id": a["condition_id"], "status_id": a["status_id"],
                           "changes": {"asset_code_mode": "AUTO", "source": "IMPORT", "import_batch": b["batch_number"],
                                       "row_no": r["row_no"]},
                           "created_at": ts, "updated_at": ts, "created_by": actor, "updated_by": actor})
        await tx.insert_many("assets", assets)
        await tx.insert_many("asset_events", events)
        await _bulk_row_update(tx, rows, "asset_id")
        await tx.update("asset_import_batches", {"company_id": cid, "id": b["id"]},
                        {"batch_state": COMMITTED, "created_count": len(assets), "committed_by": actor, "committed_at": ts,
                         "summary": {"first_code": codes[0] if codes else None, "last_code": codes[-1] if codes else None},
                         **_stamp(actor, False)})
    return {"codes": codes, "audit": {"assets_created": len(codes), "first_code": codes[0] if codes else None,
                                      "last_code": codes[-1] if codes else None}}


async def _commit_opening(ctx: AuthContext, b: Dict[str, Any], rows: List[Dict[str, Any]], mode: str) -> Dict[str, Any]:
    cid, actor = ctx.company_id, ctx.user_id
    scope = await dscope.get_scope(ctx)
    groups: Dict[str, List[Dict[str, Any]]] = {}
    for r in rows:
        groups.setdefault(r["data"]["employee_id"], []).append(r)
    ga_pic = (await get_db().users.find_one({"id": actor}, NO_ID) or {}).get("full_name")
    in_use, prepared = (await prepare_opening(ctx, [g[0]["data"]["opening_date"] for g in groups.values()])
                        if mode == "publish" else (None, None))
    ts = now()
    async with _atomic("Konflik saldo awal bersamaan: aset sudah memiliki pemegang aktif. Seluruh impor dibatalkan.") as tx:
        if prepared:
            await lock_sequences(tx, prepared)
        await _lock_batch(tx, ctx, b["id"])
        ops, items, pairs = [], [], []
        for emp_id, group in groups.items():
            d0 = group[0]["data"]
            op = {**_stamp(actor), "company_id": cid, "doc_state": DRAFT, "employee_id": emp_id, "opening_date": d0["opening_date"],
                  "project_id": d0.get("project_id"), "work_location_id": None, "ga_pic_user_id": actor, "ga_pic_name": ga_pic,
                  "manual_number": None, "manual_number_norm": None, "notes": f"Impor {b['batch_number']}",
                  "import_batch_id": b["id"], "row_version": 1}
            its = []
            for n, r in enumerate(group, 1):
                d = r["data"]
                its.append({**_stamp(actor), "company_id": cid, "opening_id": op["id"], "asset_id": d["asset_id"], "line_no": n,
                            "condition_id": d["condition_id"], "accessories": d.get("accessories"), "item_notes": d.get("item_notes")})
                r["opening_id"] = op["id"]
            ops.append(op)
            items += its
            pairs.append((op, its))
        await tx.insert_many("asset_openings", ops)
        await tx.insert_many("asset_opening_items", items)
        basts = await finalize_openings(tx, ctx, scope, pairs, in_use, prepared) if mode == "publish" else []
        await _bulk_row_update(tx, rows, "opening_id")
        await tx.update("asset_import_batches", {"company_id": cid, "id": b["id"]},
                        {"batch_state": COMMITTED, "opening_count": len(ops), "commit_mode": mode.upper(), "committed_by": actor,
                         "committed_at": ts, "summary": {**_group_summary(rows), "assets": len(items), "published": mode == "publish"},
                         **_stamp(actor, False)})
    return {"basts": basts, "audit": {"openings_created": len(ops), "assets": len(items), "mode": mode, "bast_count": len(basts)}}


# ------------------------------------------------------------------ hasil impor (Excel)
@router.get("/{batch_id}/result")
async def download_result(batch_id: str, ctx: AuthContext = Depends(require_permission("asset_import", "view"))):
    b = await _visible_batch(ctx, batch_id)
    if b.get("batch_state") != COMMITTED:
        raise HTTPException(status.HTTP_409_CONFLICT, "Hasil impor tersedia setelah batch diproses (commit).")
    rows = await _load_rows(ctx.company_id, batch_id)
    db, cid = get_db(), ctx.company_id
    wb = Workbook()
    ws = wb.active
    ws.title = "Hasil Impor"
    if b["import_type"] == MASTER:
        ws.append(["Kode Aset Lama", "Asset Code Sistem", "Nama Aset", "Serial Number", "Status"])
        assets = {}
        for chunk in _in_chunks(r.get("asset_id") for r in rows):
            assets.update({a["id"]: a for a in await db.assets.find({"company_id": cid, "id": {"$in": chunk}}, NO_ID).to_list(len(chunk))})
        for r in rows:
            a = assets.get(r.get("asset_id")) or {}
            ws.append([a.get("legacy_code"), a.get("asset_code"), a.get("name"), a.get("serial_number"),
                       svc.LIFECYCLE_STATES.get(a.get("lifecycle_state"), a.get("lifecycle_state"))])
    else:
        ws.append(["Nomor Karyawan", "Nama Karyawan", "Asset Code Sistem", "Nama Aset", "Status Dokumen", "Nomor BAST-EXS"])
        ops = {}
        for chunk in _in_chunks(r.get("opening_id") for r in rows):
            ops.update({o["id"]: o for o in await db.asset_openings.find({"company_id": cid, "id": {"$in": chunk}}, NO_ID).to_list(len(chunk))})
        bast_ids = [o.get("bast_id") for o in ops.values()]
        basts = {}
        for chunk in _in_chunks(bast_ids):
            basts.update({x["id"]: x.get("system_number") for x in await db.asset_basts.find({"company_id": cid, "id": {"$in": chunk}}, NO_ID).to_list(len(chunk))})
        for r in rows:
            d, o = r.get("data") or {}, ops.get(r.get("opening_id")) or {}
            ws.append([d.get("employee_number"), d.get("employee_name"), d.get("asset_code"), d.get("asset_name"),
                       o.get("doc_state"), basts.get(o.get("bast_id"))])
    for c in ws[1]:
        c.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    fname = f"hasil-impor-{b['batch_number'].replace('/', '-')}.xlsx"
    return Response(buf.getvalue(), media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{fname}"'})
